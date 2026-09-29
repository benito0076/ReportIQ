"""Exporta los resultados procesados a un libro de Excel con la misma
estructura de la matriz de procesamiento manual (una hoja por esquema
DH/DNH/NDH/NDNH + hoja de Comparacion con la norma)."""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .models import DIRECCIONES, ESQUEMA_LABELS

_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_HEADER_FONT = Font(bold=True)

_COLUMNAS = [
    ("No. Punto", 10), ("Punto de Monitoreo", 24), ("Direccion", 10),
    ("Inicio", 18), ("Fin", 18),
    ("Lpico dB(A)", 12), ("Lmax dB(A)", 12), ("Lmin dB(A)", 12),
    ("L90 dB(A)", 12), ("LAeq dB(A)", 12), ("LAIeq dB(A)", 12),
    ("Li dB(A)", 10), ("KI", 8), ("KT", 8), ("KS", 8),
    ("Correccion pantalla", 12), ("Correccion", 12),
    ("L90 Corregido", 14), ("LRAeq Corregido", 16), ("Tipo de ajuste", 30),
]


def _fmt(v, nd=1):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(v, nd)
    return v


def _escribir_hoja_esquema(wb, esquema, resultados_proyecto):
    ws = wb.create_sheet(esquema)
    ws.append([c[0] for c in _COLUMNAS])
    for cell, (_, ancho) in zip(ws[1], _COLUMNAS):
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for i, (_, ancho) in enumerate(_COLUMNAS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

    for punto in resultados_proyecto.proyecto.puntos:
        rp = resultados_proyecto.por_punto.get(punto.no_punto, {}).get(esquema)
        if rp is None:
            continue
        for direccion in DIRECCIONES:
            rd = rp.direcciones.get(direccion)
            if rd is None:
                continue
            ws.append([
                punto.no_punto, punto.nombre, direccion,
                str(rd.inicio) if rd.inicio else None,
                str(rd.fin) if rd.fin else None,
                _fmt(rd.lpico), _fmt(rd.lmax), _fmt(rd.lmin),
                _fmt(rd.l90), _fmt(rd.laeq), _fmt(rd.laieq),
                _fmt(rd.li), _fmt(rd.ki, 0), _fmt(rd.kt, 0), _fmt(rd.ks, 0),
                _fmt(rd.correccion_pantalla, 0), _fmt(rd.correccion, 0),
                _fmt(rd.l90_corregido), _fmt(rd.lraeq_corregido), rd.tipo_ajuste,
            ])
        # Fila de nivel resultante del punto para este esquema.
        ws.append([
            punto.no_punto, punto.nombre, "RESULTANTE (prom. energetico)",
            str(rp.inicio) if rp.inicio else None, str(rp.fin) if rp.fin else None,
            None, None, None, None, None, None, None, None, None, None, None, None,
            None, _fmt(rp.lraeq_resultante), None,
        ])
    return ws


def _escribir_hoja_comparacion(wb, resultados_proyecto):
    ws = wb.create_sheet("Comparacion Norma")
    encabezados = [
        "No. Punto", "Punto de Monitoreo", "Sector",
        "LRAeq,1h - Dia Habil", "LRAeq,1h - Dia No Habil",
        "Estandar Diurno dB(A)", "Cumple Diurno (DH)", "Cumple Diurno (DNH)",
        "LRAeq,1h - Nocturno Habil", "LRAeq,1h - Nocturno No Habil",
        "Estandar Nocturno dB(A)", "Cumple Nocturno (NDH)", "Cumple Nocturno (NDNH)",
    ]
    ws.append(encabezados)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    anchos = [10, 24, 45, 14, 16, 14, 12, 12, 16, 18, 14, 12, 12]
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = a

    for no_punto, rc in resultados_proyecto.comparacion.items():
        ws.append([
            no_punto, rc.punto.nombre, rc.punto.sector,
            _fmt(rc.lraeq_dh), _fmt(rc.lraeq_dnh),
            _fmt(rc.estandar_diurno, 0), rc.cumple_dh, rc.cumple_dnh,
            _fmt(rc.lraeq_ndh), _fmt(rc.lraeq_ndnh),
            _fmt(rc.estandar_nocturno, 0), rc.cumple_ndh, rc.cumple_ndnh,
        ])
    return ws


def _escribir_hoja_advertencias(wb, resultados_proyecto):
    if not resultados_proyecto.advertencias:
        return
    ws = wb.create_sheet("Advertencias")
    ws.append(["Punto", "Esquema", "Direccion", "Mensaje"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for a in resultados_proyecto.advertencias:
        ws.append([a.punto, a.esquema, a.direccion, a.mensaje])
    ws.column_dimensions["D"].width = 80


def exportar_resultados(resultados_proyecto, ruta_salida: str):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for esquema, etiqueta in ESQUEMA_LABELS.items():
        tiene_datos = any(
            esquema in resultados_proyecto.por_punto.get(p.no_punto, {})
            for p in resultados_proyecto.proyecto.puntos
        )
        if tiene_datos:
            _escribir_hoja_esquema(wb, esquema, resultados_proyecto)

    _escribir_hoja_comparacion(wb, resultados_proyecto)
    _escribir_hoja_advertencias(wb, resultados_proyecto)

    wb.save(ruta_salida)
    return ruta_salida


# ------------------------------------------------------------------ emision
_COLUMNAS_EMISION = [
    ("No. Punto", 10), ("ID", 6), ("Punto de Monitoreo", 26), ("Inicio", 20), ("Fin", 20),
    ("Lpico dB(A)", 11), ("Lmax dB(A)", 11), ("Lmin dB(A)", 11), ("L90 dB(A)", 11), ("LAeq dB(A)", 11),
    ("LAIeq dB(A)", 11), ("Li", 8), ("KI", 6), ("KT", 6), ("KR", 6), ("KS", 6), ("Correccion pantalla", 11),
    ("Correccion", 11), ("LRAeq Corregido", 13), ("L90 Corregido", 13), ("LRAeq Residual", 13),
    ("Origen residual", 18), ("Emision LRAeq,1h", 14), ("Incertidumbre", 12), ("Diferencia LRAeq - residual", 14),
    ("Estandar", 10), ("Cumple", 8), ("Del orden del residual (<= 3 dB)", 16),
]


def _encabezado(ws, columnas):
    ws.append([c[0] for c in columnas])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for i, (_, ancho) in enumerate(columnas, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho


def exportar_resultados_emision(resultados, ruta_salida: str):
    """Libro de resultados de emision: una hoja por jornada con la memoria de
    calculo (como las hojas DH/NH de FP-007), comparacion con la norma,
    barrido y advertencias."""
    from .models import ESQUEMAS

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    puntos = resultados.proyecto.puntos
    for esquema in ESQUEMAS:
        filas = [(p, r) for p in puntos for r in [resultados.resultado(p, esquema)] if r is not None]
        if not filas:
            continue
        ws = wb.create_sheet(esquema)
        _encabezado(ws, _COLUMNAS_EMISION)
        for p, r in filas:
            m = r.medicion
            ws.append([
                p.no_punto, f"P{p.no_punto}", p.nombre, str(m.inicio) if m.inicio else None,
                str(m.fin) if m.fin else None, _fmt(m.lpico), _fmt(m.lmax), _fmt(m.lmin), _fmt(m.l90),
                _fmt(m.laeq), _fmt(m.laieq), _fmt(m.li), _fmt(m.ki, 0), _fmt(m.kt, 0), 0, _fmt(m.ks, 0),
                _fmt(m.correccion_pantalla, 0), _fmt(m.correccion, 0), _fmt(m.lraeq_corregido),
                _fmt(m.l90_corregido), _fmt(r.residual),
                "L90 corregido" if r.residual_es_l90 else "Medido (fuente apagada)",
                _fmt(r.emision), p.incertidumbre, _fmt(r.diferencia), r.estandar, r.cumple,
                "Si" if r.del_orden_del_residual else "No",
            ])
    ws = wb.create_sheet("Comparacion Norma")
    columnas = [("No. Punto", 10), ("ID", 6), ("Punto de Monitoreo", 28)]
    for esquema in ESQUEMAS:
        columnas += [(f"Emision {esquema}", 14), (f"Estandar {esquema}", 12), (f"Cumple {esquema}", 10)]
    _encabezado(ws, columnas)
    for p in puntos:
        fila = [p.no_punto, f"P{p.no_punto}", p.nombre]
        for esquema in ESQUEMAS:
            r = resultados.resultado(p, esquema)
            fila += [_fmt(r.emision), r.estandar, r.cumple] if r else [None, None, None]
        ws.append(fila)
    if resultados.barrido:
        ws = wb.create_sheet("Barrido")
        _encabezado(ws, [("ID Punto", 16), ("Condicion de la fuente", 20), ("Inicio", 20), ("Fin", 20),
                         ("Leq dB(A)", 12), ("Seleccionado", 12)])
        for b in resultados.barrido:
            ws.append([b.nombre, b.condicion, str(b.inicio) if b.inicio else None, str(b.fin) if b.fin else None,
                       _fmt(b.leq), "Si" if b.seleccionado else "No"])
    _escribir_hoja_advertencias(wb, resultados)
    wb.save(ruta_salida)
    return ruta_salida
