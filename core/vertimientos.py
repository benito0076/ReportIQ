"""Vertimientos (matriz agua): lectura de insumos.

    - Reportes de resultados del laboratorio (formato FT-024, PDF con texto):
      ensayo, metodo, LCM, fecha de analisis, resultado, incertidumbre y unidad,
      mas los datos de la muestra (numero, punto, fecha y hora de muestreo...).
    - Plantilla FP-004 (datos de campo del muestreo compuesto): una hoja por
      punto de muestreo con la hora, el aforo volumetrico (volumen y tiempo),
      pH, temperatura, oxigeno disuelto, conductividad y solidos sedimentables
      (cada uno con su duplicado) y el tamano de la muestra compuesta.

Los nombres de los ensayos del laboratorio se asocian a un catalogo de
parametros (PARAMETROS) con clave estable, que es la que usan los limites de la
Resolucion 0631 de 2015 y el informe.
"""
from __future__ import annotations

import re
import unicodedata
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Optional

import openpyxl


class ErrorVertimientos(Exception):
    """Archivo de entrada que no se puede interpretar."""


# ----------------------------------------------------------------- parametros
@dataclass(frozen=True)
class Parametro:
    clave: str
    nombre: str  # como se escribe en el informe
    grupo: str  # capitulo del informe: insitu, fisicoquimico, hap, nitrogenado, metal, otro
    patron: str  # regex sobre el nombre del ensayo normalizado (mayusculas, sin tildes)


def _p(clave, nombre, grupo, patron):
    return Parametro(clave, nombre, grupo, patron)


PARAMETROS: tuple[Parametro, ...] = (
    _p("aceites_grasas", "Aceites y Grasas", "fisicoquimico", r"^ACEITES Y GRASAS"),
    _p("acidez", "Acidez", "fisicoquimico", r"^ACIDEZ"),
    _p("alcalinidad", "Alcalinidad", "fisicoquimico", r"^ALCALINIDAD"),
    _p("aluminio", "Aluminio", "metal", r"^ALUMINIO"),
    _p("arsenico", "Arsénico", "metal", r"^ARSENICO"),
    _p("bario", "Bario", "metal", r"^BARIO"),
    _p("berilio", "Berilio", "metal", r"^BERILIO"),
    _p("boro", "Boro", "metal", r"^BORO"),
    _p("cadmio", "Cadmio", "metal", r"^CADMIO"),
    _p("calcio", "Calcio", "metal", r"^CALCIO"),
    _p("cianuro", "Cianuro Total", "otro", r"^CIANURO"),
    _p("cloruros", "Cloruros", "fisicoquimico", r"^CLORUROS?\b"),
    _p("cobalto", "Cobalto", "metal", r"^COBALTO"),
    _p("cobre", "Cobre", "metal", r"^COBRE"),
    _p("color_436", "Color Verdadero (Real) a 436 nm", "fisicoquimico", r"^COLOR .*436"),
    _p("color_525", "Color Verdadero (Real) a 525 nm", "fisicoquimico", r"^COLOR .*525"),
    _p("color_620", "Color Verdadero (Real) a 620 nm", "fisicoquimico", r"^COLOR .*620"),
    _p("conductividad_max", "Conductividad Máxima", "insitu", r"^CONDUCTIVIDAD.*MAXIM"),
    _p("conductividad_min", "Conductividad Mínima", "insitu", r"^CONDUCTIVIDAD.*MINIM"),
    _p("conductividad", "Conductividad", "insitu", r"^CONDUCTIVIDAD"),
    _p("cromo_hexavalente", "Cromo Hexavalente", "metal", r"^CROMO HEXAVALENTE"),
    _p("cromo", "Cromo", "metal", r"^CROMO"),
    _p("dbo5", "Demanda Bioquímica de Oxígeno (DBO5)", "fisicoquimico", r"^DEMANDA BIOQUIMICA|^DBO"),
    _p("dqo", "Demanda Química de Oxígeno (DQO)", "fisicoquimico", r"^DEMANDA QUIMICA|^DQO"),
    _p("dureza_calcica", "Dureza Cálcica", "fisicoquimico", r"^DUREZA CALCICA"),
    _p("dureza_total", "Dureza Total", "fisicoquimico", r"^DUREZA TOTAL"),
    _p("estano", "Estaño", "metal", r"^ESTANO"),
    _p("fenoles", "Fenoles", "fisicoquimico", r"^FENOLES"),
    _p("formaldehido", "Formaldehído", "otro", r"^FORMALDEHIDO"),
    _p("ortofosfatos", "Fósforo Reactivo Total (Leído como Ortofosfatos)", "fisicoquimico",
       r"^FOSFORO REACTIVO|^ORTOFOSFATOS"),
    _p("fosforo_total", "Fósforo Total", "fisicoquimico", r"^FOSFORO TOTAL"),
    _p("grasas_aceites", "Grasas y Aceites", "fisicoquimico", r"^GRASAS Y ACEITES"),
    _p("hidrocarburos_totales", "Hidrocarburos Totales (HTP)", "fisicoquimico", r"^HIDROCARBUROS TOTALES"),
    _p("hidrocarburos", "Hidrocarburos", "fisicoquimico", r"^HIDROCARBUROS$|^HIDROCARBUROS \("),
    _p("hierro", "Hierro", "metal", r"^HIERRO"),
    _p("litio", "Litio", "metal", r"^LITIO"),
    _p("magnesio", "Magnesio", "metal", r"^MAGNESIO"),
    _p("manganeso", "Manganeso", "metal", r"^MANGANESO"),
    _p("mercurio", "Mercurio", "metal", r"^MERCURIO"),
    _p("molibdeno", "Molibdeno", "metal", r"^MOLIBDENO"),
    _p("niquel", "Níquel", "metal", r"^NIQUEL"),
    _p("nitratos", "Nitratos", "nitrogenado", r"^NITRATOS?\b"),
    _p("nitritos", "Nitritos", "nitrogenado", r"^NITRITOS?\b"),
    _p("nitrogeno_amoniacal", "Nitrógeno Amoniacal", "nitrogenado", r"^NITROGENO AMONIACAL|^AMONIO"),
    _p("nitrogeno_total", "Nitrógeno Total", "nitrogenado", r"^NITROGENO TOTAL"),
    _p("ph_max", "pH Máximo", "insitu", r"^PH MAXIM"),
    _p("ph_min", "pH Mínimo", "insitu", r"^PH MINIM"),
    _p("plata", "Plata", "metal", r"^PLATA"),
    _p("plomo", "Plomo", "metal", r"^PLOMO"),
    _p("potasio", "Potasio", "metal", r"^POTASIO"),
    _p("selenio", "Selenio", "metal", r"^SELENIO"),
    _p("sodio", "Sodio", "metal", r"^SODIO"),
    _p("solidos_sedimentables_max", "Sólidos Sedimentables Máximo", "insitu", r"^SOLIDOS SEDIMENTABLES.*MAXIM"),
    _p("solidos_sedimentables_min", "Sólidos Sedimentables Mínimo", "insitu", r"^SOLIDOS SEDIMENTABLES.*MINIM"),
    _p("solidos_sedimentables", "Sólidos Sedimentables", "insitu", r"^SOLIDOS SEDIMENTABLES"),
    _p("sst", "Sólidos Suspendidos Totales", "fisicoquimico", r"^SOLIDOS SUSPENDIDOS"),
    _p("sdt", "Sólidos Disueltos Totales", "fisicoquimico", r"^SOLIDOS DISUELTOS"),
    _p("sulfatos", "Sulfatos", "fisicoquimico", r"^SULFATOS?\b"),
    _p("sulfuros", "Sulfuros", "fisicoquimico", r"^SULFUROS?\b"),
    _p("saam", "Surfactantes Aniónicos (SAAM)", "fisicoquimico", r"^SURFACTANTES|^SAAM|^TENSOACTIVOS"),
    _p("temperatura_max", "Temperatura Máxima", "insitu", r"^TEMPERATURA.*MAXIM"),
    _p("temperatura_min", "Temperatura Mínima", "insitu", r"^TEMPERATURA.*MINIM"),
    _p("vanadio", "Vanadio", "metal", r"^VANADIO"),
    _p("zinc", "Zinc", "metal", r"^ZINC"),
)

# Hidrocarburos aromaticos policiclicos: cada compuesto es un parametro del grupo "hap".
_HAP = re.compile(r"^HAPS?\s*-\s*(.+)$")


def normalizar(texto: str) -> str:
    t = unicodedata.normalize("NFKD", str(texto or ""))
    t = "".join(c for c in t if not unicodedata.combining(c)).upper()
    t = re.sub(r"\(\*+\)|\*+", "", t)
    return re.sub(r"\s+", " ", t).strip()


def parametro_de(ensayo: str) -> Parametro:
    """Parametro del catalogo para el nombre de ensayo del laboratorio; si no esta
    catalogado se crea uno generico (grupo "otro") con el nombre en forma de titulo."""
    n = normalizar(ensayo)
    m = _HAP.match(n)
    if m:
        compuesto = m.group(1).strip()
        return Parametro("hap_" + re.sub(r"[^a-z0-9]+", "_", compuesto.lower()).strip("_"),
                         _titulo(compuesto), "hap", "")
    for p in PARAMETROS:
        if re.search(p.patron, n):
            return p
    return Parametro(re.sub(r"[^a-z0-9]+", "_", n.lower()).strip("_"), _titulo(ensayo), "otro", "")


_MINUSCULAS = {"de", "del", "la", "las", "los", "y", "a", "como", "en", "el"}


def _titulo(texto: str) -> str:
    palabras = re.sub(r"\s+", " ", str(texto).replace("*", "")).strip().lower().split(" ")
    salida = []
    for i, w in enumerate(palabras):
        salida.append(w if i and w in _MINUSCULAS else w[:1].upper() + w[1:])
    t = " ".join(salida)
    return re.sub(r"\(([a-z])", lambda m: "(" + m.group(1).upper(), t)


# --------------------------------------------------------------- laboratorio
@dataclass
class ResultadoLab:
    ensayo: str  # nombre tal como lo reporta el laboratorio
    metodo: str
    lcm: str
    fecha_analisis: Optional[date]
    reporte: str  # resultado tal como se reporta ("26,4", "<0,00100")
    incertidumbre: str
    unidad: str
    subcontratado: bool = False
    parametro: Optional[Parametro] = None

    @property
    def menor_que_lcm(self) -> bool:
        return self.reporte.strip().startswith("<")

    @property
    def valor(self) -> Optional[float]:
        """Valor numerico (para "<x" devuelve x, el limite de cuantificacion)."""
        return numero(self.reporte)


@dataclass
class InformeLaboratorio:
    numero: str = ""  # REPORTE DE RESULTADOS No.
    muestra: str = ""  # NUMERO DE MUESTRA
    punto: str = ""  # PUNTO DE MUESTREO
    procedencia: str = ""
    tipo_muestreo: str = ""  # Compuesta / Puntual
    tipo_muestra: str = ""  # Agua Residual no Domestica...
    plan: str = ""  # PLAN Y METODO DE MUESTREO
    fecha_muestreo: Optional[datetime] = None
    fecha_recepcion: Optional[datetime] = None
    fecha_emision: Optional[date] = None
    municipio: str = ""
    departamento: str = ""
    cliente: dict = field(default_factory=dict)  # razon_social, nit, direccion, contacto, email, telefono
    resultados: list[ResultadoLab] = field(default_factory=list)
    archivo: str = ""

    def resultado(self, clave: str) -> Optional[ResultadoLab]:
        return next((r for r in self.resultados if r.parametro and r.parametro.clave == clave), None)


def numero(texto) -> Optional[float]:
    """Numero en formato colombiano ("1.234,5", "0,00100", "<4,00"); None si no lo es."""
    if texto is None:
        return None
    if isinstance(texto, (int, float)):
        return float(texto)
    t = str(texto).strip().lstrip("<>≤≥").strip()
    if not re.fullmatch(r"-?[\d.,]+", t):
        return None
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    elif t.count(".") > 1:
        t = t.replace(".", "")
    try:
        return float(t)
    except ValueError:
        return None


_ETIQUETAS = {
    "RAZON SOCIAL": "razon_social", "NIT": "nit", "DIRECCION": "direccion", "CIUDAD/MUNICIPIO": "municipio",
    "PROYECTO": "proyecto", "SOLICITADO POR": "contacto", "E-MAIL": "email", "TELEFONO": "telefono",
    "NUMERO DE MUESTRA": "muestra", "PUNTO DE MUESTREO": "punto", "PROCEDENCIA": "procedencia",
    "DEPARTAMENTO": "departamento", "TIPO DE MUESTRA": "tipo_muestra", "TIPO DE MUESTREO": "tipo_muestreo",
    "PLAN Y METODO DE MUESTREO": "plan", "MUESTREADO POR": "muestreado_por",
    "FECHA Y HORA DE MUESTREO": "fecha_muestreo", "FECHA Y HORA DE RECEPCION": "fecha_recepcion",
    "FECHA DE EMISION DEL INFORME": "fecha_emision",
}
_FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}$")
# Lineas sueltas dentro de la tabla que no son continuacion de un ensayo (encabezados de pagina).
_NO_CONTINUACION = {"RESULTADOS", "ANALISIS", "ENSAYOS SUBCONTRATADOS", "OBSERVACIONES:"}


def _texto_pdf(ruta: str) -> list[str]:
    try:
        import pypdf
    except ImportError as exc:  # pragma: no cover
        raise ErrorVertimientos("Falta la libreria pypdf para leer los reportes de laboratorio.") from exc
    try:
        lector = pypdf.PdfReader(ruta)
        lineas = []
        for pagina in lector.pages:
            lineas += (pagina.extract_text(extraction_mode="layout") or "").splitlines()
    except Exception as exc:  # noqa: BLE001
        raise ErrorVertimientos(f"No se pudo leer el PDF: {exc}") from exc
    if not any(line.strip() for line in lineas):
        raise ErrorVertimientos("El PDF no tiene texto (¿es una imagen escaneada?). Suba el reporte "
                                "de resultados generado por el laboratorio, no una copia escaneada.")
    return lineas


def _fecha_hora(texto: str) -> Optional[datetime]:
    m = re.search(r"(\d{4}-\d{2}-\d{2})(?:\s+(\d{1,2}):(\d{2}))?", texto or "")
    if not m:
        return None
    d = datetime.strptime(m.group(1), "%Y-%m-%d")
    return d.replace(hour=int(m.group(2)), minute=int(m.group(3))) if m.group(2) else d


def leer_informe_laboratorio(ruta: str) -> InformeLaboratorio:
    """Lee un reporte de resultados del laboratorio (PDF con texto, formato FT-024)."""
    lineas = _texto_pdf(ruta)
    inf = InformeLaboratorio(archivo=ruta)
    datos: dict[str, str] = {}
    en_resultados = subcontratados = False
    ultimo: Optional[ResultadoLab] = None
    col_metodo = 0
    for linea in lineas:
        limpia = linea.strip()
        n = normalizar(limpia)
        if not limpia:
            continue
        m = re.search(r"REPORTE DE RESULTADOS NO\.?\s*([\w-]+)", n)
        if m and not inf.numero:
            inf.numero = m.group(1)
            continue
        if n.startswith("ENSAYO REALIZADO"):
            en_resultados = True
            continue
        if n.startswith("ENSAYOS SUBCONTRATADOS"):
            subcontratados = True
            continue
        if n.startswith(("LCM:", "OBSERVACIONES", "ACLARACIONES")):
            en_resultados = False
            ultimo = None
            continue
        if en_resultados:
            tokens = re.split(r"\s{2,}", limpia)
            if len(tokens) >= 7 and _FECHA.match(tokens[-4]):
                metodo = tokens[-6]
                ultimo = ResultadoLab(
                    ensayo=" ".join(tokens[:-6]), metodo=metodo, lcm=tokens[-5],
                    fecha_analisis=datetime.strptime(tokens[-4], "%Y-%m-%d").date(),
                    reporte=tokens[-3], incertidumbre=tokens[-2], unidad=tokens[-1],
                    subcontratado=subcontratados or "*" in tokens[0])
                inf.resultados.append(ultimo)
                col_metodo = linea.find(metodo)
                continue
            if ultimo is not None and len(tokens) == 1 and n not in _NO_CONTINUACION and not n.startswith("FT-"):
                # Continuacion del nombre del ensayo o del metodo (texto en dos lineas).
                sangria = len(linea) - len(linea.lstrip())
                if col_metodo and sangria >= col_metodo - 4:
                    ultimo.metodo = f"{ultimo.metodo} {limpia}"
                else:
                    ultimo.ensayo = f"{ultimo.ensayo} {limpia}"
            continue
        # Encabezado: "ETIQUETA:   valor   ETIQUETA2:   valor2"
        tokens = re.split(r"\s{2,}", limpia)
        for i, tok in enumerate(tokens[:-1]):
            clave = _ETIQUETAS.get(normalizar(tok).rstrip(":").strip())
            if clave and clave not in datos:
                siguiente = tokens[i + 1]
                if normalizar(siguiente).rstrip(":") not in _ETIQUETAS:
                    datos[clave] = siguiente.strip()
    for r in inf.resultados:
        r.ensayo = re.sub(r"\s+", " ", r.ensayo).strip()
        r.metodo = re.sub(r"\s+", " ", r.metodo).strip()
        r.parametro = parametro_de(r.ensayo)
    if not inf.resultados:
        raise ErrorVertimientos("No se encontraron resultados en el reporte del laboratorio (se esperaba la "
                                "tabla ENSAYO REALIZADO / MÉTODO / LCM / FECHA / REPORTE / INCERTIDUMBRE / UNIDAD).")
    inf.muestra = datos.get("muestra", "") or inf.numero
    inf.punto = datos.get("punto", "")
    inf.procedencia = datos.get("procedencia", "")
    inf.tipo_muestreo = datos.get("tipo_muestreo", "")
    inf.tipo_muestra = datos.get("tipo_muestra", "")
    inf.plan = datos.get("plan", "")
    inf.municipio = datos.get("municipio", "")
    inf.departamento = datos.get("departamento", "")
    inf.fecha_muestreo = _fecha_hora(datos.get("fecha_muestreo", ""))
    inf.fecha_recepcion = _fecha_hora(datos.get("fecha_recepcion", ""))
    emision = _fecha_hora(datos.get("fecha_emision", ""))
    inf.fecha_emision = emision.date() if emision else None
    inf.cliente = {k: datos[k] for k in ("razon_social", "nit", "direccion", "contacto", "email", "telefono",
                                          "proyecto") if k in datos}
    return inf


# ------------------------------------------------------------ datos de campo
@dataclass
class MedicionCampo:
    hora: Optional[time]
    volumen_ml: Optional[float]
    tiempo_s: Optional[float]
    ph: Optional[float] = None
    ph_dup: Optional[float] = None
    temperatura: Optional[float] = None
    temperatura_dup: Optional[float] = None
    oxigeno: Optional[float] = None
    oxigeno_dup: Optional[float] = None
    conductividad: Optional[float] = None
    conductividad_dup: Optional[float] = None
    sedimentables: Optional[str] = None  # texto: puede ser "<0,1"
    sedimentables_dup: Optional[str] = None

    @property
    def caudal_mls(self) -> Optional[float]:
        if self.volumen_ml is None or not self.tiempo_s:
            return None
        return self.volumen_ml / self.tiempo_s


@dataclass
class PuntoCampo:
    hoja: str
    titulo: str  # "ENTRADA DEL SISTEMA DE TRATAMIENTO"
    mediciones: list[MedicionCampo] = field(default_factory=list)
    tamano_muestra_ml: Optional[float] = None

    @property
    def suma_caudales(self) -> float:
        return sum(m.caudal_mls for m in self.mediciones if m.caudal_mls is not None)

    def fraccion(self, m: MedicionCampo) -> Optional[float]:
        s = self.suma_caudales
        return m.caudal_mls / s if m.caudal_mls is not None and s else None

    def alicuota_ml(self, m: MedicionCampo) -> Optional[float]:
        f = self.fraccion(m)
        return f * self.tamano_muestra_ml if f is not None and self.tamano_muestra_ml else None


def _num_celda(v) -> Optional[float]:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return numero(v)


def _hora(v) -> Optional[time]:
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    if isinstance(v, (int, float)) and 0 <= v < 1:
        minutos = round(v * 24 * 60)
        return time(minutos // 60, minutos % 60)
    m = re.match(r"^\s*(\d{1,2})[:.](\d{2})", str(v or ""))
    return time(int(m.group(1)), int(m.group(2))) if m else None


def _texto_celda(v, formato: str = "") -> Optional[str]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        texto = f"{v:g}".replace(".", ",")
        # Formato numerico "\<0.0" (Excel muestra "<0,1"): resultado bajo el limite.
        return "<" + texto if formato.replace('"', "").lstrip("\\").startswith("<") else texto
    return str(v).strip()


# Columnas de la hoja de punto de la FP-004 V02 (B=2 ... O=15) y fila de encabezado.
_COLUMNAS = {"hora": 2, "volumen": 3, "tiempo": 4, "ph": 6, "ph_dup": 7, "temp": 8, "temp_dup": 9,
             "od": 10, "od_dup": 11, "cond": 12, "cond_dup": 13, "ss": 14, "ss_dup": 15}


def _es_hoja_de_punto(ws) -> bool:
    enc = [normalizar(ws.cell(4, c).value) for c in range(2, 6)]
    return enc[0] == "HORA" and enc[1].startswith("VOLUMEN") and "CAUDAL" in enc[3]


def leer_fp004(ruta: str) -> list[PuntoCampo]:
    """Hojas de punto de la FP-004: las que tienen en la fila 4 Hora | Volumen | Tiempo | Caudal."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            wb = openpyxl.load_workbook(ruta, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ErrorVertimientos(f"No se pudo abrir la plantilla FP-004: {exc}") from exc
    puntos = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or not _es_hoja_de_punto(ws):
            continue
        punto = PuntoCampo(hoja=ws.title.strip(), titulo=_texto_celda(ws.cell(3, 2).value) or ws.title.strip())
        for r in range(6, min(ws.max_row, 200) + 1):
            etiqueta = normalizar(ws.cell(r, 2).value)
            if etiqueta.startswith("SUMA"):
                # "Tamano de muestra (mL)" esta junto a las filas de resumen (P24 -> Q24).
                for rr in range(r, r + 4):
                    if "TAMANO DE MUESTRA" in normalizar(ws.cell(rr, 16).value):
                        punto.tamano_muestra_ml = _num_celda(ws.cell(rr, 17).value)
                break
            celdas = {k: ws.cell(r, c) for k, c in _COLUMNAS.items()}
            fila = {k: c.value for k, c in celdas.items()}
            hora = _hora(fila["hora"])
            if hora is None:
                continue
            m = MedicionCampo(
                hora=hora, volumen_ml=_num_celda(fila["volumen"]), tiempo_s=_num_celda(fila["tiempo"]),
                ph=_num_celda(fila["ph"]), ph_dup=_num_celda(fila["ph_dup"]),
                temperatura=_num_celda(fila["temp"]), temperatura_dup=_num_celda(fila["temp_dup"]),
                oxigeno=_num_celda(fila["od"]), oxigeno_dup=_num_celda(fila["od_dup"]),
                conductividad=_num_celda(fila["cond"]), conductividad_dup=_num_celda(fila["cond_dup"]),
                sedimentables=_texto_celda(fila["ss"], celdas["ss"].number_format),
                sedimentables_dup=_texto_celda(fila["ss_dup"], celdas["ss_dup"].number_format))
            if any(v is not None for v in (m.volumen_ml, m.ph, m.temperatura, m.conductividad)):
                punto.mediciones.append(m)
        if punto.mediciones:
            puntos.append(punto)
    if not puntos:
        raise ErrorVertimientos("La plantilla no tiene hojas de punto con datos (se esperaba la tabla "
                                "Hora | Volumen (mL) | Tiempo (s) | Caudal (mL/s) en la fila 4).")
    return puntos
