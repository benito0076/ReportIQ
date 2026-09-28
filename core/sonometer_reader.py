"""Lector de memorias exportadas por sonometros Larson Davis (SoundTrack LxT y similares).

El archivo exportado (.xlsx) trae, entre otras, las hojas:
    - "Resumen": valores globales de la sesion (LAeq, LAIeq, L90, Lmax, Lmin, Lpico, horas).
    - "OBA": espectro en tercios de octava (1/3 LAeq) y otros datos.

Esta lectura es tolerante a la codificacion de caracteres especiales (los archivos
suelen tener tildes/simbolos mal codificados segun la version del software del equipo).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Optional

import openpyxl

# Bandas de tercio de octava tal como aparecen en la hoja OBA (Hz), en orden.
BANDAS_TERCIO_OCTAVA = [
    "6,3", "8,0", "10,0", "12,5", "16,0", "20,0", "25,0", "31,5", "40,0", "50,0",
    "63,0", "80,0", "100", "125", "160", "200", "250", "315", "400", "500",
    "630", "800", "1000", "1250", "1600", "2000", "2500", "3150", "4000", "5000",
    "6300", "8000", "10000", "12500", "16000", "20000",
]

# Filas de interes en "Resumen": clave normalizada -> etiquetas exactas conocidas.
# El caracter de ponderacion "A" a veces se exporta mal codificado (aparece como
# un glifo distinto, p.ej. la letra griega omega) segun la fuente usada por el
# software del sonometro; por eso ademas se intenta con los patrones regex de
# _RESUMEN_PATRONES, que aceptan cualquier caracter en esa posicion.
_RESUMEN_CLAVES = {
    "inicio": ["Tiempo de inicio"],
    "fin": ["Para el tiempo"],
    "lpico": ["LApk (dB)", "Lpk (dB)", "LZpk (dB)"],
    "lmax": ["LASmax (dB)", "LSmax (dB)"],
    "lmin": ["LASmin (dB)", "LSmin (dB)"],
    "l90": ["LAS 90,0 (dB)"],
    "laeq": ["LAeq (dB)"],
    "laieq": ["LAIeq (dB)"],
}

_RESUMEN_PATRONES = {
    "lpico": re.compile(r"^L.pk \(dB\)$"),
    "lmax": re.compile(r"^L.Smax \(dB\)$"),
    "lmin": re.compile(r"^L.Smin \(dB\)$"),
    "laeq": re.compile(r"^L.eq \(dB\)$"),
}


def _norm(s):
    if s is None:
        return ""
    return str(s).strip()


class ErrorLecturaMemoria(Exception):
    pass


@dataclass
class DatosMemoria:
    ruta: str
    inicio: Optional[object] = None
    fin: Optional[object] = None
    lpico: Optional[float] = None
    lmax: Optional[float] = None
    lmin: Optional[float] = None
    l90: Optional[float] = None
    laeq: Optional[float] = None
    laieq: Optional[float] = None
    numero_serie: Optional[str] = None
    modelo: Optional[str] = None
    espectro_laeq: dict = field(default_factory=dict)  # banda(str) -> dB


def leer_memoria(ruta: str) -> DatosMemoria:
    """Lee una memoria exportada de sonometro y extrae los valores necesarios."""
    try:
        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ErrorLecturaMemoria(f"No se pudo abrir el archivo '{ruta}': {exc}") from exc

    if "Resumen" not in wb.sheetnames:
        raise ErrorLecturaMemoria(f"El archivo '{ruta}' no tiene una hoja 'Resumen'.")

    ws = wb["Resumen"]
    valores = {}
    numero_serie = None
    modelo = None
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        etiqueta = _norm(row[0])
        if not etiqueta:
            continue
        for clave, alias in _RESUMEN_CLAVES.items():
            if clave in valores:
                continue
            if etiqueta in alias:
                valores[clave] = row[1] if len(row) > 1 else None
        for clave, patron in _RESUMEN_PATRONES.items():
            if clave in valores:
                continue
            if patron.match(etiqueta):
                valores[clave] = row[1] if len(row) > 1 else None

        # Numero de serie del equipo: dos formatos conocidos.
        #  - SoundTrack LxT: fila "Numero de Serie" con el valor en la col. B.
        #  - SoundExpert 821: fila "Medidor" con Modelo en col. B y Serie en col. C.
        if numero_serie is None and re.match(r"^N.mero de Serie$", etiqueta):
            numero_serie = row[1] if len(row) > 1 else None
        if numero_serie is None and etiqueta == "Medidor" and len(row) > 2:
            modelo = row[1]
            numero_serie = row[2]
        if modelo is None and etiqueta == "Modelo" and len(row) > 1:
            modelo = row[1]

    datos = DatosMemoria(ruta=ruta)
    datos.inicio = valores.get("inicio")
    datos.fin = valores.get("fin")
    datos.numero_serie = _norm(numero_serie) or None
    datos.modelo = _norm(modelo) or None
    for campo in ("lpico", "lmax", "lmin", "l90", "laeq", "laieq"):
        val = valores.get(campo)
        try:
            setattr(datos, campo, float(val) if val is not None else None)
        except (TypeError, ValueError):
            setattr(datos, campo, None)

    if "OBA" in wb.sheetnames:
        datos.espectro_laeq = _leer_espectro(wb["OBA"])

    wb.close()
    return datos


def _leer_espectro(ws) -> dict:
    """Extrae la fila '1/3 LAeq' de la hoja OBA, indexada por banda en Hz."""
    filas = list(ws.iter_rows(values_only=True))
    fila_frecuencias = None
    fila_laeq = None
    for row in filas:
        if not row:
            continue
        etiqueta = _norm(row[0])
        if etiqueta == "Frecuencia (Hz)" and fila_frecuencias is None:
            fila_frecuencias = row
        if etiqueta == "1/3 LAeq" and fila_laeq is None:
            fila_laeq = row
        if fila_frecuencias is not None and fila_laeq is not None:
            break

    espectro = {}
    if fila_frecuencias is None or fila_laeq is None:
        return espectro

    for i, banda in enumerate(fila_frecuencias[1:], start=1):
        banda_str = _norm(banda)
        if not banda_str:
            continue
        try:
            val = float(fila_laeq[i])
        except (TypeError, ValueError, IndexError):
            val = None
        espectro[banda_str] = val
    return espectro
