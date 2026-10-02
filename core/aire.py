"""Calidad del aire (Resolucion 2254 de 2017).

Lee las plantillas de procesamiento que diligencia el laboratorio (una hoja
CA-n / ESTACION n por estacion) y recalcula todo con sus mismas formulas:

    FP-031  PM10 con equipos de alto volumen (Hi-Vol)
    FP-032  PM2.5 con equipos de bajo volumen
    FP-033  SO2 manual (pararrosanilina)
    FP-035  Compuestos organicos volatiles (COV)
    FP-021  Equipos automaticos: CO, NO2 y O3 (datos horarios)

Se recalcula a partir de los datos de campo (no de los resultados guardados en
el Excel), de modo que la plantilla puede venir con o sin formulas calculadas.
Unica diferencia deliberada con las plantillas: en la FP-032 el caudal de
referencia usa la temperatura ambiente (la plantilla V03 usaba el caudal).
"""
from __future__ import annotations

import math
import re
import statistics
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Optional

import openpyxl

PM10, PM25, SO2, NO2, CO, O3, COV = "PM10", "PM2.5", "SO2", "NO2", "CO", "O3", "COV"
CONTAMINANTES = (PM10, PM25, SO2, NO2, CO, O3, COV)
MANUALES = (PM10, PM25, SO2)  # una muestra de 24 h por dia
AUTOMATICOS = (CO, NO2, O3)  # datos horarios

NOMBRES = {
    PM10: "Partículas Menores a 10 Micras (PM10)",
    PM25: "Partículas Menores a 2,5 Micras (PM2,5)",
    SO2: "Dióxido de Azufre (SO2)",
    NO2: "Dióxido de Nitrógeno (NO2)",
    CO: "Monóxido de Carbono (CO)",
    O3: "Ozono (O3)",
    COV: "Compuestos Orgánicos Volátiles (COV)",
}

# Niveles maximos permisibles (Res. 2254 de 2017, Art. 2), en µg/m3 a 25 °C y 760 mmHg.
LIMITES = {
    PM10: {"24h": 75, "anual": 50},
    PM25: {"24h": 37, "anual": 25},
    SO2: {"24h": 50, "1h": 100},
    NO2: {"1h": 200, "anual": 60},
    CO: {"8h": 5000, "1h": 35000},
    O3: {"8h": 100},
}

# Factores de conversion (Protocolo de calidad del aire, Tabla 29), como en la FP-021.
# CO: ppm -> mg/m3 (x1000 -> µg/m3); los demas: ppb -> µg/m3.
FACTORES = {CO: 1.14 * 1000, NO2: 1.88, O3: 1.96, SO2: 2.62}

COMPUESTOS_COV = ("Tolueno", "Benceno", "Etilbenceno", "m,p-Xileno", "o-Xileno")

# Limites de cuantificacion del laboratorio (masa en µg). Una masa igual o
# inferior se reporta con "<".
LC_POR_DEFECTO = {
    SO2: 2.50,
    "Tolueno": 0.030, "Benceno": 0.030, "Etilbenceno": 0.030, "m,p-Xileno": 0.060, "o-Xileno": 0.030,
}

# Puntos de corte del ICA (Res. 2254 de 2017, Art. 20, Tabla 6), como en las plantillas.
CATEGORIAS_ICA = (
    (0, 50, "Buena", "Verde"),
    (51, 100, "Aceptable", "Amarillo"),
    (101, 150, "Dañina a la salud para grupos sensibles", "Naranja"),
    (151, 200, "Dañina a la salud", "Rojo"),
    (201, 300, "Muy dañina a la salud", "Púrpura"),
    (301, 500, "Peligrosa", "Marrón"),
)
PUNTOS_CORTE_ICA = {
    PM10: ((0, 54.99), (55, 154.99), (155, 254.99), (255, 354.99), (355, 424.99), (425, 604)),
    PM25: ((0, 12.99), (13, 37.99), (38, 55.99), (56, 150.99), (151, 250.99), (251, 500)),
    CO: ((0, 5094.99), (5095, 10819.99), (10820, 14254.99), (14255, 17688.99), (17689, 34862.99), (34863, 57703)),
    NO2: ((0, 100.99), (101, 189.99), (190, 677.99), (678, 1221.99), (1222, 2349.99), (2350, 3853)),
    O3: ((0, 106.99), (107, 138.99), (139, 167.99), (168, 207.99), (208, 393.99)),
}


class ErrorPlantillaAire(Exception):
    """La plantilla no tiene la estructura esperada."""


# ----------------------------------------------------------------- utilidades
def _num(v) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    t = str(v).strip().replace(",", ".").lstrip("<").strip()
    try:
        return float(t)
    except ValueError:
        return None


def _hora(v) -> Optional[time]:
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    if isinstance(v, timedelta):
        return (datetime.min + v).time()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 <= v < 1:
        return (datetime.min + timedelta(days=float(v))).time().replace(microsecond=0)
    if isinstance(v, str):
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.strptime(v.strip(), fmt).time()
            except ValueError:
                continue
    return None


def _dia(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _promedio(*vals) -> Optional[float]:
    v = [x for x in vals if x is not None]
    return sum(v) / len(v) if v else None


def _texto(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


class _Celda:
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class _Hoja:
    """Hoja ya leida en memoria (solo valores) con la parte de la interfaz de
    openpyxl que usan los lectores: title, sheet_state, max_row, cell e iter_rows."""

    def __init__(self, ws):
        self.title = ws.title
        self.sheet_state = getattr(ws, "sheet_state", "visible")
        self._ws = ws
        self._filas = None

    def _cargar(self) -> list:
        if self._filas is None:
            if hasattr(self._ws, "reset_dimensions"):
                self._ws.reset_dimensions()  # no confiar en la dimension declarada en el archivo
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")  # formatos condicionales no soportados
                self._filas = [tuple(f) for f in self._ws.iter_rows(values_only=True)]
            self._ws = None
        return self._filas

    @property
    def max_row(self) -> int:
        return len(self._cargar())

    def cell(self, fila: int, columna: int) -> _Celda:
        filas = self._cargar()
        valores = filas[fila - 1] if 0 < fila <= len(filas) else ()
        return _Celda(valores[columna - 1] if 0 < columna <= len(valores) else None)

    def __getitem__(self, coordenada: str) -> _Celda:
        from openpyxl.utils.cell import coordinate_to_tuple

        return self.cell(*coordinate_to_tuple(coordenada))

    def iter_rows(self, min_row: int = 1, max_row: Optional[int] = None, values_only: bool = True):
        yield from self._cargar()[min_row - 1:max_row]


class _Libro:
    def __init__(self, wb):
        self.worksheets = [_Hoja(ws) for ws in wb.worksheets]


def _abrir(ruta: str) -> _Libro:
    """Abre la plantilla en modo de solo lectura: sin estilos ni celdas combinadas,
    que en las FP costaban ~6 s de CPU por archivo (el plan de Render tiene 0,15 CPU).
    Cada hoja se lee una sola vez, al usarla, y queda en memoria solo con valores."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # formatos condicionales e imagenes WMF
        try:
            return _Libro(openpyxl.load_workbook(ruta, data_only=True, read_only=True))
        except Exception as exc:  # noqa: BLE001
            raise ErrorPlantillaAire(f"No se pudo abrir el archivo: {exc}") from exc


def _hojas_estacion(wb, patron: str) -> list:
    """Hojas visibles de estacion (CA-1, CA2, 'CA3,'...) ordenadas por numero."""
    hojas = []
    for ws in wb.worksheets:
        m = re.fullmatch(patron, ws.title.strip(), flags=re.IGNORECASE)
        if m and ws.sheet_state == "visible":
            hojas.append((int(m.group(1)), ws))
    return sorted(hojas, key=lambda h: h[0])


def _fila_encabezado(ws, columna: int, texto: str, desde: int = 1, hasta: int = 400) -> Optional[int]:
    for r in range(desde, min(hasta, ws.max_row) + 1):
        if _texto(ws.cell(r, columna).value).lower() == texto.lower():
            return r
    return None


def _regresion(xs: list, ys: list) -> tuple[float, float, float]:
    """Pendiente, intercepto y coeficiente de correlacion (SLOPE/INTERCEPT/CORREL de Excel)."""
    n = len(xs)
    mx, my = sum(xs) / n, sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    m = sxy / sxx
    r = sxy / math.sqrt(sxx * syy) if syy else 1.0
    return m, my - m * mx, r


def redondear(valor: Optional[float], decimales: int = 2) -> Optional[float]:
    """Redondeo como lo muestra Excel (15 cifras significativas y mitad hacia arriba):
    2258,6249999999995 -> 2258,63."""
    if valor is None:
        return None
    d = Decimal(f"{valor:.15g}").quantize(Decimal(1).scaleb(-decimales), rounding=ROUND_HALF_UP)
    return float(d)


def percentil_exc(valores: list, p: float) -> float:
    """PERCENTILE.EXC de Excel (interpolacion sobre el rango p*(n+1))."""
    v = sorted(valores)
    n = len(v)
    rango = p * (n + 1)
    if rango <= 1:
        return v[0]
    if rango >= n:
        return v[-1]
    k = int(rango)
    return v[k - 1] + (rango - k) * (v[k] - v[k - 1])


# ---------------------------------------------------------------- estadistica
@dataclass
class Estadistica:
    """Analisis estadistico de las plantillas (deteccion de datos atipicos)."""
    n: int
    promedio: float
    desviacion: float
    cv: float
    mediana: float
    q1: float
    q3: float
    iqr: float
    lim_inf: float  # m - 3*IQR
    lim_sup: float  # m + 3*IQR
    maximo: float
    fecha_max: Optional[date]
    minimo: float
    fecha_min: Optional[date]
    atipicos: int
    bajo_lc: bool = False  # todos los datos bajo el limite de cuantificacion ("<")


def estadistica(datos: list, poblacional: bool = False, bajo_lc: bool = False) -> Optional[Estadistica]:
    """`datos`: [(fecha, valor)]. Las plantillas manuales usan STDEVA (muestral)
    y la de equipos automaticos STDEV.P (poblacional)."""
    vals = [v for _, v in datos if v is not None]
    if len(vals) < 2:
        return None
    prom = sum(vals) / len(vals)
    desv = statistics.pstdev(vals) if poblacional else statistics.stdev(vals)
    med = statistics.median(vals)
    q1, q3 = percentil_exc(vals, 0.25), percentil_exc(vals, 0.75)
    iqr = q3 - q1
    maximo, minimo = max(vals), min(vals)
    f_max = next(f for f, v in datos if v == maximo)
    f_min = next(f for f, v in datos if v == minimo)
    lim_inf, lim_sup = med - 3 * iqr, med + 3 * iqr
    # Como las plantillas: atipica la concentracion por encima de m + 3*IQR.
    atipicos = sum(1 for v in vals if v > lim_sup)
    return Estadistica(
        n=len(vals), promedio=prom, desviacion=desv, cv=desv / prom if prom else 0.0, mediana=med,
        q1=q1, q3=q3, iqr=iqr, lim_inf=lim_inf, lim_sup=lim_sup, maximo=maximo,
        fecha_max=f_max.date() if isinstance(f_max, datetime) else f_max, minimo=minimo,
        fecha_min=f_min.date() if isinstance(f_min, datetime) else f_min, atipicos=atipicos, bajo_lc=bajo_lc,
    )


# ------------------------------------------------------------------------ ICA
def ica(contaminante: str, concentracion: Optional[float]) -> Optional[tuple[float, str]]:
    """Indice de calidad del aire (interpolacion lineal entre puntos de corte)."""
    cortes = PUNTOS_CORTE_ICA.get(contaminante)
    if not cortes or concentracion is None or concentracion < 0:
        return None
    for i, (pc_bajo, pc_alto) in enumerate(cortes):
        siguiente = cortes[i + 1][0] if i + 1 < len(cortes) else pc_alto
        if concentracion < siguiente or i == len(cortes) - 1:
            if concentracion > pc_alto and i == len(cortes) - 1:
                return None  # fuera de la tabla (p.ej. O3 8 h > 394 µg/m3)
            i_bajo, i_alto = CATEGORIAS_ICA[i][0], CATEGORIAS_ICA[i][1]
            valor = (i_alto - i_bajo) / (pc_alto - pc_bajo) * (min(concentracion, pc_alto) - pc_bajo) + i_bajo
            return valor, categoria_ica(valor)
    return None


def categoria_ica(valor: float) -> str:
    for _, alto, nombre, _ in CATEGORIAS_ICA:
        if valor <= alto:
            return nombre
    return CATEGORIAS_ICA[-1][2]


# --------------------------------------------------------- muestras de 24 h
@dataclass
class Muestra:
    """Una muestra integrada (24 h para PM10/PM2.5/SO2, 1 h para COV)."""
    inicio: Optional[datetime]
    fin: Optional[datetime]
    codigo: str  # filtro o numero de muestra de laboratorio
    minutos: Optional[float]
    temperatura: Optional[float]  # °C, promedio inicial/final
    presion: Optional[float]  # mmHg
    caudal: Optional[float]  # m3/min (PM10), L/min (PM2.5, COV), mL/min (SO2)
    masa: Optional[float]  # µg
    volumen: Optional[float]  # m3 a condiciones actuales
    concentracion_actual: Optional[float]  # µg/m3
    concentracion: Optional[float]  # µg/m3 a condiciones de referencia
    valida: bool = True
    bajo_lc: bool = False

    @property
    def fecha(self) -> Optional[date]:
        """Fecha final de monitoreo (la que se reporta)."""
        return (self.fin or self.inicio).date() if (self.fin or self.inicio) else None


@dataclass
class SerieManual:
    contaminante: str  # PM10, PM2.5, SO2 o un compuesto COV
    estacion: int
    nombre_estacion: str
    muestras: list = field(default_factory=list)
    advertencias: list = field(default_factory=list)

    @property
    def datos(self) -> list:
        return [(m.fecha, m.concentracion) for m in self.muestras if m.concentracion is not None]

    @property
    def bajo_lc(self) -> bool:
        return bool(self.muestras) and all(m.bajo_lc for m in self.muestras)

    @property
    def pct_validas(self) -> Optional[float]:
        return 100 * sum(m.valida for m in self.muestras) / len(self.muestras) if self.muestras else None

    def estadistica(self) -> Optional[Estadistica]:
        return estadistica(self.datos, bajo_lc=self.bajo_lc)


def _inicio_fin(fecha_ini, hora_ini, fecha_fin, hora_fin) -> tuple[Optional[datetime], Optional[datetime]]:
    d0, h0, h1 = _dia(fecha_ini), _hora(hora_ini), _hora(hora_fin)
    d1 = _dia(fecha_fin) or (d0 + timedelta(days=1) if d0 else None)
    inicio = datetime.combine(d0, h0 or time()) if d0 else None
    fin = datetime.combine(d1, h1 or time()) if d1 else None
    return inicio, fin


def _filas_datos(ws, fila_inicio: int):
    """Filas de datos consecutivas (columna B con el numero de muestra)."""
    r = fila_inicio
    while r <= ws.max_row and _num(ws.cell(r, 2).value) is not None:
        yield r
        r += 1


def _v(ws, celda):
    return ws[celda].value


def leer_pm10_hivol(ruta: str) -> list:
    """FP-031: PM10 con muestreador de alto volumen. Una hoja CA-n por estacion."""
    wb = _abrir(ruta)
    hojas = _hojas_estacion(wb, r"CA-?\s*(\d+),?")
    if not hojas:
        raise ErrorPlantillaAire("No se encontraron hojas de estacion (CA-1, CA-2...) en la plantilla FP-031 de PM10.")
    series = []
    for n, ws in hojas:
        fila = _fila_encabezado(ws, 2, "ID", 60)
        if fila is None or "fecha de inicio" not in _texto(ws.cell(fila, 3).value).lower():
            raise ErrorPlantillaAire(f"La hoja '{ws.title}' no tiene el formato de la FP-031 (PM10 alto volumen).")
        serie = SerieManual(PM10, n, _texto(_v(ws, "I10")))
        # Calibracion del equipo (orificio de transferencia): Qa y lectura corregida
        # de cada punto, y las rectas de la plantilla (E42/E43 y M42/M43).
        p_amb, t_amb = _num(_v(ws, "E17")), _num(_v(ws, "J18"))
        m_qa, b_qa = _num(_v(ws, "K24")), _num(_v(ws, "K25"))
        qa, i_corr, i_act = [], [], []
        for r in range(36, 41):
            dh, lectura = _num(ws.cell(r, 5).value), _num(ws.cell(r, 8).value)
            if None in (dh, lectura, p_amb, t_amb, m_qa, b_qa):
                continue
            tk = t_amb + 273.15
            qa.append((1 / m_qa) * (math.sqrt(dh * tk / p_amb) - b_qa))
            i_corr.append(lectura * math.sqrt(tk / p_amb))
            i_act.append(lectura)
        if len(qa) < 3:
            serie.advertencias.append(f"{ws.title}: la calibracion del Hi-Vol tiene menos de 3 puntos; no se calcula PM10.")
            series.append(serie)
            continue
        pend, inter, corr = _regresion(qa, i_corr)
        pend_m, inter_m, _ = _regresion(qa, i_act)
        if corr < 0.99:
            serie.advertencias.append(f"{ws.title}: el coeficiente de correlacion de la calibracion ({corr:.4f}) es menor que 0,99.")

        for r in _filas_datos(ws, fila + 2):
            c = lambda col: ws.cell(r, col).value  # noqa: E731
            inicio, fin = _inicio_fin(c(3), c(5), c(4), c(6))
            h0, h1 = _num(c(7)), _num(c(8))
            minutos = (h1 - h0) * 60 if h0 is not None and h1 is not None else None
            manometro = _promedio(_num(c(16)), _num(c(17)))
            temp = _promedio(_num(c(20)), _num(c(21)))
            pres = _promedio(_num(c(23)), _num(c(24)))
            w0, w1 = _num(c(27)), _num(c(28))
            masa = (w1 - w0) * 1e6 if w0 is not None and w1 is not None else None
            caudal = (manometro - inter_m) / pend_m if manometro is not None else None
            vol = vol_ref = None
            if None not in (caudal, temp, pres, minutos):
                vol = (1 / pend) * (caudal * math.sqrt((temp + 273) / pres) - inter) * minutos
                vol_ref = vol * (pres * 298.15) / (760 * (temp + 273.15))
            valida = (minutos is not None and 1379 <= minutos <= 1501 and caudal is not None
                      and 1.02 <= caudal <= 1.24 and corr >= 0.99)
            serie.muestras.append(Muestra(
                inicio=inicio, fin=fin, codigo=_texto(c(26)), minutos=minutos, temperatura=temp, presion=pres,
                caudal=caudal, masa=masa, volumen=vol,
                concentracion_actual=masa / vol if masa is not None and vol else None,
                concentracion=masa / vol_ref if masa is not None and vol_ref else None, valida=valida,
            ))
        series.append(serie)
    return series


def leer_pm25_lowvol(ruta: str) -> list:
    """FP-032: PM2.5 con muestreador de bajo volumen (16,67 L/min)."""
    wb = _abrir(ruta)
    hojas = _hojas_estacion(wb, r"CA-?\s*(\d+),?")
    if not hojas:
        raise ErrorPlantillaAire("No se encontraron hojas de estacion (CA-1, CA-2...) en la plantilla FP-032 de PM2.5.")
    series = []
    for n, ws in hojas:
        fila = _fila_encabezado(ws, 2, "ID", 30)
        if fila is None or _texto(ws.cell(fila, 3).value).lower() != "inicial":
            raise ErrorPlantillaAire(f"La hoja '{ws.title}' no tiene el formato de la FP-032 (PM2.5 bajo volumen).")
        serie = SerieManual(PM25, n, _texto(_v(ws, "G41")))
        for r in _filas_datos(ws, fila + 2):
            c = lambda col: ws.cell(r, col).value  # noqa: E731
            inicio, fin = _inicio_fin(c(3), c(4), c(5), c(6))
            minutos = (fin - inicio).total_seconds() / 60 if inicio and fin else None
            temp = _promedio(_num(c(9)), _num(c(10)))
            pres = _promedio(_num(c(12)), _num(c(13)))
            caudal = _promedio(_num(c(15)), _num(c(16)))
            w0, w1 = _num(c(23)), _num(c(24))
            masa = (w1 - w0) * 1e6 if w0 is not None and w1 is not None else None
            vol = vol_ref = None
            if None not in (caudal, minutos):
                horas = minutos / 60
                vol = caudal * horas * 0.06
                if None not in (temp, pres):
                    # Caudal a condiciones de referencia (25 °C, 760 mmHg).
                    q_ref = caudal * (pres / (temp + 273.15)) * (298.15 / 760)
                    vol_ref = q_ref * horas * 0.06
            valida = (minutos is not None and 1379 <= minutos <= 1501 and vol is not None and 23 <= vol <= 25
                      and caudal is not None and 15.84 <= caudal <= 17.5)
            serie.muestras.append(Muestra(
                inicio=inicio, fin=fin, codigo=_texto(c(22)), minutos=minutos, temperatura=temp, presion=pres,
                caudal=caudal, masa=masa, volumen=vol,
                concentracion_actual=masa / vol if masa is not None and vol else None,
                concentracion=masa / vol_ref if masa is not None and vol_ref else None, valida=valida,
            ))
        series.append(serie)
    return series


def leer_so2(ruta: str, lc: float = LC_POR_DEFECTO[SO2]) -> list:
    """FP-033: SO2 manual (tren de burbujeo, 180-220 mL/min durante 24 h)."""
    wb = _abrir(ruta)
    hojas = _hojas_estacion(wb, r"CA-?\s*(\d+),?")
    if not hojas:
        raise ErrorPlantillaAire("No se encontraron hojas de estacion (CA-1, CA-2...) en la plantilla FP-033 de SO2.")
    series = []
    for n, ws in hojas:
        fila = _fila_encabezado(ws, 2, "ID", 20)
        if fila is None or "fecha de inicio" not in _texto(ws.cell(fila, 3).value).lower():
            raise ErrorPlantillaAire(f"La hoja '{ws.title}' no tiene el formato de la FP-033 (SO2).")
        serie = SerieManual(SO2, n, _texto(_v(ws, "B10")))
        for r in _filas_datos(ws, fila + 2):
            c = lambda col: ws.cell(r, col).value  # noqa: E731
            inicio, fin = _inicio_fin(c(3), c(5), c(4), c(6))
            h0, h1 = _num(c(7)), _num(c(8))
            minutos = (h1 - h0) * 60 if h0 is not None and h1 is not None else None
            q0, q1 = _num(c(10)), _num(c(11))
            caudal = _promedio(q0, q1)
            temp = _promedio(_num(c(14)), _num(c(15)))
            pres = _promedio(_num(c(17)), _num(c(18)))
            masa = _num(c(21))
            actual = ref = None
            if None not in (masa, caudal, minutos) and caudal and minutos:
                actual = masa * 1e6 / (caudal * minutos)
                if None not in (temp, pres):
                    ref = actual / ((pres / (temp + 273.15)) * (298.15 / 760))
            dif = abs((q0 - q1) / q1) * 100 if q0 is not None and q1 else None
            valida = (minutos is not None and 1379 <= minutos <= 1501 and caudal is not None
                      and 180 <= caudal <= 220 and dif is not None and dif <= 5)
            serie.muestras.append(Muestra(
                inicio=inicio, fin=fin, codigo=_texto(c(20)), minutos=minutos, temperatura=temp, presion=pres,
                caudal=caudal, masa=masa, volumen=caudal * minutos / 1e6 if caudal and minutos else None,
                concentracion_actual=actual, concentracion=ref, valida=valida,
                bajo_lc=masa is not None and masa <= lc + 1e-9 or _texto(c(21)).startswith("<"),
            ))
        series.append(serie)
    return series


def _compuesto(titulo: str) -> Optional[str]:
    t = _texto(titulo).lower().replace(" ", "")
    if "m,p" in t or "m+p" in t or "mp-x" in t:
        return "m,p-Xileno"
    if "o-xileno" in t or "o-xilen" in t:
        return "o-Xileno"
    for nombre in ("Etilbenceno", "Benceno", "Tolueno"):
        if nombre.lower() in t:
            return nombre
    return None


def leer_cov(ruta: str, lc: Optional[dict] = None) -> list:
    """FP-035: COV en tubos de carbon activado (1 h). Un bloque por compuesto;
    se reporta el tubo de alto flujo si tiene dato, si no el de bajo flujo."""
    lc = {**LC_POR_DEFECTO, **(lc or {})}
    wb = _abrir(ruta)
    hojas = _hojas_estacion(wb, r"CA-?\s*(\d+),?")
    if not hojas:
        raise ErrorPlantillaAire("No se encontraron hojas de estacion (CA1, CA2...) en la plantilla FP-035 de COV.")
    series = []
    for n, ws in hojas:
        nombre = _texto(_v(ws, "H27"))
        primera = None  # filas del primer bloque: los demas pueden referenciar sus caudales
        for r in range(1, ws.max_row + 1):
            titulo = _texto(ws.cell(r, 2).value)
            if not titulo.upper().startswith(("VOC", "COV")):
                continue
            compuesto = _compuesto(titulo)
            if compuesto is None or _texto(ws.cell(r + 1, 2).value).upper() != "ID":
                continue
            serie = SerieManual(compuesto, n, nombre)
            filas = list(_filas_datos(ws, r + 2))
            primera = primera or filas
            for k, fr in enumerate(filas):
                def c(col, fr=fr, k=k):
                    v = ws.cell(fr, col).value
                    if v is None and primera and k < len(primera):
                        v = ws.cell(primera[k], col).value
                    return v
                inicio, fin = _inicio_fin(c(3), c(5), c(4), c(6))
                minutos = _num(c(12))
                temp = _promedio(_num(c(16)), _num(c(17)))
                pres = _promedio(_num(c(13)), _num(c(14)))
                masa_alto, masa_bajo = _num(ws.cell(fr, 9).value), _num(ws.cell(fr, 8).value)
                q_alto = _promedio(_num(c(31)), _num(c(32)))
                q_bajo = _promedio(_num(c(26)), _num(c(27)))
                if masa_alto is not None and q_alto:
                    masa, caudal = masa_alto, q_alto / 1000
                else:
                    masa, caudal = masa_bajo, (q_bajo / 1000 if q_bajo else None)
                actual = ref = None
                if None not in (masa, caudal, minutos) and caudal and minutos:
                    actual = masa * 1000 / (caudal * minutos)
                    if None not in (temp, pres):
                        ref = masa * 1000 / (caudal * minutos * (pres / (temp + 273)) * (298 / 760))
                serie.muestras.append(Muestra(
                    inicio=inicio, fin=fin, codigo=_texto(c(7)), minutos=minutos, temperatura=temp, presion=pres,
                    caudal=caudal, masa=masa, volumen=caudal * minutos / 1000 if caudal and minutos else None,
                    concentracion_actual=actual, concentracion=ref,
                    bajo_lc=masa is not None and masa <= lc.get(compuesto, 0) + 1e-9,
                ))
            series.append(serie)
    if not series:
        raise ErrorPlantillaAire("No se encontraron bloques de compuestos (VOC´S - TOLUENO...) en la plantilla FP-035.")
    return series


# --------------------------------------------------------- equipos automaticos
@dataclass
class DiaGas:
    fecha: date
    max_horario: Optional[float]  # µg/m3
    max_8h: Optional[float]  # maxima media movil de 8 h del dia (µg/m3)
    horas: int = 0


@dataclass
class SerieAutomatica:
    contaminante: str  # CO, NO2, O3
    estacion: int
    nombre_estacion: str
    horas: list = field(default_factory=list)  # [(datetime, µg/m3)]
    filas: int = 0  # registros horarios de la plantilla (con o sin dato)
    medias_8h: list = field(default_factory=list)  # [(datetime de la ultima hora, µg/m3)]
    advertencias: list = field(default_factory=list)

    @property
    def dias(self) -> list:
        por_dia: dict = {}
        for t, v in self.horas:
            por_dia.setdefault(t.date(), [[], []])[0].append(v)
        for t, v in self.medias_8h:
            por_dia.setdefault(t.date(), [[], []])[1].append(v)
        return [DiaGas(d, max(h) if h else None, max(m) if m else None, len(h))
                for d, (h, m) in sorted(por_dia.items())]

    @property
    def pct_validos(self) -> Optional[float]:
        return 100 * len(self.horas) / self.filas if self.filas else None

    def estadistica_1h(self) -> Optional[Estadistica]:
        return estadistica(self.horas, poblacional=True)

    def estadistica_8h(self) -> Optional[Estadistica]:
        return estadistica(self.medias_8h, poblacional=True)


def serie_automatica(contaminante: str, estacion: int, nombre: str, lecturas: list,
                     filas: Optional[int] = None) -> SerieAutomatica:
    """`lecturas`: [(datetime, valor en ppm (CO) o ppb)]. La media movil de 8 h
    toma las 8 lecturas consecutivas que terminan en cada hora (como la FP-021)."""
    serie = SerieAutomatica(contaminante, estacion, nombre)
    factor = FACTORES[contaminante]
    serie.horas = [(t, v * factor) for t, v in sorted(lecturas)]
    vals = [v for _, v in serie.horas]
    serie.medias_8h = [(serie.horas[i][0], sum(vals[i - 7:i + 1]) / 8) for i in range(7, len(vals))]
    serie.filas = max(filas or 0, len(serie.horas))
    return serie


_HOJA_FP021 = r"ESTACI[OÓ]N\s*(\d+)\s*_\s*(CO|NO2|O3)"


def leer_fp021(ruta: str) -> list:
    """FP-021: una hoja visible por estacion y gas (ESTACION 1_CO, ESTACIÓN 2_NO2...)."""
    wb = _abrir(ruta)
    series = []
    for ws in wb.worksheets:
        m = re.fullmatch(_HOJA_FP021, ws.title.strip(), flags=re.IGNORECASE)
        if not m or ws.sheet_state != "visible":
            continue
        n, gas = int(m.group(1)), m.group(2).upper()
        enc = None
        for r in range(1, 20):
            fila = [_texto(ws.cell(r, c).value).lower() for c in range(1, 12)]
            if any(t.startswith("fecha") for t in fila) and any(t.startswith("hora") for t in fila):
                enc = r, fila
                break
        if enc is None:
            raise ErrorPlantillaAire(f"La hoja '{ws.title}' no tiene las columnas Fecha / Hora / Concentración.")
        r0, fila = enc
        col_f = next(i for i, t in enumerate(fila) if t.startswith("fecha")) + 1
        col_h = next(i for i, t in enumerate(fila) if t.startswith("hora")) + 1
        col_c = next((i for i, t in enumerate(fila) if t.startswith("concentraci") and ("ppm" in t or "ppb" in t)), None)
        if col_c is None:
            raise ErrorPlantillaAire(f"La hoja '{ws.title}' no tiene la columna de concentración en ppm/ppb.")
        nombre = ""
        for r in range(1, r0):
            if "nombre de la estaci" in _texto(ws.cell(r, 2).value).lower():
                nombre = next((_texto(ws.cell(r, c).value) for c in range(3, 12) if _texto(ws.cell(r, c).value)), "")
        lecturas, dia, filas = [], None, 0
        for r in range(r0 + 1, ws.max_row + 1):
            d = _dia(ws.cell(r, col_f).value)
            dia = d or dia
            h = _hora(ws.cell(r, col_h).value)
            v = _num(ws.cell(r, col_c + 1).value)
            if h is None:
                if lecturas:
                    break
                continue
            filas += 1
            if dia is not None and v is not None:
                lecturas.append((datetime.combine(dia, h), v))
        series.append(serie_automatica(gas, n, nombre, lecturas, filas))
    if not series:
        raise ErrorPlantillaAire("No se encontraron hojas de estación (ESTACION 1_CO, ESTACIÓN 1_NO2, ESTACION 1_O3...) "
                                 "en la plantilla FP-021.")
    return sorted(series, key=lambda s: (AUTOMATICOS.index(s.contaminante), s.estacion))


def leer_reporte_analizador(ruta: str, contaminante: str, estacion: int, nombre: str = "") -> SerieAutomatica:
    """Exportacion directa del analizador: columna 'Time' y la del gas (ppm/ppb)."""
    wb = _abrir(ruta)
    ws = wb.worksheets[0]
    enc = [_texto(c).lower() for c in next(ws.iter_rows(max_row=1, values_only=True))]
    clave = {CO: "co", NO2: "no2", O3: "o3"}[contaminante]
    col = next((i for i, t in enumerate(enc) if re.match(rf"^{clave}\b", t)), None)
    if not enc or not enc[0].startswith(("time", "fecha")) or col is None:
        raise ErrorPlantillaAire(f"El reporte del analizador no tiene las columnas 'Time' y '{contaminante}'.")
    lecturas = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        t, v = fila[0], _num(fila[col]) if col < len(fila) else None
        if isinstance(t, datetime) and v is not None:
            lecturas.append((t, v))
    return serie_automatica(contaminante, estacion, nombre, lecturas)


# ---------------------------------------------------------------- proyecto
@dataclass
class EstacionAire:
    numero: int  # CA-n / ESTACION n en las plantillas
    nombre: str = ""
    codigo: str = ""  # ID interno (E1_Lis_Tmpst)
    codigo_anla: str = ""
    longitud: str = ""  # grados o Este (Origen Nacional)
    latitud: str = ""  # grados o Norte
    descripcion: str = ""
    foto_ruta: str = ""  # foto de la estacion (informe Word)


@dataclass
class ProyectoAire:
    nombre_proyecto: str = ""
    codigo: str = ""
    cliente: str = ""
    estaciones: list = field(default_factory=list)  # [EstacionAire]
    # contaminante (PM10, PM2.5, SO2, COV, AUTOMATICOS) -> ruta de la plantilla FP
    plantillas: dict = field(default_factory=dict)
    # (contaminante, numero de estacion) -> ruta del reporte del analizador
    reportes_analizador: dict = field(default_factory=dict)
    limites_cuantificacion: dict = field(default_factory=dict)


@dataclass
class ResultadosAire:
    proyecto: ProyectoAire
    manuales: dict = field(default_factory=dict)  # (contaminante, estacion) -> SerieManual
    automaticos: dict = field(default_factory=dict)  # (contaminante, estacion) -> SerieAutomatica
    cov: dict = field(default_factory=dict)  # (compuesto, estacion) -> SerieManual
    advertencias: list = field(default_factory=list)

    def estaciones(self) -> list:
        """Estaciones con datos, completando nombres desde las plantillas."""
        por_numero = {e.numero: e for e in self.proyecto.estaciones}
        numeros = {n for (_, n) in [*self.manuales, *self.automaticos, *self.cov]} | set(por_numero)
        salida = []
        for n in sorted(numeros):
            e = por_numero.get(n) or EstacionAire(numero=n)
            if not e.nombre:
                serie = next((s for (c, k), s in [*self.manuales.items(), *self.automaticos.items(), *self.cov.items()]
                              if k == n and s.nombre_estacion), None)
                e = EstacionAire(**{**e.__dict__, "nombre": serie.nombre_estacion if serie else f"Estación {n}"})
            salida.append(e)
        return salida

    def contaminantes(self) -> list:
        presentes = {c for c, _ in self.manuales} | {c for c, _ in self.automaticos}
        if self.cov:
            presentes.add(COV)
        return [c for c in CONTAMINANTES if c in presentes]

    def dias_ica(self, estacion: int) -> dict:
        """ICA diario por contaminante: {contaminante: [(fecha, concentracion, ica, categoria)]}."""
        salida = {}
        for c in (PM10, PM25):
            s = self.manuales.get((c, estacion))
            if s:
                salida[c] = [(f, v, *(ica(c, v) or (None, None))) for f, v in s.datos]
        for c, attr in ((CO, "max_8h"), (NO2, "max_horario"), (O3, "max_8h")):
            s = self.automaticos.get((c, estacion))
            if s:
                salida[c] = [(d.fecha, getattr(d, attr), *(ica(c, getattr(d, attr)) or (None, None)))
                             for d in s.dias if getattr(d, attr) is not None]
        return salida


def procesar_aire(proyecto: ProyectoAire) -> ResultadosAire:
    res = ResultadosAire(proyecto)
    lc = {**LC_POR_DEFECTO, **proyecto.limites_cuantificacion}
    lectores = {PM10: leer_pm10_hivol, PM25: leer_pm25_lowvol, SO2: lambda r: leer_so2(r, lc[SO2])}
    for contaminante, lector in lectores.items():
        ruta = proyecto.plantillas.get(contaminante)
        if not ruta:
            continue
        try:
            for s in lector(ruta):
                res.manuales[(contaminante, s.estacion)] = s
                res.advertencias.extend(s.advertencias)
        except ErrorPlantillaAire as exc:
            res.advertencias.append(f"{contaminante}: {exc}")
    if proyecto.plantillas.get(COV):
        try:
            for s in leer_cov(proyecto.plantillas[COV], lc):
                res.cov[(s.contaminante, s.estacion)] = s
        except ErrorPlantillaAire as exc:
            res.advertencias.append(f"COV: {exc}")
    if proyecto.plantillas.get("AUTOMATICOS"):
        try:
            for s in leer_fp021(proyecto.plantillas["AUTOMATICOS"]):
                res.automaticos[(s.contaminante, s.estacion)] = s
        except ErrorPlantillaAire as exc:
            res.advertencias.append(f"Equipos automáticos: {exc}")
    # Un reporte directo del analizador reemplaza la hoja de la FP-021 de esa estacion.
    for (contaminante, n), ruta in proyecto.reportes_analizador.items():
        try:
            res.automaticos[(contaminante, n)] = leer_reporte_analizador(ruta, contaminante, n)
        except ErrorPlantillaAire as exc:
            res.advertencias.append(f"{contaminante} estación {n}: {exc}")

    for (c, n), s in sorted(res.manuales.items()):
        invalidas = [m.fecha for m in s.muestras if not m.valida]
        if invalidas:
            res.advertencias.append(
                f"{c} – {s.nombre_estacion or f'estación {n}'}: {len(invalidas)} muestra(s) fuera de los criterios de "
                "validez (tiempo de muestreo, caudal o volumen): " + ", ".join(f.isoformat() for f in invalidas if f) + ".")
    for (c, n), s in sorted(res.automaticos.items()):
        # El primer y el ultimo dia suelen ser parciales (instalacion y retiro del equipo).
        incompletos = [d.fecha.isoformat() for d in s.dias[1:-1] if d.horas < 18]
        if incompletos:
            s.advertencias.append(f"{c} – {s.nombre_estacion or f'estación {n}'}: día(s) con menos de 18 datos "
                                  "horarios (75 %): " + ", ".join(incompletos) + ".")
            res.advertencias.extend(s.advertencias)
    return res
