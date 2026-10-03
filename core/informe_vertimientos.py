"""Informe Word de vertimientos (Res. 0631 de 2015).

Plantilla: templates/informe_vertimientos_template.docx (formato FP-023, a partir
del informe EA-144-26). Se conservan las partes fijas (portada, cuadro de
control, laboratorio, metodologia de muestreo, regla de decision,
documentacion y anexos) y se redactan con los datos del proyecto la
introduccion, los objetivos, el cliente, la descripcion de los puntos, la
tabla de metodos, la normatividad, los resultados (in situ y de laboratorio),
el analisis por parametro con sus graficas y las conclusiones.
"""
from __future__ import annotations

import copy
import datetime as dt
import os
import re
import textwrap
from types import SimpleNamespace
from typing import Optional

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from . import res0631, res0631_datos
from . import vertimientos_graficas as gr
from . import vertimientos_textos as tx
from .docx_utils import _sin_tildes, blips, encontrar_tabla, reemplazar_imagen, set_cell_text, set_paragraph_text
from .informe_aire import (
    EMPRESA, MESES, Cursor, _cuerpo, _es_leyenda, _estilo, _quitar, _quitar_entradas_huerfanas_de_indices,
    _renumerar, _texto, fecha_larga, fmt, parsear_coordenada, quitar_relaciones_huerfanas,
)
from .informe_textos import _anio_fuentes, _cliente, _cuadro_control, actualizar_indices_al_abrir, cantidad, lista
from .vertimientos import FilaComparacion, ResultadoPunto, ResultadosVertimiento, numero

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ACREDITACION = "Resolución 1004 del 24 de agosto de 2026"
ACREDITACION_ANTERIOR = re.compile(r"Resolución 1096 del 11 de octubre de 2024")
LAB_PROPIO = f"{EMPRESA} acreditado bajo la {ACREDITACION}"
LAB_SUBCONTRATADO = "Laboratorio subcontratado acreditado por el IDEAM"
NO_CUMPLE = "FFC7CE"
SUBENCABEZADO = "D9E7A1"
ANCHO = Cm(15.5)
TIPOS_AGUA = {"ARnD": "Agua Residual no Doméstica", "ARD": "Agua Residual Doméstica"}


# ------------------------------------------------------------------ formato
def num(v: Optional[float], dec: int = 2) -> str:
    return "---" if v is None else fmt(v, dec, miles=abs(v) >= 10000)


def hora(t) -> str:
    return t.strftime("%H:%M") if t else ""


def _decimales(v: float) -> int:
    return 0 if abs(v) >= 100 else 1 if abs(v) >= 10 else 2


def _es_titulo_n(el, nivel: int, texto: Optional[str] = None) -> bool:
    if el.tag != f"{_W}p" or _estilo(el) != f"Ttulo{nivel}":
        return False
    return texto is None or texto in _sin_tildes(_texto(el))


def _nivel(el) -> Optional[int]:
    m = re.fullmatch(r"Ttulo(\d)", _estilo(el)) if el.tag == f"{_W}p" else None
    return int(m.group(1)) if m else None


def _titulo(doc, nivel: int, texto: str):
    return next((el for el in _cuerpo(doc) if _es_titulo_n(el, nivel, texto)), None)


def _vaciar(titulo) -> None:
    """Quita el contenido de la seccion (hasta el siguiente titulo de igual o mayor nivel)."""
    nivel = _nivel(titulo)
    el = titulo.getnext()
    while el is not None:
        siguiente = el.getnext()
        n = _nivel(el)
        if n is not None and n <= nivel:
            break
        if el.tag == f"{_W}sectPr":
            break
        if el.tag in (f"{_W}p", f"{_W}tbl"):
            _quitar(el)
        el = siguiente


def _numpr(doc, desde_titulo) -> Optional[object]:
    """Numeracion (vineta) del primer parrafo de lista de la seccion, para reutilizarla."""
    el = desde_titulo.getnext() if desde_titulo is not None else None
    while el is not None and _nivel(el) is None:
        if el.tag == f"{_W}p" and el.find(f"{_W}pPr/{_W}numPr") is not None:
            return copy.deepcopy(el.find(f"{_W}pPr/{_W}numPr"))
        el = el.getnext()
    return None


def _vineta(cur: Cursor, texto: str, numpr) -> Paragraph:
    p = cur.parrafo(contraer(texto), "List Paragraph")
    if numpr is not None:
        p._p.get_or_add_pPr().append(copy.deepcopy(numpr))
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def minus(nombre: str) -> str:
    """'Demanda Bioquímica de Oxígeno (DBO5)' -> 'demanda bioquímica de oxígeno (DBO5)' (respeta siglas y pH)."""
    palabras = []
    for w in nombre.split(" "):
        letras = re.sub(r"[^A-Za-zÁÉÍÓÚÑáéíóúñ]", "", w)
        sigla = sum(c.isupper() for c in letras) >= 2 or any(c.isdigit() for c in w) \
            or (letras[1:2].isupper() if len(letras) > 1 else False)
        palabras.append(w if sigla else w.lower())
    return " ".join(palabras)


def contraer(texto: str) -> str:
    """'de el' -> 'del', 'a el' -> 'al'."""
    def cambio(m):
        nuevo = "del" if m.group(1).lower() == "de" else "al"
        return nuevo.capitalize() if m.group(1)[0].isupper() else nuevo

    return re.sub(r"\b([Dd]e|[Aa]) el\b", cambio, texto)


ARTICULO = {
    "acidez": "la acidez", "alcalinidad": "la alcalinidad", "cloruros": "los cloruros", "fenoles": "los fenoles",
    "sulfatos": "los sulfatos", "sulfuros": "los sulfuros", "sst": "los sólidos suspendidos totales",
    "sdt": "los sólidos disueltos totales", "saam": "los surfactantes (SAAM)", "cianuro": "el cianuro total",
    "formaldehido": "el formaldehído", "dbo5": "la DBO5", "dqo": "la DQO", "aceites": "los aceites y grasas",
    "color": "el color verdadero", "dureza": "la dureza", "fosforo": "el fósforo total y los ortofosfatos",
    "nitrogenados": "los compuestos nitrogenados", "conductividad": "la conductividad eléctrica",
    "ph": "el pH", "oxigeno": "el oxígeno disuelto", "sedimentables": "los sólidos sedimentables",
    "temperatura": "la temperatura",
}


def _ref(texto: str, n: int) -> str:
    """Agrega "(ver Gráfica n)" antes del punto final."""
    return texto.rstrip().rstrip(".") + f" (ver Gráfica {n})."


def _p(cur: Cursor, texto: str) -> Paragraph:
    p = cur.parrafo(contraer(texto.replace("S.A.S..", "S.A.S.")))
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def _combinar(t: Table, r0: int, c0: int, r1: int, c1: int, texto: str, relleno: Optional[str] = None,
              negrita: bool = False, izquierda: bool = False):
    a = t.cell(r0, c0)
    if (r0, c0) != (r1, c1):
        a = a.merge(t.cell(r1, c1))
    set_cell_text(a, texto)
    for p in a.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if izquierda else WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(9)
            run.bold = negrita or None
    if relleno:
        tcpr = a._tc.get_or_add_tcPr()
        for viejo in tcpr.findall(qn("w:shd")):
            tcpr.remove(viejo)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), relleno)
        tcpr.append(shd)
    return a


def _fuente_y_notas(cur: Cursor, notas: list[str], anio: int):
    for n in notas:
        cur.parrafo(n, "Nota Tabla")
    cur.fuente(anio=anio)


# ---------------------------------------------------------------- contexto
class Contexto:
    def __init__(self, res: ResultadosVertimiento, datos, hoy: Optional[dt.date] = None):
        self.res = res
        self.proyecto = res.proyecto
        self.datos = datos
        self.puntos = res.puntos
        self.con_lab = [rp for rp in res.puntos if rp.informe]
        self.evaluados = [rp for rp in self.con_lab if rp.punto.evaluar]
        self.con_campo = [rp for rp in res.puntos if rp.campo and rp.campo.mediciones]
        lab = self.con_lab[0].informe if self.con_lab else None
        cliente_lab = lab.cliente if lab else {}
        self.cliente = (res.proyecto.cliente or cliente_lab.get("razon_social", "")).strip()
        fechas = [rp.informe.fecha_muestreo for rp in self.con_lab if rp.informe.fecha_muestreo]
        self.fecha_muestreo = min(fechas).date() if fechas else None
        fecha = None
        if getattr(datos, "fecha", ""):
            try:
                fecha = dt.date.fromisoformat(datos.fecha)
            except ValueError:
                fecha = None
        self.fecha = fecha or hoy or dt.date.today()
        self.anio = self.fecha.year
        self.municipio = (getattr(datos, "municipio", "") or (lab.municipio if lab else "")).strip()
        self.departamento = (getattr(datos, "departamento", "") or (lab.departamento if lab else "")).strip()
        tipos = sorted({rp.punto.tipo_agua for rp in res.puntos} or {"ARnD"})
        self.siglas = " y ".join(tipos)
        self.tipo_largo = " y ".join(TIPOS_AGUA.get(t, t) for t in tipos)
        self.compuesto = any("compuest" in _sin_tildes(rp.informe.tipo_muestreo) for rp in self.con_lab) \
            or bool(self.con_campo)
        # Datos del cliente: los del formulario y, si faltan, los del reporte del laboratorio.
        self.datos_cliente = SimpleNamespace(**{k: getattr(datos, k, "") for k in (
            "cliente_nit", "cliente_direccion", "cliente_contacto", "cliente_ciudad", "cliente_departamento",
            "cliente_actividad")})
        for campo, clave in (("cliente_nit", "nit"), ("cliente_direccion", "direccion"),
                             ("cliente_contacto", "contacto")):
            if not getattr(self.datos_cliente, campo).strip():
                setattr(self.datos_cliente, campo, cliente_lab.get(clave, ""))
        if not self.datos_cliente.cliente_ciudad.strip():
            self.datos_cliente.cliente_ciudad = self.municipio
        if not self.datos_cliente.cliente_departamento.strip():
            self.datos_cliente.cliente_departamento = self.departamento
        self.columnas = res.columnas
        self.aplican = [(i, c) for i, c in enumerate(res.columnas) if c.aplica]

    @property
    def ubicacion(self) -> str:
        return ", ".join(x for x in (self.municipio, self.departamento) if x)

    def nombres(self, puntos) -> str:
        return lista([f"«{rp.punto.nombre}»" for rp in puntos])

    def los_puntos(self, puntos, denominados: bool = True) -> str:
        if len(puntos) == 1:
            return f"el punto de monitoreo{' denominado' if denominados else ''} {self.nombres(puntos)}"
        return f"los puntos de monitoreo{' denominados' if denominados else ''} {self.nombres(puntos)}"

    def articulos(self) -> str:
        acts = [res0631.ACTIVIDADES[c.actividad] for c in self.columnas if not c.alcantarillado]
        vistos, salida = set(), []
        for a in acts:
            if a["clave"] not in vistos:
                vistos.add(a["clave"])
                salida.append(f"el Artículo {a['articulo']} ({a['actividad']})")
        return lista(salida)


# ------------------------------------------------------------ partes fijas
def _portada_y_encabezado(doc, ctx: Contexto, foto: Optional[str], foto_punto: str):
    titulo = f"INFORME TÉCNICO DE CALIDAD DE {ctx.tipo_largo.upper()} ({ctx.siglas})"
    titulo_corto = f"INFORME TÉCNICO DE CALIDAD DE {ctx.tipo_largo.upper()}"
    ref = ctx.fecha_muestreo or ctx.fecha
    mes, anio = MESES[ref.month - 1].upper(), str(ref.year)
    cliente = (ctx.cliente or ctx.proyecto.nombre_proyecto).upper()
    for txbx in doc.element.body.iter(f"{_W}txbxContent"):
        ps = [Paragraph(p, None) for p in txbx.iter(f"{_W}p")]
        con_texto = [p for p in ps if p.text.strip()]
        texto = _sin_tildes(" ".join(p.text for p in con_texto)).strip()
        if not texto or texto.startswith("fp-0"):
            continue
        if texto.startswith("informe tecnico"):
            set_paragraph_text(con_texto[0], titulo)
        elif texto.startswith("fotografia"):
            set_paragraph_text(con_texto[0], f"Fotografía: {foto_punto}" if foto_punto else "Fotografía:")
            for p in con_texto[1:]:
                set_paragraph_text(p, "")
        elif re.fullmatch(r"([a-z]+|mes) de\s*(\d{4}|ano)", texto.replace("\n", " ")):
            if len(con_texto) >= 2:
                set_paragraph_text(con_texto[0], f"{mes} DE")
                set_paragraph_text(con_texto[1], anio)
            else:
                set_paragraph_text(con_texto[0], f"{mes} DE {anio}")
        else:
            set_paragraph_text(con_texto[0], cliente)
            for p in con_texto[1:]:
                set_paragraph_text(p, "")
    # Foto de la portada (marcador de la plantilla junto a "Fotografía:").
    if foto and os.path.exists(foto):
        for el in list(doc.element.body.iterchildren())[:20]:
            if "fotografia" in _sin_tildes(_texto(el)) and blips(el):
                reemplazar_imagen(doc, blips(el)[0], foto)
                break
    version = (getattr(ctx.datos, "version", "") or "1.0").strip()
    cliente_previo = ""
    codigo = ctx.proyecto.codigo.strip()
    vistos = set()
    for seccion in doc.sections:
        for hdr in (seccion.header, seccion.first_page_header, seccion.even_page_header):
            if hdr.is_linked_to_previous or id(hdr._element) in vistos:
                continue
            vistos.add(id(hdr._element))
            for tabla in hdr.tables:
                for fila in tabla.rows:
                    for celda in fila.cells:
                        ps = [p for p in celda.paragraphs if p.text.strip()]
                        t0 = _sin_tildes(ps[0].text.strip()) if ps else ""
                        if t0.startswith("informe tecnico de calidad de agua"):
                            set_paragraph_text(ps[0], titulo_corto)
                        elif t0.startswith("codigo") and len(ps) >= 3:
                            set_paragraph_text(ps[1], codigo or "---")
                            set_paragraph_text(ps[2], f"VERSIÓN {version}")
                        elif re.fullmatch(r"([a-z]+|mes)", t0) and len(ps) == 2 and re.search(r"\d{4}|ano", _sin_tildes(ps[1].text)):
                            set_paragraph_text(ps[0], mes)
                            set_paragraph_text(ps[1], f"DE {anio}")
                        elif t0 == "nombre del cliente" or (cliente_previo and t0 == cliente_previo):
                            set_paragraph_text(ps[0], cliente)


def _acreditacion(doc):
    for el in _cuerpo(doc):
        if el.tag != f"{_W}p" or not ACREDITACION_ANTERIOR.search(_texto(el)):
            continue
        p = Paragraph(el, doc)
        set_paragraph_text(p, ACREDITACION_ANTERIOR.sub(ACREDITACION, p.text))


# ------------------------------------------------------------ coordenadas
_LAT, _LON = (-5.0, 14.0), (-83.0, -66.0)


def coordenadas(punto) -> dict:
    """Latitud/longitud (grados decimales) y Norte/Este (Origen Nacional) del punto."""
    from .geo_colombia import a_geografica, es_origen_nacional, geografica_a_origen_nacional

    salida = {"lat": "---", "lon": "---", "norte": "---", "este": "---", "dd": None}
    x, y = parsear_coordenada(punto.longitud), parsear_coordenada(punto.latitud)
    if x is None or y is None:
        return salida
    plana = es_origen_nacional(x, y)
    lat, lon = a_geografica(x, y)
    if not (_LAT[0] <= lat <= _LAT[1] and _LON[0] <= lon <= _LON[1]):
        return salida
    if plana:
        este, norte = x, y
    else:
        este, norte = geografica_a_origen_nacional(lat, lon)
    salida.update(lat=f"{lat:.5f}", lon=f"{lon:.5f}", norte=fmt(norte, 3, True), este=fmt(este, 3, True),
                  dd=(lat, lon))
    return salida


# ------------------------------------------------------------ introduccion
def _parametros_por_laboratorio(ctx: Contexto) -> list[tuple[str, str]]:
    """(laboratorio, parametro) en el orden de los reportes, con los HAP y los maximos/minimos agrupados."""
    filas, vistos = [], set()
    agrupar = {"ph_max": "pH", "ph_min": "pH", "conductividad_max": "Conductividad",
               "conductividad_min": "Conductividad", "solidos_sedimentables_max": "Sólidos Sedimentables",
               "solidos_sedimentables_min": "Sólidos Sedimentables", "temperatura_max": "Temperatura",
               "temperatura_min": "Temperatura"}
    for f in ctx.res.filas:
        if f.parametro.grupo == "hap":
            nombre = "Hidrocarburos Aromáticos Policíclicos (HAP)"
        else:
            nombre = agrupar.get(f.parametro.clave, f.parametro.nombre)
        if nombre in vistos:
            continue
        vistos.add(nombre)
        filas.append((LAB_SUBCONTRATADO if f.subcontratado else LAB_PROPIO, nombre))
    return sorted(filas, key=lambda x: (x[0] != LAB_PROPIO,))


def _introduccion(doc, ctx: Contexto, faltantes):
    t = _titulo(doc, 1, "introduccion")
    if t is None:
        faltantes.append("Capítulo de introducción de la plantilla")
        return
    _vaciar(t)
    cur = Cursor(doc, t)
    d = ctx.datos_cliente
    domicilio = f", con domicilio en {d.cliente_direccion.strip()}" if d.cliente_direccion.strip() else ""
    ubic = f", ubicada en el municipio de {ctx.ubicacion}" if ctx.ubicacion else ""
    quien = f"La empresa {ctx.cliente}" if ctx.cliente else "El cliente"
    n_tabla = cur.numero("Tabla")
    _p(cur, f"{quien} contrató a la firma {EMPRESA} para la toma y análisis de muestras de {ctx.tipo_largo} "
            f"({ctx.siglas}){domicilio}{ubic}. La presente caracterización fisicoquímica hace referencia a las "
            f"muestras de {ctx.tipo_largo} ({ctx.siglas}) tomadas en "
            + ("el punto de muestreo indicado" if len(ctx.puntos) == 1 else "los puntos de muestreo indicados")
            + f" en la Tabla {n_tabla}.")
    cur.leyenda("Tabla", "Puntos de monitoreo" if len(ctx.puntos) > 1 else "Punto de monitoreo")
    filas = []
    for rp in ctx.puntos:
        c = coordenadas(rp.punto)
        filas.append([rp.punto.nombre, c["lat"], c["lon"], c["norte"], c["este"]])
    cur.tabla([[("Punto de monitoreo", 1, 2), ("Coordenadas Geográficas", 2, 1),
                ("Coordenadas Origen Nacional Único", 2, 1)],
               ["Latitud (N)", "Longitud (W)", "Norte (m)", "Este (m)"]], filas, izquierda=(0,))
    cur.fuente(anio=ctx.anio)
    if any(coordenadas(rp.punto)["dd"] is None for rp in ctx.puntos):
        faltantes.append("Coordenadas: registre la latitud/longitud o el Este/Norte de todos los puntos")
    plan = f" El muestreo se realizó siguiendo las directrices plasmadas en el plan de monitoreo con código " \
           f"{ctx.proyecto.codigo.strip()}." if ctx.proyecto.codigo.strip() else ""
    _p(cur, "La planeación y ejecución del muestreo, la cadena de custodia y los análisis fisicoquímicos de las "
            f"muestras tomadas en campo fueron desarrolladas y ejecutadas por {EMPRESA}, laboratorio debidamente "
            f"acreditado ante el IDEAM mediante la {ACREDITACION}, por la cual se mantiene y modifica el alcance de "
            "la acreditación a la sociedad para producir información cuantitativa física, química y biótica; y se "
            "toman otras determinaciones." + plan)
    params = _parametros_por_laboratorio(ctx)
    sub = any(lab == LAB_SUBCONTRATADO for lab, _ in params)
    cuando = f" el {fecha_larga(ctx.fecha_muestreo)}" if ctx.fecha_muestreo else ""
    if not ctx.fecha_muestreo:
        faltantes.append("Fecha de muestreo: no se encontró en los reportes del laboratorio")
    en = f" en la empresa {ctx.cliente}" if ctx.cliente else ""
    n2 = cur.numero("Tabla")
    _p(cur, f"Se realizó un muestreo {'compuesto' if ctx.compuesto else 'puntual'}{en}{cuando} por el personal de "
            f"{EMPRESA}. Los análisis fisicoquímicos de las muestras de agua fueron realizados por {EMPRESA}"
            + (" y laboratorios subcontratados" if sub else "") + f", según se especifica en la Tabla {n2}.")
    cur.leyenda("Tabla", "Parámetros medidos para análisis de agua residual")
    t2 = cur.tabla([["Laboratorio encargado del análisis", "Parámetro"]], [list(x) for x in params],
                   anchos=[Cm(8), Cm(7.5)], izquierda=(1,))
    _agrupar_columna(t2, 0, 1)
    cur.fuente(anio=ctx.anio)
    _p(cur, _texto_comparacion(ctx))


def _agrupar_columna(t: Table, col: int, n_enc: int):
    """Combina verticalmente las celdas iguales consecutivas de una columna."""
    filas = list(t.rows)
    ini = n_enc
    while ini < len(filas):
        fin = ini
        valor = filas[ini].cells[col].text
        while fin + 1 < len(filas) and filas[fin + 1].cells[col].text == valor:
            fin += 1
        if fin > ini:
            _combinar(t, ini, col, fin, col, valor)
        ini = fin + 1


def _texto_comparacion(ctx: Contexto) -> str:
    if not ctx.columnas:
        return ("Los resultados obtenidos se reportan sin comparación normativa, dado que no se definieron las "
                "actividades de la Resolución 0631 de 2015 aplicables al vertimiento.")
    partes = []
    for c in ctx.columnas:
        if c.alcantarillado:
            continue
        a = res0631.ACTIVIDADES[c.actividad]
        partes.append(f"el Artículo {a['articulo']} de la Resolución 0631 de 2015, para la actividad «{a['actividad']}» "
                      f"({a['sector'].lower()})")
    texto = ("Los resultados obtenidos de las muestras tomadas fueron comparados con los límites establecidos en "
             + lista(partes) + ", expedida por el Ministerio de Ambiente y Desarrollo Sostenible (MADS).")
    if ctx.proyecto.alcantarillado:
        texto += (" Al tratarse de un vertimiento realizado al alcantarillado público, se aplica el Artículo 16 "
                  "«Vertimientos puntuales de aguas residuales no domésticas (ARnD) al alcantarillado público», que "
                  "ajusta los límites de la actividad para este tipo de descarga.")
    return texto


def _objetivos(doc, ctx: Contexto):
    general = _titulo(doc, 2, "objetivo general")
    especificos = _titulo(doc, 2, "objetivos especificos")
    numpr = _numpr(doc, especificos)
    donde = f" ubicada en el municipio de {ctx.ubicacion}" if ctx.ubicacion else ""
    quien = f"la empresa {ctx.cliente}" if ctx.cliente else "el cliente"
    if general is not None:
        _vaciar(general)
        _p(Cursor(doc, general), f"Determinar la calidad y/o características fisicoquímicas del {ctx.tipo_largo} "
                                 f"({ctx.siglas}) proveniente de {quien}{donde}, con el fin de dar cumplimiento a lo "
                                 "establecido por la Resolución 0631 de 2015 del Ministerio de Ambiente y Desarrollo "
                                 "Sostenible (MADS).")
    if especificos is not None:
        _vaciar(especificos)
        cur = Cursor(doc, especificos)
        tipo = "compuesto" if ctx.compuesto else "puntual"
        ubicados = "ubicado" if len(ctx.puntos) == 1 else "ubicados"
        _vineta(cur, f"Realizar la toma de muestra de tipo {tipo} en {ctx.los_puntos(ctx.puntos)}, {ubicados} en las "
                     f"instalaciones de {quien}.", numpr)
        _vineta(cur, "Determinar la calidad fisicoquímica del vertimiento de acuerdo con los resultados obtenidos para "
                     "cada uno de los parámetros evaluados.", numpr)
        _vineta(cur, "Presentar las respectivas conclusiones de acuerdo con los resultados reportados en el presente "
                     "monitoreo.", numpr)
        if ctx.columnas:
            texto = ("Comparar los resultados con los límites máximos permisibles para vertimientos puntuales de "
                     f"{ctx.tipo_largo} ({ctx.siglas}) establecidos en la Resolución 0631 del 17 de marzo de 2015 del "
                     f"Ministerio de Ambiente y Desarrollo Sostenible (MADS); específicamente, lo estipulado en "
                     f"{ctx.articulos()}")
            texto += (", así como en el Artículo 16, debido a que el vertimiento se realiza al alcantarillado público."
                      if ctx.proyecto.alcantarillado else ".")
            _vineta(cur, texto, numpr)


# ------------------------------------------------------- puntos de muestreo
def _descripcion_puntos(doc, ctx: Contexto, carpeta: str, faltantes):
    t = _titulo(doc, 2, "descripcion y localizacion")
    if t is None:
        faltantes.append("Sección de descripción y localización de los puntos de la plantilla")
        return
    _vaciar(t)
    cur = Cursor(doc, t)
    n_tabla = cur.numero("Tabla")
    n_img = cur.numero("Imagen")
    _p(cur, f"La Tabla {n_tabla} relaciona la descripción y localización de {ctx.los_puntos(ctx.puntos)}. "
            f"Adicionalmente, la distribución en el espacio se presenta en la Imagen {n_img}.")
    _p(cur, "El registro fotográfico ampliado de los puntos de monitoreo se relaciona en el Anexo 3, mientras que en el "
            "Anexo 1 se registran los formatos de campo como evidencia de la información recolectada.")
    cur.leyenda("Tabla", "Descripción de los puntos de monitoreo")
    sin_desc = []
    for rp in ctx.puntos:
        c = coordenadas(rp.punto)
        fecha = ""
        if rp.informe and rp.informe.fecha_muestreo:
            fecha = rp.informe.fecha_muestreo.date().isoformat()
        if rp.campo and rp.campo.mediciones:
            ms = [m for m in rp.campo.mediciones if m.hora]
            if ms:
                fecha = f"{fecha} / {hora(ms[0].hora)}–{hora(ms[-1].hora)}".strip(" /")
        elif rp.informe and rp.informe.fecha_muestreo:
            fecha = f"{fecha} / {rp.informe.fecha_muestreo.strftime('%H:%M')}"
        muestra = rp.informe.muestra if rp.informe else "---"
        desc = rp.punto.descripcion.strip()
        if not desc:
            sin_desc.append(rp.punto.nombre)
        foto = rp.punto.foto_ruta if rp.punto.foto_ruta and os.path.exists(rp.punto.foto_ruta) else None
        filas = [
            [rp.punto.nombre, fecha or "---", rp.punto.tipo_agua, muestra, c["lat"], c["lon"]],
            ["", "", "", "", "", ""], ["", "", "", "", "", ""],
            ["", "", "", "", c["norte"], c["este"]],
            ["", "", "", "", "", ""], ["", "", "", "", "", ""],
        ]
        if foto:
            filas.append(["", "", "", "", "", ""])
        tabla = cur.tabla([[("Punto de monitoreo", 1, 2), ("Fecha y hora de muestreo", 1, 2), ("Tipo de Agua", 1, 2),
                            ("Muestra", 1, 2), ("Coordenadas Geográficas", 2, 1)], ["Latitud", "Longitud"]], filas,
                          anchos=[Cm(3.2), Cm(2.6), Cm(1.6), Cm(1.8), Cm(3.15), Cm(3.15)])
        for col, valor in enumerate([rp.punto.nombre, fecha or "---", rp.punto.tipo_agua, muestra]):
            _combinar(tabla, 2, col, 5, col, valor)
        _combinar(tabla, 3, 4, 3, 5, "Coordenadas Origen Nacional Único", SUBENCABEZADO, negrita=True)
        set_cell_text(tabla.cell(4, 4), "Norte (m)")
        set_cell_text(tabla.cell(4, 5), "Este (m)")
        for col in (4, 5):
            _combinar(tabla, 4, col, 4, col, tabla.cell(4, col).text, SUBENCABEZADO, negrita=True)
        _combinar(tabla, 6, 0, 6, 5, "Descripción del punto de monitoreo", SUBENCABEZADO, negrita=True)
        _combinar(tabla, 7, 0, 7, 5, desc or "---", izquierda=True)
        if foto:
            celda = _combinar(tabla, 8, 0, 8, 5, "")
            run = celda.paragraphs[0].add_run()
            run.add_picture(foto, width=Cm(7))
        cur.fuente(anio=ctx.anio)
    if sin_desc:
        faltantes.append("Descripción de los puntos: agregue la descripción de " + ", ".join(sin_desc))
    cur.leyenda("Imagen", "Localización de los puntos de monitoreo")
    ruta = os.path.join(carpeta, "localizacion_vertimientos.png")
    try:
        from .isophones import generar_mapa_localizacion

        puntos = []
        for k, rp in enumerate(ctx.puntos, start=1):
            c = coordenadas(rp.punto)
            if c["dd"]:
                puntos.append(SimpleNamespace(no_punto=k, nombre=rp.punto.nombre, longitud=str(c["dd"][1]),
                                              latitud=str(c["dd"][0])))
        if not puntos:
            raise ValueError("sin coordenadas")
        proyecto = SimpleNamespace(nombre_proyecto=ctx.proyecto.nombre_proyecto, cliente=ctx.cliente, puntos=puntos)
        generar_mapa_localizacion(proyecto, ruta, con_basemap=os.environ.get("ENGINE_BASEMAP", "1") != "0",
                                  titulo="LOCALIZACIÓN DE PUNTOS - VERTIMIENTOS")
        cur.imagen(ruta, ANCHO)
    except Exception:  # noqa: BLE001
        cur.parrafo("[Plano de localización: registre las coordenadas de los puntos]", alinear=WD_ALIGN_PARAGRAPH.CENTER)
        faltantes.append("Imagen de localización: registre las coordenadas de los puntos")
    cur.fuente(anio=ctx.anio)


# ------------------------------------------------------- metodos de analisis
def _tabla_metodos(doc, ctx: Contexto, faltantes):
    """Tabla 'Metodología de preservación, técnicas y métodos utilizados' con los ensayos del proyecto."""
    leyenda = next((el for el in _cuerpo(doc) if _es_leyenda(el)
                    and "metodologia de preservacion" in _sin_tildes(_texto(el))), None)
    if leyenda is None:
        faltantes.append("Tabla de métodos y preservación de la plantilla")
        return
    # Quita la tabla y las notas de la plantilla (se reconstruyen).
    el = leyenda.getnext()
    while el is not None:
        siguiente = el.getnext()
        if el.tag not in (f"{_W}p", f"{_W}tbl"):
            if el.tag == f"{_W}sectPr":
                break
            el = siguiente  # marcadores (bookmarkEnd) entre la tabla y las notas
            continue
        if el.tag == f"{_W}p" and (_nivel(el) is not None or (_texto(el).strip() and _estilo(el) not in (
                "NotaTabla", "Fuente"))):
            break
        _quitar(el)
        el = siguiente
    cur = Cursor(doc, leyenda)
    filas, hap = [], []
    sub = False
    for f in ctx.res.filas:
        if f.parametro.grupo == "insitu":
            continue
        tecnica, preservacion = tx.tecnica_preservacion(f.parametro.clave, f.parametro.grupo)
        nombre = f.parametro.nombre + (" (1)" if f.subcontratado else "")
        sub = sub or f.subcontratado
        fila = [nombre, f.unidad, f.metodo, f.lcm, f.incertidumbre, tecnica, preservacion]
        (hap if f.parametro.grupo == "hap" else filas).append(fila)
    seccion_hap = None
    if hap:
        seccion_hap = len(filas)
        filas.append([""] * 7)
        filas += hap
    t = cur.tabla([["Parámetro", "Unidad", "Método", "L.C.M.", "Incertidumbre (%)", "Técnica", "Preservación"]],
                  filas, anchos=[Cm(3.6), Cm(1.8), Cm(2.6), Cm(1.5), Cm(1.7), Cm(2.2), Cm(2.6)], izquierda=(0,))
    if seccion_hap is not None:
        _combinar(t, 1 + seccion_hap, 0, 1 + seccion_hap, 6, "Hidrocarburos Aromáticos Policíclicos (HAP)",
                  SUBENCABEZADO, negrita=True)
    notas = ["(1) Parámetro subcontratado con laboratorio acreditado por el IDEAM."] if sub else []
    notas.append("N.A.: No Aplica. / N.E.: No Especifica. / L.C.M.: Límite de cuantificación del método.")
    for n in notas:
        cur.parrafo(n, "Nota Tabla")
    cur.parrafo("Fuente: Standard Methods for the Examination of Water and Wastewater.", "Fuente")


# ---------------------------------------------------------------- normatividad
def _normatividad(doc, ctx: Contexto, n_tabla_lab: Optional[int]):
    t = _titulo(doc, 1, "normatividad")
    if t is None:
        return
    _vaciar(t)
    cur = Cursor(doc, t)
    if not ctx.columnas:
        _p(cur, "Para el presente estudio no se definieron las actividades de la Resolución 0631 de 2015 aplicables al "
                "vertimiento; los resultados se presentan sin comparación normativa.")
        return
    donde = ctx.los_puntos(ctx.evaluados) if ctx.evaluados else "los puntos de monitoreo"
    de = f" de {ctx.cliente}" if ctx.cliente else ""
    for k, c in enumerate([c for c in ctx.columnas if not c.alcantarillado]):
        a = res0631.ACTIVIDADES[c.actividad]
        inicio = (f"Los resultados de las muestras tomadas en {donde}{de} fueron comparados con los límites "
                  f"establecidos en el Artículo {a['articulo']}" if k == 0 else
                  f"Así mismo, se tuvieron en cuenta los límites del Artículo {a['articulo']}")
        _p(cur, f"{inicio} de la Resolución 0631 de 2015, expedida por el Ministerio de Ambiente y Desarrollo "
                f"Sostenible (MADS), que define los parámetros fisicoquímicos a monitorear y sus valores límites "
                f"máximos permisibles en los vertimientos puntuales a cuerpos de aguas superficiales de "
                f"{a['sector'][0].lower() + a['sector'][1:]}, para la actividad «{a['actividad']}».")
    if ctx.proyecto.alcantarillado:
        _p(cur, "Al tratarse de un vertimiento al alcantarillado público, se aplica el Artículo 16 de la misma "
                "resolución, según el cual los vertimientos puntuales de aguas residuales no domésticas (ARnD) al "
                "alcantarillado público deben cumplir los valores límites máximos permisibles de la actividad "
                "correspondiente multiplicados por un factor de 1,50 para DQO, DBO5, SST, SSED, grasas y aceites y "
                "los compuestos de fósforo y nitrógeno, y las mismas exigencias de la actividad para los demás "
                "parámetros; el pH debe estar entre 5,00 y 9,00 unidades.")
    if ctx.proyecto.consumo_humano:
        _p(cur, "Dado que el cuerpo receptor tiene uso para consumo humano y doméstico, y pecuario, los Hidrocarburos "
                "Aromáticos Policíclicos (HAP) deben cumplir un valor límite máximo permisible de 0,01 mg/L.")
    _p(cur, "Adicionalmente, el Artículo 5 de la Resolución 0631 de 2015 establece un valor máximo permisible de "
            "temperatura de 40,00 °C para los vertimientos puntuales.")
    if n_tabla_lab:
        _p(cur, f"La Tabla {n_tabla_lab} presenta los límites de referencia de los parámetros objeto de estudio.")


# ---------------------------------------------------------------- resultados
def _serie(campo, atributo):
    salida = []
    for m in campo.mediciones:
        v = getattr(m, atributo)
        if isinstance(v, str):
            v = numero(v)
        salida.append(v)
    return salida


def _horas(campo) -> list[str]:
    return [hora(m.hora) for m in campo.mediciones]


def _extremos(horas, valores):
    pares = [(h, v) for h, v in zip(horas, valores) if v is not None]
    if not pares:
        return None
    mn = min(pares, key=lambda x: x[1])
    mx = max(pares, key=lambda x: x[1])
    return mn, mx


def _tabla_campo(cur: Cursor, ctx: Contexto, rp: ResultadoPunto):
    campo = rp.campo
    con_od = any(m.oxigeno is not None for m in campo.mediciones)
    enc = [("Hora", 1, 2), ("Volumen (mL)", 1, 2), ("Tiempo (s)", 1, 2), ("Caudal (mL/s)", 1, 2),
           ("pH (Unidades)", 1, 2), ("pH (Unidades) Duplicado", 1, 2), ("Temperatura (°C)", 1, 2),
           ("Temperatura (°C) Duplicado", 1, 2), ("Conductividad (µS/cm)", 1, 2),
           ("Conductividad (µS/cm) Duplicado", 1, 2)]
    if con_od:
        enc += [("OD (mg O2/L)", 1, 2), ("OD (mg O2/L) Duplicado", 1, 2)]
    enc += [("Sólidos sedimentables (mL/L-h)", 1, 2), ("Sólidos sedimentables (mL/L-h) Duplicado", 1, 2),
            ("Alícuotas", 2, 1)]
    ncol = len(enc) + 1
    filas = []
    for m in campo.mediciones:
        fr, al = campo.fraccion(m), campo.alicuota_ml(m)
        fila = [hora(m.hora), num(m.volumen_ml, 0), num(m.tiempo_s, 2), num(m.caudal_mls, 1), num(m.ph, 2),
                num(m.ph_dup, 2), num(m.temperatura, 2), num(m.temperatura_dup, 2), num(m.conductividad, 0),
                num(m.conductividad_dup, 0)]
        if con_od:
            fila += [num(m.oxigeno, 2), num(m.oxigeno_dup, 2)]
        fila += [m.sedimentables or "-", m.sedimentables_dup or "-",
                 num(fr * 100 if fr is not None else None, 2), num(al, 1)]
        filas.append([x.replace("---", "-") for x in fila])
    # Resumen: suma, promedio, maximo y minimo por columna numerica.
    columnas = [("caudal_mls", 1), ("ph", 2), ("ph_dup", 2), ("temperatura", 2), ("temperatura_dup", 2),
                ("conductividad", 0), ("conductividad_dup", 0)]
    if con_od:
        columnas += [("oxigeno", 2), ("oxigeno_dup", 2)]
    columnas += [("sedimentables", 1), ("sedimentables_dup", 1)]
    series = {a: [v for v in _serie(campo, a) if v is not None] for a, _ in columnas}
    total_al = sum(campo.alicuota_ml(m) or 0 for m in campo.mediciones)
    for etiqueta, f in (("Suma", None), ("Promedio", lambda xs: sum(xs) / len(xs)), ("Máximo", max),
                        ("Mínimo", min)):
        fila = [etiqueta, "", ""]
        for a, dec in columnas:
            xs = series[a]
            if etiqueta == "Suma":
                fila.append(num(sum(xs), dec) if a == "caudal_mls" and xs else "N.A.")
            elif etiqueta == "Mínimo" and a.startswith("sedimentables") and xs:
                # Un minimo "<0,1" se reporta como tal (no como 0,1).
                textos = [getattr(m, a) for m in campo.mediciones if getattr(m, a)]
                bajo = [t for t in textos if t.strip().startswith("<") and numero(t) == min(xs)]
                fila.append(bajo[0].strip() if bajo else num(min(xs), dec))
            else:
                fila.append(num(f(xs), dec) if xs else "-")
        fila += (["100,00", num(total_al, 0)] if etiqueta == "Suma" else
                 ["Tamaño de muestra (mL)", num(campo.tamano_muestra_ml, 0)] if etiqueta == "Promedio" else ["", ""])
        filas.append(fila)
    t = cur.tabla([[(rp.punto.nombre.upper(), ncol, 1)], enc, ["Fracción (%)", "Volumen teórico (mL)"]], filas)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7)
    n = len(campo.mediciones) + 3
    for k in range(4):
        _combinar(t, n + k, 0, n + k, 2, filas[len(campo.mediciones) + k][0], negrita=True)
    _combinar(t, n + 1, ncol - 2, n + 3, ncol - 2, "Tamaño de muestra (mL)")
    _combinar(t, n + 1, ncol - 1, n + 3, ncol - 1, num(campo.tamano_muestra_ml, 0))
    return t


def _texto_tabla_lab(ctx: Contexto) -> str:
    return (f"presenta los resultados del análisis fisicoquímico realizado al agua proveniente de "
            f"{ctx.los_puntos(ctx.con_lab)}, siguiendo los métodos establecidos por el Standard Methods for the "
            "Examination of Water and Wastewater"
            + (", y comparándolos con los límites normativos establecidos por los artículos mencionados."
               if ctx.columnas else "."))


def _tabla_laboratorio(cur: Cursor, ctx: Contexto) -> int:
    puntos = ctx.con_lab
    cols = ctx.columnas
    n = cur.numero("Tabla")
    _p(cur, f"La Tabla {n} {_texto_tabla_lab(ctx)}")
    cur.leyenda("Tabla", "Resultados del análisis de laboratorio vs Resolución 0631 de 2015" if cols
                else "Resultados del análisis de laboratorio")
    enc0 = [("Parámetros", 1, 2), ("Unidades", 1, 2)] + [rp.punto.nombre for rp in puntos] \
        + [(c.titulo, 1, 2) for c in cols]
    enc1 = [rp.informe.muestra for rp in puntos]
    filas, rellenos, secciones = [], {}, []
    hap_inicio = None
    for f in ctx.res.filas:
        if f.parametro.grupo == "hap" and hap_inicio is None:
            hap_inicio = len(filas)
            secciones.append(len(filas))
            filas.append([""] * (2 + len(puntos) + len(cols)))
        nombre = f.parametro.nombre + (" (1)" if f.subcontratado else "")
        fila = [nombre, f.unidad]
        for k, rp in enumerate(puntos):
            r = f.resultados.get(rp.punto.nombre)
            fila.append(r.reporte if r else "---")
            estados = [f.conformidad.get((rp.punto.nombre, i)) for i, _ in ctx.aplican]
            if "No cumple" in estados:
                rellenos[(len(filas), 2 + k)] = NO_CUMPLE
        for lim in f.limites:
            fila.append(res0631.texto_limite(f.parametro.clave, lim) + lim.marca)
        filas.append(fila)
    anchos = None
    t = cur.tabla([enc0, enc1] if puntos else [enc0], filas, rellenos, anchos, izquierda=(0,))
    n_enc = 2 if puntos else 1
    for s in secciones:
        _combinar(t, n_enc + s, 0, n_enc + s, len(filas[0]) - 1, "Hidrocarburos Aromáticos Policíclicos (HAP)",
                  SUBENCABEZADO, negrita=True)
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    notas = []
    if any(c.alcantarillado for c in cols):
        notas += ["*Se aplican las mismas exigencias establecidas para el parámetro respectivo en la actividad "
                  "específica para los vertimientos puntuales a cuerpos de agua superficial multiplicados por un "
                  "factor de 1,50.",
                  "**Se aplican las mismas exigencias establecidas para el parámetro respectivo en la actividad "
                  "específica para los vertimientos puntuales a cuerpos de agua superficial."]
    if any(f.subcontratado for f in ctx.res.filas):
        notas.append("(1) Parámetro subcontratado con laboratorio acreditado por el IDEAM.")
    if cols:
        notas.append("Temperatura: el Artículo 5 de la Resolución 0631 de 2015 establece un valor máximo permisible "
                     "de 40,00 °C.")
        notas.append("N.E.: No Especifica. Las celdas sombreadas indican resultados que no cumplen el límite "
                     "aplicable.")
    _fuente_y_notas(cur, notas, ctx.anio)
    inc = ctx.res.incumplimientos()
    if ctx.evaluados and ctx.aplican:
        if inc:
            detalles = [f"{minus(f.parametro.nombre)} en el punto «{p}» ({f.resultados[p].reporte} {f.unidad} "
                        f"frente a {res0631.texto_limite(f.parametro.clave, f.limites[c])} {f.unidad}, "
                        f"{ctx.columnas[c].titulo})" for p, f, c in inc]
            _p(cur, "De acuerdo con la regla de decisión de aceptación simple, no cumplen con los límites máximos "
                    f"permisibles: {lista(detalles)}. Los demás parámetros con límite establecido cumplen con la "
                    "Resolución 0631 de 2015.")
        else:
            _p(cur, "De acuerdo con la regla de decisión de aceptación simple, todos los parámetros con límite máximo "
                    f"permisible establecido cumplen con la Resolución 0631 de 2015 en {ctx.los_puntos(ctx.evaluados)}.")
    return n


# ------------------------------------------------------------- in situ
def _limites_insitu(ctx: Contexto, clave: str, unidad: str) -> list[tuple[str, float]]:
    """Lineas de limite para las graficas de campo (solo de los puntos que se evaluan)."""
    if clave == "temperatura":
        return [("Límite máximo permisible: 40,00 °C (Art. 5)", 40.0)]
    por_valor: dict = {}
    for _, c in ctx.aplican:
        lim = res0631.limite(c.actividad, clave, c.alcantarillado, ctx.proyecto.consumo_humano)
        if lim.tipo == "rango":
            por_valor.setdefault(("Límite mínimo", lim.minimo), []).append(c.titulo)
            por_valor.setdefault(("Límite máximo", lim.maximo), []).append(c.titulo)
        elif lim.tipo == "max":
            por_valor.setdefault(("Límite máximo permisible", lim.maximo), []).append(c.titulo)
    return [(f"{etiqueta}: {res0631.fmt_limite(v)} {unidad} ({' / '.join(titulos)})", v)
            for (etiqueta, v), titulos in por_valor.items()]


def _analisis_insitu(cur: Cursor, ctx: Contexto, carpeta: str):
    puntos = ctx.con_campo
    cur.titulo("ANÁLISIS DE PARÁMETROS IN SITU", 2)
    # Caudal.
    cur.titulo("Caudal", 3)
    frases = []
    for rp in puntos:
        hs, qs = _horas(rp.campo), [m.caudal_mls for m in rp.campo.mediciones]
        e = _extremos(hs, qs)
        if not e:
            continue
        ms = [m for m in rp.campo.mediciones if m.hora]
        rango = f"entre las {hora(ms[0].hora)} y las {hora(ms[-1].hora)} horas" if ms else ""
        if e[0][1] == e[1][1]:
            frases.append(f"en el punto «{rp.punto.nombre}» ({rango}) el caudal se mantuvo en {num(e[0][1])} mL/s "
                          "durante toda la jornada")
        else:
            frases.append(f"en el punto «{rp.punto.nombre}» ({rango}) el caudal varió entre {num(e[0][1])} mL/s "
                          f"({e[0][0]} horas) y {num(e[1][1])} mL/s ({e[1][0]} horas)")
    n = cur.numero("Gráfica")
    refs = lista([f"la Gráfica {n + k}" for k in range(len(puntos))])
    _p(cur, "La medición de caudal se efectuó por el método volumétrico durante el muestreo compuesto: "
            + "; ".join(frases) + f", como se observa en {refs}.")
    for rp in puntos:
        ruta = gr.caudal(_horas(rp.campo), [m.caudal_mls for m in rp.campo.mediciones],
                         gr.nombre_archivo(carpeta, "caudal", rp.punto.nombre))
        cur.leyenda("Gráfica", f"Comportamiento del caudal – {rp.punto.nombre}")
        cur.imagen(ruta, ANCHO)
        cur.fuente(anio=ctx.anio)
    # Conductividad, pH, OD, solidos sedimentables y temperatura.
    params = [("conductividad", "Conductividad eléctrica", "µS/cm", "conductividad", 0, "Conductividad (µS/cm)"),
              ("ph", "Potencial de hidrogeniones (pH)", "Unidades de pH", "ph", 2, "pH (Unidades de pH)"),
              ("oxigeno", "Oxígeno disuelto", "mg O2/L", "oxigeno", 2, "Oxígeno disuelto (mg O2/L)"),
              ("sedimentables", "Sólidos Sedimentables", "mL/L-h", "solidos_sedimentables", 1,
               "Sólidos sedimentables (mL/L-h)"),
              ("temperatura", "Temperatura", "°C", "temperatura", 2, "Temperatura (°C)")]
    for atributo, titulo, unidad, clave, dec, eje in params:
        if not any(v is not None for rp in puntos for v in _serie(rp.campo, atributo)):
            continue
        cur.titulo(titulo, 3)
        intro = tx.INSITU[atributo if atributo != "sedimentables" else "sedimentables"]
        n_ph = None
        if atributo == "ph":
            n_ph = cur.numero("Tabla")
            intro += f" La Tabla {n_ph} presenta los rangos de pH."
        _p(cur, intro)
        if atributo == "ph":
            cur.leyenda("Tabla", "Rangos de pH")
            cur.tabla([["pH", "Rango"]], [[r[2], r[3]] for r in tx.RANGOS_PH], anchos=[Cm(4), Cm(6)])
            cur.parrafo("Fuente: adaptado de la clasificación de reacción del agua (pH).", "Fuente")
        n = cur.numero("Gráfica")
        frases = []
        for k, rp in enumerate(puntos):
            hs, vs = _horas(rp.campo), _serie(rp.campo, atributo)
            e = _extremos(hs, vs)
            if not e:
                continue
            (h0, v0), (h1, v1) = e
            if atributo == "sedimentables":
                textos = [m.sedimentables for m in rp.campo.mediciones if m.sedimentables]
                bajo = [t for t in textos if t.strip().startswith("<")]
                minimo = f"inferior al límite de cuantificación del método ({bajo[0]} {unidad})" if bajo \
                    else f"{num(v0, dec)} {unidad} ({h0} horas)"
                frase = f"en el punto «{rp.punto.nombre}» se reportaron valores entre {minimo} y {num(v1, dec)} " \
                        f"{unidad} ({h1} horas)"
            else:
                frase = f"en el punto «{rp.punto.nombre}» los valores oscilaron entre {num(v0, dec)} {unidad} " \
                        f"({h0} horas) y {num(v1, dec)} {unidad} ({h1} horas)"
            if atributo == "ph":
                c0, c1 = tx.clase_ph(v0), tx.clase_ph(v1)
                frase += f", con una condición {c0}" if c0 == c1 else f", con condiciones entre {c0} y {c1}"
            frase += f" (ver Gráfica {n + k})"
            frases.append(frase)
        texto = "Durante el presente monitoreo, " + "; ".join(frases) + "."
        evaluados = [rp for rp in puntos if rp.punto.evaluar]
        limites = _limites_insitu(ctx, clave, unidad) if atributo in ("ph", "sedimentables", "temperatura") \
            and evaluados else []
        if limites:
            if atributo == "temperatura":
                texto += (" Desde el punto de vista normativo, el Artículo 5 de la Resolución 0631 de 2015 establece "
                          "un límite máximo permisible de 40,00 °C.")
            else:
                texto += (" Los límites de la Resolución 0631 de 2015 aplicables a "
                          f"{ctx.los_puntos(evaluados, False)} se muestran en las gráficas.")
        _p(cur, texto)
        for rp in puntos:
            bajos = [bool(m.sedimentables and m.sedimentables.strip().startswith("<"))
                     for m in rp.campo.mediciones] if atributo == "sedimentables" else None
            ruta = gr.in_situ(_horas(rp.campo), _serie(rp.campo, atributo), eje,
                              gr.nombre_archivo(carpeta, atributo, rp.punto.nombre),
                              limites if rp.punto.evaluar else None, decimales=dec, bajos=bajos)
            cur.leyenda("Gráfica", contraer(f"Comportamiento de {ARTICULO.get(atributo, minus(titulo))} – "
                                            f"{rp.punto.nombre}"))
            cur.imagen(ruta, ANCHO)
            cur.fuente(anio=ctx.anio)


# ----------------------------------------------------------- laboratorio
def _valor_texto(r, unidad: str) -> str:
    if r.menor_que_lcm:
        return f"un valor inferior al límite de cuantificación del método ({r.reporte} {unidad})"
    return f"{r.reporte} {unidad}"


def _frase_resultados(ctx: Contexto, f: FilaComparacion) -> str:
    rs = [(rp, f.resultados.get(rp.punto.nombre)) for rp in ctx.con_lab]
    rs = [(rp, r) for rp, r in rs if r is not None]
    if not rs:
        return ""
    nombre = minus(f.parametro.nombre)
    if len(rs) > 1 and all(r.menor_que_lcm for _, r in rs) and len({r.reporte for _, r in rs}) == 1:
        return (f"Para el parámetro {nombre} se reportó {_valor_texto(rs[0][1], f.unidad)} en todos los puntos de "
                "monitoreo.")
    partes = [f"{_valor_texto(r, f.unidad)} en el punto «{rp.punto.nombre}»" for rp, r in rs]
    return f"Para el parámetro {nombre} se reportó {lista(partes)}."


def _frase_norma(ctx: Contexto, f: FilaComparacion) -> str:
    if not ctx.evaluados or not ctx.aplican:
        return ""
    con_juicio = [(i, c) for i, c in ctx.aplican if f.limites[i].con_juicio]
    if not con_juicio:
        return ("Los artículos aplicables de la Resolución 0631 de 2015 no establecen un límite máximo permisible para "
                "este parámetro (análisis y reporte o no especificado), razón por la cual no se emite juicio "
                "normativo; sin embargo, se realiza su respectivo análisis y reporte.")
    limites = lista([f"{res0631.texto_limite(f.parametro.clave, f.limites[i])} {f.unidad} en el {c.titulo}"
                     for i, c in con_juicio])
    estados = []
    for rp in ctx.evaluados:
        es = [f.conformidad.get((rp.punto.nombre, i)) for i, _ in con_juicio]
        if "No cumple" in es:
            estados.append(f"el punto «{rp.punto.nombre}» no cumple con el límite máximo permisible")
        elif es and all(e == "Cumple" for e in es):
            estados.append(f"el punto «{rp.punto.nombre}» cumple con el límite máximo permisible")
    texto = f"El límite máximo permisible establecido es de {limites}"
    return texto + (f"; de acuerdo con lo anterior, {lista(estados)}." if estados else ".")


def _limites_grafica(ctx: Contexto, f: FilaComparacion) -> list[tuple[str, float]]:
    if not ctx.evaluados:
        return []
    salida, vistos = [], {}
    for i, c in ctx.aplican:
        lim = f.limites[i]
        if lim.tipo == "max" and lim.maximo is not None:
            vistos.setdefault(lim.maximo, []).append(c.titulo)
    for v, titulos in vistos.items():
        salida.append((f"{' / '.join(titulos)} ({res0631.fmt_limite(v)} {f.unidad})", v))
    return salida


def _etiquetas_puntos(ctx: Contexto) -> list[str]:
    return ["\n".join(textwrap.wrap(rp.punto.nombre, 24)) for rp in ctx.con_lab]


def _grafica_lab(cur: Cursor, ctx: Contexto, filas: list[FilaComparacion], titulo: str, carpeta: str,
                 con_limites: bool = True):
    series = []
    for f in filas:
        valores, bajos = [], []
        for rp in ctx.con_lab:
            r = f.resultados.get(rp.punto.nombre)
            valores.append(r.valor if r else None)
            bajos.append(bool(r and r.menor_que_lcm))
        series.append((f.parametro.nombre, valores, bajos))
    unidades = {f.unidad for f in filas}
    unidad = unidades.pop() if len(unidades) == 1 else "mg/L"
    limites = _limites_grafica(ctx, filas[0]) if con_limites and len(filas) == 1 else []
    ruta = gr.laboratorio(series, _etiquetas_puntos(ctx), unidad, gr.nombre_archivo(carpeta, titulo), limites)
    if ruta is None:
        return None
    n = cur.leyenda("Gráfica", contraer(f"Comportamiento de {titulo}"))
    cur.imagen(ruta, ANCHO)
    cur.fuente(anio=ctx.anio)
    return n


def _seccion_parametros(cur: Cursor, ctx: Contexto, titulo: str, descripcion: str, filas: list[FilaComparacion],
                        carpeta: str, clave: str):
    cur.titulo(titulo, 3)
    if descripcion:
        _p(cur, descripcion)
    n = cur.numero("Gráfica")
    if clave == "dbo_dqo":
        for f in filas:
            _p(cur, _ref(" ".join(x for x in (_frase_resultados(ctx, f), _frase_norma(ctx, f)) if x),
                         cur.numero("Gráfica")))
            _grafica_lab(cur, ctx, [f], ARTICULO.get(f.parametro.clave, minus(f.parametro.nombre)), carpeta)
        return
    normas = [_frase_norma(ctx, f) for f in filas]
    if len(filas) > 1 and len(set(normas)) == 1:
        textos = [_frase_resultados(ctx, f) for f in filas]
        if normas[0]:
            textos.append(normas[0].replace("para este parámetro", "para estos parámetros"))
    else:
        textos = [" ".join(x for x in (_frase_resultados(ctx, f), n) if x) for f, n in zip(filas, normas)]
    if clave == "dureza":
        totales = next((f for f in filas if f.parametro.clave == "dureza_total"), None)
        if totales:
            clases = []
            for rp in ctx.con_lab:
                r = totales.resultados.get(rp.punto.nombre)
                if r and r.valor is not None:
                    clases.append(f"{tx.clase_dureza(r.valor)} en el punto «{rp.punto.nombre}»")
            n_t = cur.numero("Tabla")
            textos.append(f"De acuerdo con la Tabla {n_t}, el agua se clasifica como {lista(clases)}.")
    principal = [f for f in filas if f.parametro.clave in ("aceites_grasas", "grasas_aceites")] if clave == "aceites" \
        else filas
    graficar = principal or filas
    _p(cur, _ref(" ".join(textos), n))
    if clave == "dureza" and any(f.parametro.clave == "dureza_total" for f in filas):
        cur.leyenda("Tabla", "Clasificación de los valores de dureza")
        cur.tabla([[("Dureza Total", 2, 1)], ["Clasificación", "Dureza (mg CaCO3/L)"]],
                  [list(x) for x in tx.TABLA_DUREZA], anchos=[Cm(5), Cm(5)])
        cur.parrafo("Fuente: clasificación de la dureza del agua (mg CaCO3/L).", "Fuente")
    _grafica_lab(cur, ctx, graficar, ARTICULO.get(clave, minus(titulo)), carpeta,
                 con_limites=len(graficar) == 1)


def _hap(cur: Cursor, ctx: Contexto, filas: list[FilaComparacion], descripcion: str):
    cur.titulo("Hidrocarburos aromáticos policíclicos", 3)
    _p(cur, descripcion)
    cuantificados = []
    lcm = None
    for f in filas:
        for p, r in f.resultados.items():
            if r.menor_que_lcm:
                lcm = lcm or r.reporte
            elif r.valor is not None:
                cuantificados.append(f"{f.parametro.nombre} ({r.reporte} {f.unidad} en el punto «{p}»)")
    if cuantificados:
        texto = (f"Para el presente estudio se analizaron {cantidad(len(filas), 'compuesto', 'compuestos')}; se "
                 f"cuantificaron {lista(cuantificados)}, y los demás compuestos reportaron concentraciones inferiores "
                 "al límite de cuantificación del método.")
    else:
        cuales = "del compuesto analizado" if len(filas) == 1 else \
            f"de los {cantidad(len(filas), 'compuesto', 'compuestos')} analizados"
        texto = (f"Para el presente estudio, la concentración {cuales} fue inferior al límite de cuantificación del método ({lcm or '<LCM'} mg/L) en todos los "
                 "puntos de monitoreo.")
    norma = _frase_norma(ctx, filas[0]) if filas else ""
    _p(cur, f"{texto} {norma}".strip())


def _metales(cur: Cursor, ctx: Contexto, filas: list[FilaComparacion], descripcion: str, carpeta: str):
    cur.titulo("Metales", 3)
    _p(cur, descripcion)
    bajo = [f for f in filas if all(r.menor_que_lcm for r in f.resultados.values())]
    cuant = [f for f in filas if f not in bajo]
    texto = ("Para el presente estudio se analizaron las concentraciones de los metales: "
             + lista([minus(f.parametro.nombre) for f in filas]) + ".")
    if bajo:
        texto += (" De estos, " + lista([minus(f.parametro.nombre) for f in bajo])
                  + (" reportó" if len(bajo) == 1 else " reportaron")
                  + " una concentración inferior al límite de cuantificación del método de análisis en todos los "
                    "puntos de monitoreo.")
    if cuant:
        texto += " A continuación, se presenta el análisis correspondiente para los metales cuantificados."
    _p(cur, texto)
    if bajo and ctx.evaluados and ctx.aplican:
        inc = [f for f in bajo if any(e == "No cumple" for e in f.conformidad.values())]
        if not inc:
            _p(cur, "Los metales con resultados inferiores al límite de cuantificación cumplen con los límites "
                    "máximos permisibles establecidos para ellos en la norma de referencia.")
    for f in cuant:
        cur.titulo(f.parametro.nombre, 4)
        n = cur.numero("Gráfica")
        _p(cur, _ref(" ".join(x for x in (_frase_resultados(ctx, f), _frase_norma(ctx, f)) if x), n))
        _grafica_lab(cur, ctx, [f], f"el {minus(f.parametro.nombre)}", carpeta)


def _analisis_laboratorio(cur: Cursor, ctx: Contexto, carpeta: str):
    filas = [f for f in ctx.res.filas if f.parametro.grupo != "insitu"]
    if not filas:
        return
    cur.titulo("ANÁLISIS DE PARÁMETROS FISICOQUÍMICOS", 2)
    usados = set()
    for clave, titulo, claves, descripcion in tx.SECCIONES:
        if clave == "hap":
            grupo = [f for f in filas if f.parametro.grupo == "hap"]
            if grupo:
                _hap(cur, ctx, grupo, descripcion)
                usados.update(f.parametro.clave for f in grupo)
            continue
        if clave == "metales":
            grupo = [f for f in filas if f.parametro.grupo == "metal"]
            if grupo:
                _metales(cur, ctx, grupo, descripcion, carpeta)
                usados.update(f.parametro.clave for f in grupo)
            continue
        grupo = [f for f in filas if f.parametro.clave in claves]
        if not grupo:
            continue
        _seccion_parametros(cur, ctx, titulo, descripcion, grupo, carpeta, clave)
        usados.update(f.parametro.clave for f in grupo)
    for f in filas:
        if f.parametro.clave in usados:
            continue
        _seccion_parametros(cur, ctx, f.parametro.nombre, "", [f], carpeta, f.parametro.clave)


def _resultados(doc, ctx: Contexto, carpeta: str, faltantes) -> Optional[int]:
    t = _titulo(doc, 1, "resultados")
    if t is None:
        faltantes.append("Capítulo de resultados de la plantilla")
        return None
    _vaciar(t)
    cur = Cursor(doc, t)
    if ctx.con_campo:
        cur.titulo("RESULTADOS DE PARÁMETROS IN SITU", 2)
        n = cur.numero("Tabla")
        refs = lista([f"{'La' if k == 0 else 'la'} Tabla {n + k}" for k in range(len(ctx.con_campo))])
        ms = ctx.con_campo[0].campo.mediciones
        horas = [m.hora for m in ms if m.hora]
        detalle = ""
        if len(horas) >= 2:
            base = dt.date(2000, 1, 1)
            inicio, fin = dt.datetime.combine(base, horas[0]), dt.datetime.combine(base, horas[-1])
            intervalo = (dt.datetime.combine(base, horas[1]) - inicio).seconds // 60
            duracion = round((fin - inicio).seconds / 3600)
            horas_txt = "una (1) hora" if duracion == 1 else cantidad(duracion, "hora", "horas")
            detalle = (f" El muestreo se realizó durante {horas_txt}, con mediciones cada "
                       f"{intervalo} minutos para caudal, conductividad, temperatura y pH.")
        _p(cur, f"{refs} "
                f"{'presenta' if len(ctx.con_campo) == 1 else 'presentan'} los resultados obtenidos en campo durante "
                f"el muestreo en {ctx.los_puntos(ctx.con_campo)}.{detalle}")
        for rp in ctx.con_campo:
            cur.leyenda("Tabla", f"Resultados de parámetros in situ – {rp.punto.nombre}")
            _tabla_campo(cur, ctx, rp)
            _fuente_y_notas(cur, ["N.A.: No Aplica."], ctx.anio)
    elif ctx.compuesto:
        faltantes.append("Datos de campo: suba la FP-004 para las tablas y gráficas de parámetros in situ")
    n_lab = None
    if ctx.con_lab:
        cur.titulo("RESULTADOS MEDIDOS EN EL LABORATORIO", 2)
        n_lab = _tabla_laboratorio(cur, ctx)
    if ctx.con_campo:
        _analisis_insitu(cur, ctx, carpeta)
    if ctx.con_lab:
        _analisis_laboratorio(cur, ctx, carpeta)
    return n_lab


# --------------------------------------------------------------- conclusiones
def _conclusiones(doc, ctx: Contexto):
    t = _titulo(doc, 1, "conclusiones")
    if t is None:
        return
    numpr = _numpr(doc, t)
    _vaciar(t)
    cur = Cursor(doc, t)
    quien = f"la empresa {ctx.cliente}" if ctx.cliente else "el cliente"
    donde = f" en el municipio de {ctx.ubicacion}" if ctx.ubicacion else ""
    _p(cur, f"Una vez realizado el monitoreo solicitado por {quien}{donde}"
            + (" y de acuerdo con los límites establecidos en la Resolución 0631 de 2015 expedida por el Ministerio "
               "de Ambiente y Desarrollo Sostenible (MADS)" if ctx.columnas else "")
            + ", se puede concluir lo siguiente:")
    inc = ctx.res.incumplimientos()
    titulos = lista([f"el {t}" for t in sorted({c.titulo for _, c in ctx.aplican})])
    for rp in ctx.evaluados:
        if not ctx.aplican:
            break
        propios = sorted({minus(f.parametro.nombre) for p, f, _ in inc if p == rp.punto.nombre})
        if propios:
            texto = (f"La mayoría de los parámetros analizados en el punto «{rp.punto.nombre}» reportaron resultados "
                     f"dentro de los límites máximos permisibles establecidos en {titulos} de la Resolución 0631 de "
                     f"2015, con excepción de {lista(propios)}, "
                     + ("cuyo resultado no cumple" if len(propios) == 1 else "cuyos resultados no cumplen")
                     + " con el límite permisible.")
        else:
            texto = (f"Todos los parámetros con límite establecido en el punto «{rp.punto.nombre}» reportaron "
                     f"resultados dentro de los límites máximos permisibles establecidos en {titulos} de la "
                     "Resolución 0631 de 2015.")
        _vineta(cur, texto, numpr)
    sin_norma = [rp for rp in ctx.con_lab if not rp.punto.evaluar]
    if sin_norma and ctx.aplican:
        _vineta(cur, f"Los resultados de {ctx.los_puntos(sin_norma)} se reportan sin declaración de conformidad, por "
                     "corresponder a la caracterización del agua antes de su vertimiento.", numpr)
    frases = []
    for rp in ctx.con_campo:
        e = _extremos(_horas(rp.campo), [m.caudal_mls for m in rp.campo.mediciones])
        if e:
            (h0, v0), (h1, v1) = e
            frases.append(f"de {num(v0)} mL/s durante toda la jornada en el punto «{rp.punto.nombre}»" if v0 == v1
                          else f"entre {num(v0)} mL/s ({h0} horas) y {num(v1)} mL/s ({h1} horas) en el punto "
                               f"«{rp.punto.nombre}»")
    if frases:
        _vineta(cur, f"Los valores de caudal registraron una variación {lista(frases)}.", numpr)
    for rp in ctx.con_campo:
        if not rp.punto.evaluar:
            continue
        e = _extremos(_horas(rp.campo), _serie(rp.campo, "ph"))
        if e:
            (h0, v0), (h1, v1) = e
            c0, c1 = tx.clase_ph(v0), tx.clase_ph(v1)
            clase = f"un comportamiento {c0}" if c0 == c1 else f"un comportamiento entre {c0} y {c1}"
            _vineta(cur, f"El punto «{rp.punto.nombre}» registró valores de pH que oscilaron entre {num(v0)} unidades "
                         f"de pH ({h0} horas) y {num(v1)} unidades de pH ({h1} horas), evidenciando {clase}.", numpr)


# ------------------------------------------------------------------ principal
def limpiar_plantilla(doc):
    """Deja la plantilla sin datos de un proyecto (idempotente)."""
    _acreditacion(doc)
    for nivel, texto in ((1, "introduccion"), (2, "objetivo general"), (2, "descripcion y localizacion"),
                         (1, "normatividad"), (1, "resultados")):
        t = _titulo(doc, nivel, texto)
        if t is not None:
            _vaciar(t)
    for nivel, texto in ((2, "objetivos especificos"), (1, "conclusiones")):
        t = _titulo(doc, nivel, texto)
        if t is None:
            continue
        # Se conserva una sola vineta vacia: de ella se toma la numeracion.
        primera = True
        el = t.getnext()
        while el is not None and _nivel(el) is None and el.tag != f"{_W}sectPr":
            siguiente = el.getnext()
            if el.tag in (f"{_W}p", f"{_W}tbl"):
                if primera and el.tag == f"{_W}p" and el.find(f"{_W}pPr/{_W}numPr") is not None:
                    set_paragraph_text(Paragraph(el, doc), "")
                    primera = False
                else:
                    _quitar(el)
            el = siguiente
    _quitar_entradas_huerfanas_de_indices(doc)
    quitar_relaciones_huerfanas(doc)


def generar_informe_vertimiento(res: ResultadosVertimiento, ruta_plantilla: str, ruta_salida: str, carpeta: str,
                                datos, hoy: Optional[dt.date] = None) -> tuple[str, list]:
    """Genera el informe Word. Devuelve (ruta, partes a revisar)."""
    doc = docx.Document(ruta_plantilla)
    _acreditacion(doc)
    ctx = Contexto(res, datos, hoy)
    faltantes: list = []
    os.makedirs(carpeta, exist_ok=True)
    if not ctx.con_lab:
        faltantes.append("Resultados del laboratorio: ningún punto tiene reporte")

    foto_rp = next((rp for rp in ctx.evaluados + ctx.puntos
                    if rp.punto.foto_ruta and os.path.exists(rp.punto.foto_ruta)), None)
    _portada_y_encabezado(doc, ctx, foto_rp.punto.foto_ruta if foto_rp else None,
                          foto_rp.punto.nombre if foto_rp else "")
    _cuadro_control(doc, SimpleNamespace(datos=datos, fecha=ctx.fecha))
    _introduccion(doc, ctx, faltantes)
    _objetivos(doc, ctx)
    if ctx.cliente:
        _cliente(doc, SimpleNamespace(cliente=ctx.cliente, fecha=ctx.fecha, datos=ctx.datos_cliente), faltantes)
    else:
        faltantes.append("Información del cliente: el proyecto no tiene cliente")
    _descripcion_puntos(doc, ctx, carpeta, faltantes)
    _tabla_metodos(doc, ctx, faltantes)
    n_lab = _resultados(doc, ctx, carpeta, faltantes)
    _normatividad(doc, ctx, n_lab)
    _conclusiones(doc, ctx)
    _renumerar(doc)
    _anio_fuentes(doc, SimpleNamespace(fecha=ctx.fecha))
    _quitar_entradas_huerfanas_de_indices(doc)
    actualizar_indices_al_abrir(doc)
    quitar_relaciones_huerfanas(doc)
    doc.save(ruta_salida)
    return ruta_salida, faltantes
