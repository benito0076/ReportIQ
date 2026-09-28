"""Utilidades genericas para manipular tablas de un documento Word (.docx)
con python-docx: ubicar tablas por su encabezado, ajustar el numero de filas
de datos y escribir valores preservando el formato de la fila plantilla."""
from __future__ import annotations

import copy
import unicodedata


def _sin_tildes(s: str) -> str:
    if not s:
        return ""
    nfkd = unicodedata.normalize("NFKD", s)
    return "".join(c for c in nfkd if not unicodedata.combining(c)).lower()


def texto_fila(row) -> str:
    return _sin_tildes(" ".join(c.text for c in row.cells))


def encontrar_tabla(doc, requeridos, excluidos=None, filas_encabezado=3, desde=0):
    """Devuelve la primera tabla (a partir del indice `desde`) cuyo encabezado
    (primeras `filas_encabezado` filas) contiene todos los textos de
    `requeridos` y ninguno de `excluidos` (comparacion sin tildes, minusculas)."""
    excluidos = excluidos or []
    req = [_sin_tildes(r) for r in requeridos]
    exc = [_sin_tildes(e) for e in excluidos]
    for i, tabla in enumerate(doc.tables):
        if i < desde:
            continue
        texto = " ".join(texto_fila(r) for r in tabla.rows[:filas_encabezado])
        if all(r in texto for r in req) and not any(e in texto for e in exc):
            return i, tabla
    return None, None


def encontrar_tablas(doc, requeridos, excluidos=None, filas_encabezado=3):
    """Como encontrar_tabla pero devuelve todas las coincidencias en orden."""
    excluidos = excluidos or []
    req = [_sin_tildes(r) for r in requeridos]
    exc = [_sin_tildes(e) for e in excluidos]
    encontradas = []
    for i, tabla in enumerate(doc.tables):
        texto = " ".join(texto_fila(r) for r in tabla.rows[:filas_encabezado])
        if all(r in texto for r in req) and not any(e in texto for e in exc):
            encontradas.append((i, tabla))
    return encontradas


def texto_precedente(doc, tabla, max_parrafos=8) -> str:
    """Texto (sin tildes, minusculas) de los parrafos con contenido
    inmediatamente anteriores a `tabla` en el cuerpo del documento. Sirve
    para distinguir tablas que comparten el mismo encabezado pero estan en
    secciones distintas (p.ej. dos tablas 'Nombre/Codigo/Serial': una de
    equipos de medicion y otra de calibradores)."""
    from docx.text.paragraph import Paragraph

    body = tabla._element.getparent()
    hermanos = list(body)
    try:
        pos = hermanos.index(tabla._element)
    except ValueError:
        return ""

    encontrados = []
    i = pos - 1
    while i >= 0 and len(encontrados) < max_parrafos:
        el = hermanos[i]
        if el.tag.endswith("}p"):
            texto = Paragraph(el, None).text.strip()
            if texto:
                encontrados.append(texto)
        elif el.tag.endswith("}tbl"):
            break  # no cruzar hacia la tabla anterior
        i -= 1
    return _sin_tildes(" ".join(encontrados))


def encontrar_tabla_por_contexto(doc, requeridos, contexto_requerido=None, contexto_excluido=None,
                                  filas_encabezado=1):
    """Como encontrar_tabla, pero ademas exige que el texto de los parrafos
    precedentes contenga/no contenga ciertas palabras clave, para distinguir
    tablas con encabezados identicos ubicadas en secciones distintas."""
    contexto_requerido = [_sin_tildes(c) for c in (contexto_requerido or [])]
    contexto_excluido = [_sin_tildes(c) for c in (contexto_excluido or [])]
    for i, tabla in enumerate(encontrar_tablas(doc, requeridos, filas_encabezado=filas_encabezado)):
        _, t = tabla
        ctx = texto_precedente(doc, t)
        if all(c in ctx for c in contexto_requerido) and not any(c in ctx for c in contexto_excluido):
            return t
    return None


def set_cell_text(cell, texto):
    """Reemplaza el texto de una celda conservando el formato del primer run
    del primer parrafo (fuente, tamano, negrita, etc.)."""
    texto = "" if texto is None else str(texto)
    parrafos = cell.paragraphs
    p0 = parrafos[0]
    # Elimina parrafos adicionales (celdas con saltos de linea previos).
    for p in parrafos[1:]:
        p._element.getparent().remove(p._element)
    if p0.runs:
        p0.runs[0].text = texto
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.add_run(texto)


def ensure_row_count(tabla, filas_encabezado: int, n_filas_datos: int):
    """Asegura que la tabla tenga exactamente `n_filas_datos` filas de datos
    despues de las `filas_encabezado` filas de encabezado, clonando la
    ultima fila de datos existente (o quitando filas sobrantes)."""
    total_actual = len(tabla.rows)
    datos_actual = total_actual - filas_encabezado
    if datos_actual <= 0:
        return

    if n_filas_datos == datos_actual:
        return

    if n_filas_datos < datos_actual:
        for _ in range(datos_actual - n_filas_datos):
            fila = tabla.rows[-1]
            fila._element.getparent().remove(fila._element)
        return

    fila_modelo = tabla.rows[-1]._element
    for _ in range(n_filas_datos - datos_actual):
        nueva = copy.deepcopy(fila_modelo)
        fila_modelo.addnext(nueva)
        fila_modelo = nueva


def escribir_filas(tabla, filas_encabezado: int, filas_valores):
    """Ajusta la tabla al numero de filas de `filas_valores` y escribe cada
    lista de valores en las celdas correspondientes."""
    ensure_row_count(tabla, filas_encabezado, len(filas_valores))
    for i, valores in enumerate(filas_valores):
        fila = tabla.rows[filas_encabezado + i]
        for cell, valor in zip(fila.cells, valores):
            set_cell_text(cell, valor)


def ajustar_bloques_de_filas(tabla, filas_por_bloque: int, n_bloques: int):
    """Como ensure_row_count pero clonando/quitando BLOQUES completos de
    `filas_por_bloque` filas contiguas (preservando las combinaciones de
    celdas -vMerge/gridSpan- dentro de cada bloque, ya que se copian como
    unidad). Util para tablas de 'tarjetas' repetidas (una tarjeta = un
    bloque de varias filas), a diferencia de escribir_filas que clona una
    sola fila."""
    total_filas = len(tabla.rows)
    bloques_actuales = total_filas // filas_por_bloque
    if bloques_actuales == 0:
        return

    if n_bloques < bloques_actuales:
        for _ in range(bloques_actuales - n_bloques):
            for _ in range(filas_por_bloque):
                fila = tabla.rows[-1]
                fila._element.getparent().remove(fila._element)
        return

    if n_bloques > bloques_actuales:
        filas_modelo = [tabla.rows[i]._element for i in range(filas_por_bloque)]
        ancla = tabla.rows[-1]._element
        for _ in range(n_bloques - bloques_actuales):
            for fila_modelo in filas_modelo:
                nueva = copy.deepcopy(fila_modelo)
                ancla.addnext(nueva)
                ancla = nueva


def set_cell_image(cell, ruta_imagen: str, ancho_emu=None):
    """Reemplaza todo el contenido de una celda por una sola imagen."""
    from docx.shared import Emu

    parrafos = cell.paragraphs
    for p in parrafos[1:]:
        p._element.getparent().remove(p._element)
    p0 = parrafos[0]
    for r in list(p0.runs):
        r._element.getparent().remove(r._element)
    run = p0.add_run()
    if ancho_emu:
        run.add_picture(ruta_imagen, width=Emu(ancho_emu))
    else:
        run.add_picture(ruta_imagen)


def eliminar_tabla(tabla):
    tabla._element.getparent().remove(tabla._element)


def insertar_texto_despues_de(doc, texto_heading, texto_nuevo, negrita=False):
    """Inserta un parrafo de texto justo despues del primer heading cuyo
    texto contenga `texto_heading`. Util para agregar una leyenda/titulo
    antes de una imagen insertada con insertar_imagen_despues_de."""
    objetivo = _sin_tildes(texto_heading)
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and objetivo in _sin_tildes(p.text):
            nuevo_p = p.insert_paragraph_before("")
            p._element.addnext(nuevo_p._element)
            run = nuevo_p.add_run(texto_nuevo)
            run.bold = negrita
            return True
    return False


def insertar_imagen_despues_de(doc, texto_heading, ruta_imagen, ancho_emu=None):
    """Inserta una imagen en un nuevo parrafo justo despues del primer
    heading cuyo texto contenga `texto_heading` (sin tildes, insensible a
    mayusculas). Si no se encuentra el heading, no hace nada."""
    from docx.shared import Emu

    objetivo = _sin_tildes(texto_heading)
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and objetivo in _sin_tildes(p.text):
            nuevo_p = p.insert_paragraph_before("")
            # mover nuevo_p para que quede DESPUES del heading, no antes
            p._element.addnext(nuevo_p._element)
            run = nuevo_p.add_run()
            if ancho_emu:
                run.add_picture(ruta_imagen, width=Emu(ancho_emu))
            else:
                run.add_picture(ruta_imagen)
            return True
    return False


def set_paragraph_text(parrafo, texto):
    """Reemplaza el texto de un parrafo conservando el formato de su primer run."""
    runs = parrafo.runs
    if runs:
        runs[0].text = texto
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        parrafo.add_run(texto)


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def reemplazar_imagen(doc, blip, ruta_imagen):
    """Cambia la imagen de un <a:blip> existente por `ruta_imagen`,
    conservando el ancho y la posicion del dibujo y ajustando el alto a la
    proporcion de la nueva imagen."""
    rid, imagen = doc.part.get_or_add_image(ruta_imagen)
    blip.set(f"{_R}embed", rid)
    ancho_px, alto_px = imagen.px_width, imagen.px_height
    if not ancho_px or not alto_px:
        return
    dibujo = blip
    while dibujo is not None and dibujo.tag not in (f"{_WP}inline", f"{_WP}anchor"):
        dibujo = dibujo.getparent()
    if dibujo is None:
        return
    extent = dibujo.find(f"{_WP}extent")
    if extent is None:
        return
    cx = int(extent.get("cx"))
    cy = int(cx * alto_px / ancho_px)
    extent.set("cy", str(cy))
    for ext in dibujo.iter(f"{_A}ext"):
        if ext.get("cx") is not None and ext.getparent().tag == f"{_A}xfrm":
            ext.set("cx", str(cx))
            ext.set("cy", str(cy))


def blips(elemento):
    return list(elemento.iter(f"{_A}blip"))
