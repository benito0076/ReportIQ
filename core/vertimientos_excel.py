"""Excel de resultados de vertimientos: comparacion con la Res. 0631, datos de campo y norma aplicada."""
from __future__ import annotations

from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import res0631, res0631_datos
from .vertimientos import PuntoCampo, ResultadosVertimiento

_ENC = PatternFill("solid", fgColor="BAD405")
_ROJO = PatternFill("solid", fgColor="F8CBAD")
_VERDE = PatternFill("solid", fgColor="E2EFDA")
_BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _encabezado(ws, fila: int, textos: list[str], anchos: Optional[list[int]] = None):
    for c, t in enumerate(textos, start=1):
        celda = ws.cell(fila, c, t)
        celda.fill, celda.font, celda.alignment, celda.border = _ENC, Font(bold=True), _CENTRO, _BORDE
    for c, w in enumerate(anchos or [], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[fila].height = 45
    ws.freeze_panes = ws.cell(fila + 1, 2)


def _celda(ws, fila: int, col: int, valor, fill=None, izquierda=False):
    c = ws.cell(fila, col, valor)
    c.border = _BORDE
    c.alignment = Alignment(horizontal="left" if izquierda else "center", vertical="center", wrap_text=True)
    if fill:
        c.fill = fill
    return c


def _hoja_resultados(wb: Workbook, res: ResultadosVertimiento):
    ws = wb.active
    ws.title = "Resultados vs Res. 0631"
    puntos = [rp for rp in res.puntos if rp.informe]
    evaluados = [(rp.punto.nombre, i, c) for rp in puntos if rp.punto.evaluar
                 for i, c in enumerate(res.columnas) if c.aplica]
    titulos = (["Parámetro", "Unidades", "Método", "LCM", "Incertidumbre (µc), %"]
               + [f"{rp.punto.nombre}\n{rp.informe.muestra}" for rp in puntos]
               + [c.titulo for c in res.columnas]
               + [f"Conformidad {p}\n({c.titulo})" for p, _, c in evaluados])
    _encabezado(ws, 1, titulos, [38, 14, 24, 10, 12] + [16] * len(puntos) + [22] * len(res.columnas)
                + [20] * len(evaluados))
    for r, f in enumerate(res.filas, start=2):
        valores = [f.parametro.nombre + (" (*)" if f.subcontratado else ""), f.unidad, f.metodo, f.lcm,
                   f.incertidumbre]
        for c, v in enumerate(valores, start=1):
            _celda(ws, r, c, v, izquierda=c in (1, 3))
        col = len(valores) + 1
        for rp in puntos:
            rl = f.resultados.get(rp.punto.nombre)
            _celda(ws, r, col, rl.reporte if rl else "—")
            col += 1
        for lim in f.limites:
            _celda(ws, r, col, res0631.texto_limite(f.parametro.clave, lim) + lim.marca)
            col += 1
        for punto, i, _ in evaluados:
            estado = f.conformidad.get((punto, i))
            _celda(ws, r, col, estado or "—", _ROJO if estado == "No cumple" else _VERDE if estado else None)
            col += 1
    fila = len(res.filas) + 3
    notas = ["(*) Ensayo subcontratado.", "N.E.: No establecido por la norma."]
    if res.proyecto.alcantarillado:
        notas += ["* Art. 16: exigencias de la actividad para cuerpos de agua superficial multiplicadas por 1,50.",
                  "** Art. 16: mismas exigencias de la actividad para cuerpos de agua superficial."]
    notas.append("Temperatura: valor límite máximo de 40,00 °C (Artículo 5 de la Resolución 0631 de 2015).")
    for k, n in enumerate(notas):
        ws.cell(fila + k, 1, n).font = Font(italic=True, size=9)


def _hoja_campo(wb: Workbook, nombre: str, campo: PuntoCampo):
    ws = wb.create_sheet(f"Campo {campo.hoja}"[:31])
    ws.cell(1, 1, f"{campo.titulo} (hoja «{campo.hoja}» de la FP-004)").font = Font(bold=True)
    titulos = ["Hora", "Volumen (mL)", "Tiempo (s)", "Caudal (mL/s)", "Fracción", "Volumen alícuota (mL)",
               "pH", "pH dup.", "Temperatura (°C)", "Temperatura dup.", "OD (mg O2/L)", "OD dup.",
               "Conductividad (µS/cm)", "Conductividad dup.", "Sól. sedimentables (mL/L-h)", "Sól. sed. dup."]
    _encabezado(ws, 3, titulos, [9, 11, 9, 11, 9, 12] + [11] * 10)
    for r, m in enumerate(campo.mediciones, start=4):
        fr, al = campo.fraccion(m), campo.alicuota_ml(m)
        fila = [m.hora.strftime("%H:%M") if m.hora else "", m.volumen_ml, m.tiempo_s,
                round(m.caudal_mls, 2) if m.caudal_mls is not None else None,
                round(fr, 4) if fr is not None else None, round(al, 1) if al is not None else None,
                m.ph, m.ph_dup, m.temperatura, m.temperatura_dup, m.oxigeno, m.oxigeno_dup, m.conductividad,
                m.conductividad_dup, m.sedimentables, m.sedimentables_dup]
        for c, v in enumerate(fila, start=1):
            _celda(ws, r, c, v)
    r = 4 + len(campo.mediciones)
    caudales = [m.caudal_mls for m in campo.mediciones if m.caudal_mls is not None]
    resumen = [("Suma de caudales (mL/s)", sum(caudales) if caudales else None),
               ("Caudal promedio (mL/s)", sum(caudales) / len(caudales) if caudales else None),
               ("Tamaño de la muestra compuesta (mL)", campo.tamano_muestra_ml)]
    for k, (t, v) in enumerate(resumen):
        ws.cell(r + 1 + k, 1, t).font = Font(bold=True)
        ws.cell(r + 1 + k, 4, round(v, 2) if isinstance(v, float) else v)


def _hoja_norma(wb: Workbook, res: ResultadosVertimiento):
    if not res.columnas:
        return
    ws = wb.create_sheet("Norma aplicada")
    _encabezado(ws, 1, ["Parámetro (Res. 0631 de 2015)"] + [c.titulo for c in res.columnas],
                [36] + [24] * len(res.columnas))
    nombres = _nombres_norma()
    for r, p in enumerate(res0631_datos.ORDEN, start=2):
        _celda(ws, r, 1, nombres.get(p, p), izquierda=True)
        for c, col in enumerate(res.columnas, start=2):
            lim = res0631.limite(col.actividad, p, col.alcantarillado, res.proyecto.consumo_humano)
            _celda(ws, r, c, lim.texto + lim.marca)


def _nombres_norma() -> dict:
    from .vertimientos import PARAMETROS

    n = {p.clave: p.nombre for p in PARAMETROS}
    n.update({"ph": "pH", "solidos_sedimentables": "Sólidos Sedimentables (SSED)",
              "compuestos_fenolicos": "Compuestos Semivolátiles Fenólicos", "fenoles": "Fenoles Totales",
              "hidrocarburos_totales": "Hidrocarburos Totales (HTP)", "hap": "Hidrocarburos Aromáticos Policíclicos (HAP)",
              "btex": "BTEX (Benceno, Tolueno, Etilbenceno y Xileno)",
              "aox": "Compuestos Orgánicos Halogenados Adsorbibles (AOX)", "fluoruros": "Fluoruros",
              "antimonio": "Antimonio", "titanio": "Titanio", "aceites_grasas": "Grasas y Aceites",
              "color": "Color Real (436, 525 y 620 nm)"})
    return n


def exportar_resultados_vertimiento(res: ResultadosVertimiento, ruta: str) -> str:
    wb = Workbook()
    _hoja_resultados(wb, res)
    for rp in res.puntos:
        if rp.campo:
            _hoja_campo(wb, rp.punto.nombre, rp.campo)
    _hoja_norma(wb, res)
    if res.advertencias:
        ws = wb.create_sheet("Advertencias")
        for r, a in enumerate(res.advertencias, start=1):
            ws.cell(r, 1, a)
        ws.column_dimensions["A"].width = 120
    wb.save(ruta)
    return ruta
