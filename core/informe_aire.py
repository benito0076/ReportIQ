"""Informe Word de calidad del aire (Res. 2254 de 2017).

Plantilla: templates/informe_aire_template.docx (formato FP-024, a partir del
informe EC-041/042-26). Se conservan las partes fijas (portada, cuadro de
control, marco normativo, procedimientos, laboratorio) y se regeneran por
completo los capitulos que dependen de los datos: meteorologia, niveles de
inmision por contaminante, analisis estadistico, ICA y conclusiones, con
tantas estaciones y contaminantes como tenga el proyecto.
"""
from __future__ import annotations

import copy
import datetime as dt
import os
import re
from types import SimpleNamespace
from typing import Optional

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from . import aire_graficas as gr
from .aire import (
    CATEGORIAS_ICA, CO, COMPUESTOS_COV, COV, LIMITES, NO2, O3, PM10, PM25, SO2, ResultadosAire, redondear,
)
from .docx_utils import _sin_tildes, blips, encontrar_tabla, reemplazar_imagen, set_cell_text, set_paragraph_text
from .informe_textos import _anio_fuentes, _cliente, _cuadro_control, actualizar_indices_al_abrir, cantidad, lista

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]
EMPRESA = "Ambienciq Ingenieros S.A.S."
ENCABEZADO = "BAD405"
ANCHO_GRAFICA = Cm(15.5)

# Nombre en el texto, sigla, palabras para reconocer filas de las tablas de la plantilla.
INFO = {
    PM10: ("Material Particulado menor a 10 micras (PM10)", "PM10", ("10 micras", "pm10", "alto volumen")),
    PM25: ("Material Particulado menor a 2.5 micras (PM2.5)", "PM2.5", ("2.5 micras", "2,5 micras", "pm2,5", "pm2.5",
                                                                        "bajo volumen")),
    SO2: ("Dióxido de Azufre (SO2)", "SO2", ("azufre", "so2")),
    NO2: ("Dióxido de Nitrógeno (NO2)", "NO2", ("nitrogeno", "no2", "nox")),
    CO: ("Monóxido de Carbono (CO)", "CO", ("monoxido", "(co)", " co")),
    O3: ("Ozono (O3)", "O3", ("ozono", "o3")),
    COV: ("Compuestos Orgánicos Volátiles (COV)", "COV", ("volatiles", "cov", "voc")),
}
TITULOS = {
    PM10: "Partículas Menores a 10 Micras (PM10)",
    PM25: "Partículas Menores a 2,5 Micras (PM2,5)",
    SO2: "Dióxido de Azufre (SO2)",
    NO2: "Dióxido de Nitrógeno (NO2)",
    CO: "Monóxido de Carbono (CO)",
    O3: "Ozono (O3)",
    COV: "Compuestos Orgánicos Volátiles (COV)",
}


# ------------------------------------------------------------------ formato
def fmt(v: Optional[float], dec: int = 2, miles: bool = False) -> str:
    if v is None:
        return "---"
    r = redondear(v, dec)
    texto = f"{r:,.{dec}f}" if miles else f"{r:.{dec}f}"
    return texto.replace(",", "_").replace(".", ",").replace("_", ".")


def conc(v, bajo_lc=False, dec=2, miles=False) -> str:
    return ("<" if bajo_lc and v is not None else "") + fmt(v, dec, miles)


def fecha_larga(d) -> str:
    if d is None:
        return ""
    d = d.date() if isinstance(d, dt.datetime) else d
    return f"{d.day:02d} de {MESES[d.month - 1]} de {d.year}"


def _rango_fechas(d0, d1) -> str:
    if d0 == d1:
        return f"el {fecha_larga(d0)}"
    return f"entre el {fecha_larga(d0)} y el {fecha_larga(d1)}"


# ------------------------------------------------------------- utilidades xml
def _estilo(el) -> str:
    ppr = el.find(f"{_W}pPr")
    st = ppr.find(f"{_W}pStyle") if ppr is not None else None
    return st.get(qn("w:val")) if st is not None else ""


def _texto(el) -> str:
    return "".join(t.text or "" for t in el.iter(f"{_W}t"))


def _cuerpo(doc):
    return [el for el in doc.element.body.iterchildren() if el.tag in (f"{_W}p", f"{_W}tbl")]


def _es_titulo(el, nivel: int, texto: Optional[str] = None) -> bool:
    if el.tag != f"{_W}p" or _estilo(el) != f"Ttulo{nivel}":
        return False
    return texto is None or texto in _sin_tildes(_texto(el))


def _es_leyenda(el) -> bool:
    return el.tag == f"{_W}p" and _estilo(el) == "Descripcin"


def _tiene_imagen(el) -> bool:
    return bool(blips(el)) or next(el.iter(f"{_W}object"), None) is not None


def _quitar(el):
    padre = el.getparent()
    if padre is not None:
        # Un salto de seccion dentro del parrafo se conserva en el siguiente parrafo vacio.
        sect = el.find(f"{_W}pPr/{_W}sectPr") if el.tag == f"{_W}p" else None
        if sect is not None:
            p = OxmlElement("w:p")
            ppr = OxmlElement("w:pPr")
            ppr.append(sect)
            p.append(ppr)
            el.addnext(p)
        padre.remove(el)


def _es_accesorio(el) -> bool:
    """Tablas, imagenes, lineas vacias y 'Fuente'/'Nota' que acompanan una leyenda."""
    if el.tag == f"{_W}tbl":
        return True
    if el.tag != f"{_W}p" or _estilo(el).startswith("Ttulo") or _es_leyenda(el) and not _tiene_imagen(el):
        return False
    t = _sin_tildes(_texto(el)).strip()
    return not t or _tiene_imagen(el) or t.startswith(("fuente", "nota"))


def _quitar_bloque(leyenda):
    """Quita una leyenda con su tabla/imagen y sus lineas de fuente y notas."""
    siguiente = leyenda.getnext()
    _quitar(leyenda)
    while siguiente is not None and _es_accesorio(siguiente):
        actual, siguiente = siguiente, siguiente.getnext()
        _quitar(actual)


def _sin_bloque(texto: str) -> str:
    return re.sub(r"\s*[–-]\s*(Bloque|Campo)?\s*Lisama\s*$", "", texto.strip())


# -------------------------------------------------- limpieza de la plantilla
def limpiar_plantilla(doc):
    """Deja la plantilla lista para un solo proyecto: sin los bloques del
    segundo campo (Llanito) y sin el contenido de los capitulos que se
    regeneran. Es idempotente (la plantilla incluida ya viene limpia)."""
    for el in _cuerpo(doc):
        if el.getparent() is None:
            continue
        if _es_leyenda(el) and "llanito" in _sin_tildes(_texto(el)):
            _quitar_bloque(el)
        elif _es_titulo(el, 3, "bloque llanito"):
            _quitar(el)

    cuerpo = _cuerpo(doc)
    i_meteo = next((i for i, el in enumerate(cuerpo) if _es_titulo(el, 1, "meteorolog")), None)
    i_concl = next((i for i, el in enumerate(cuerpo) if _es_titulo(el, 1, "conclusiones")), None)
    i_anexos = next((i for i, el in enumerate(cuerpo) if _es_titulo(el, 1, "anexos")), None)
    if i_meteo is not None and i_concl is not None:
        conservar = set()
        for k in range(i_meteo + 1, i_concl):
            el = cuerpo[k]
            if _es_titulo(el, 1, "resultados") or _es_titulo(el, 2, "niveles de inmision") \
                    or _es_titulo(el, 2, "analisis estadistico") or _es_titulo(el, 2, "indice de calidad"):
                conservar.add(k)
            elif _es_leyenda(el) and "beaufort" in _sin_tildes(_texto(el)):
                conservar.add(k)
                j = k + 1
                while j < i_concl and _es_accesorio(cuerpo[j]):
                    conservar.add(j)
                    j += 1
        for k in range(i_meteo + 1, i_concl):
            if k not in conservar:
                _quitar(cuerpo[k])
    if i_concl is not None and i_anexos is not None:
        vinetas = [el for el in cuerpo[i_concl + 1:i_anexos] if el.getparent() is not None
                   and el.tag == f"{_W}p" and el.find(f"{_W}pPr/{_W}numPr") is not None]
        for el in vinetas[1:]:
            _quitar(el)
    # Leyendas (y sus entradas en los indices) sin el sufijo del bloque.
    for el in _cuerpo(doc):
        texto = _sin_tildes(_texto(el))
        es_indice = el.tag == f"{_W}p" and _estilo(el).lower().startswith(("tabladeilustraciones", "tdc", "toc"))
        if es_indice and "llanito" in texto:
            _quitar(el)
        elif (_es_leyenda(el) or es_indice) and "lisama" in texto:
            _quitar_sufijo_bloque(el)
    _quitar_entradas_huerfanas_de_indices(doc)
    quitar_relaciones_huerfanas(doc)


def _quitar_entradas_huerfanas_de_indices(doc):
    """Entradas de la tabla de contenido y de los indices de tablas, graficas y
    figuras que apuntan a titulos eliminados (Word los regenera al abrir)."""
    def clave(texto: str) -> str:
        t = _sin_tildes(texto)
        t = re.sub(r"^\s*(tabla|grafica|figura|ecuacion)\s*\d+\s*\.?", "", t)
        t = re.sub(r"^\s*[\d.]+", "", t)
        t = re.sub(r"\d+\s*$", "", t)
        return re.sub(r"[^a-z0-9]+", "", t)

    marcadores = {b.get(qn("w:name")) for b in doc.element.body.iter(f"{_W}bookmarkStart")}
    vigentes = {clave(_texto(el)) for el in _cuerpo(doc)
                if el.tag == f"{_W}p" and (_es_leyenda(el) or _estilo(el).startswith("Ttulo"))}
    for p in list(doc.element.body.iter(f"{_W}p")):
        if not _estilo(p).lower().startswith(("tabladeilustraciones", "tdc", "toc")):
            continue
        anclas = [h.get(qn("w:anchor")) for h in p.iter(f"{_W}hyperlink") if h.get(qn("w:anchor"))]
        huerfana = bool(anclas) and all(a not in marcadores for a in anclas)
        if huerfana or (_texto(p).strip() and clave(_texto(p)) not in vigentes):
            _quitar(p)


def _quitar_sufijo_bloque(el):
    """'... – Bloque Lisama' aunque Word lo tenga partido en varios fragmentos."""
    ts = [t for t in el.iter(f"{_W}t")]
    completo = "".join(t.text or "" for t in ts)
    m = re.search(r"\s*[–-]\s*(Bloque|Campo)?\s*Lisama", completo)
    if not m:
        return
    pos = 0
    for t in ts:
        texto = t.text or ""
        ini, fin = pos, pos + len(texto)
        pos = fin
        a, b = max(m.start(), ini), min(m.end(), fin)
        if a < b:
            t.text = texto[:a - ini] + texto[b - ini:]


def quitar_relaciones_huerfanas(doc):
    """Quita del paquete las imagenes/objetos que ya no se usan (la plantilla
    trae graficas EMF pesadas del informe original)."""
    usados = set()
    for el in doc.element.iter():
        for k, v in el.attrib.items():
            if k.startswith(_R):
                usados.add(v)
    part = doc.part
    for rid, rel in list(part.rels.items()):
        if rel.is_external or rid in usados:
            continue
        if any(x in rel.reltype for x in ("/image", "/oleObject", "/chart", "/package")):
            part.rels.pop(rid)


# ------------------------------------------------------------- construccion
class Cursor:
    """Inserta parrafos, tablas e imagenes uno tras otro despues de un elemento."""

    def __init__(self, doc, despues_de):
        self.doc = doc
        self.ultimo = despues_de

    def _poner(self, el):
        self.ultimo.addnext(el)
        self.ultimo = el

    def numero(self, tipo: str) -> int:
        """Numero que tendra la siguiente leyenda 'Tabla' / 'Gráfica' / 'Figura'."""
        n = 0
        for el in self.doc.element.body.iterchildren():
            if _es_leyenda(el) and _texto(el).strip().startswith(tipo):
                n += 1
            if el is self.ultimo:
                break
        return n + 1

    def parrafo(self, texto: str = "", estilo: str = "Normal", negrita: bool = False, alinear=None) -> Paragraph:
        p = self.doc.add_paragraph(style=estilo)
        if texto:
            run = p.add_run(texto)
            run.bold = negrita or None
        if alinear is not None:
            p.alignment = alinear
        self._poner(p._p)
        return p

    def titulo(self, texto: str, nivel: int) -> Paragraph:
        return self.parrafo(texto, f"Heading {nivel}")

    def subtitulo(self, texto: str) -> Paragraph:
        return self.parrafo(texto, "TITULO 3")

    def leyenda(self, tipo: str, texto: str) -> int:
        n = self.numero(tipo)
        p = self.doc.add_paragraph(style="Caption")
        p.add_run(f"{tipo} ")
        _campo(p, f" SEQ {tipo} \\* ARABIC ", str(n))
        p.add_run(f". {texto}")
        self._poner(p._p)
        return n

    def fuente(self, texto: Optional[str] = None, anio: Optional[int] = None):
        self.parrafo(texto or f"Fuente: {EMPRESA}, {anio or dt.date.today().year}.", "Fuente")

    def nota(self, texto: str):
        self.parrafo(texto, "Fuente")

    def imagen(self, ruta: Optional[str], ancho=ANCHO_GRAFICA):
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if ruta and os.path.exists(ruta):
            p.add_run().add_picture(ruta, width=ancho)
        self._poner(p._p)

    def tabla(self, encabezado: list, filas: list, rellenos: Optional[dict] = None, anchos: Optional[list] = None,
              izquierda: tuple = ()) -> Table:
        t = construir_tabla(self.doc, encabezado, filas, rellenos, anchos, izquierda)
        self._poner(t._tbl)
        return t


def _campo(p: Paragraph, instruccion: str, resultado: str):
    def run(*hijos):
        r = OxmlElement("w:r")
        for h in hijos:
            r.append(h)
        p._p.append(r)

    def fld(tipo):
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), tipo)
        return f

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruccion
    texto = OxmlElement("w:t")
    texto.text = resultado
    run(fld("begin"))
    run(instr)
    run(fld("separate"))
    run(texto)
    run(fld("end"))


def _formato_celda(celda, encabezado: bool, relleno: Optional[str], izquierda: bool):
    tcpr = celda._tc.get_or_add_tcPr()
    bordes = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "auto")
        bordes.append(b)
    tcpr.append(bordes)
    color = ENCABEZADO if encabezado else relleno
    if color:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcpr.append(shd)
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcpr.append(va)
    for p in celda.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if izquierda and not encabezado else WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        for r in p.runs:
            r.font.size = Pt(9)
            r.bold = encabezado or None


def construir_tabla(doc, encabezado: list, filas: list, rellenos: Optional[dict] = None,
                    anchos: Optional[list] = None, izquierda: tuple = ()) -> Table:
    """`encabezado`: filas de celdas (texto, columnas, filas_que_ocupa); `filas`:
    listas de textos. `rellenos`: {(fila, columna): color} en las filas de datos."""
    ncol = sum(c[1] if isinstance(c, tuple) else 1 for c in encabezado[0]) if encabezado else len(filas[0])
    t = doc.add_table(rows=len(encabezado) + len(filas), cols=ncol)
    tblpr = t._tbl.tblPr
    ancho = OxmlElement("w:tblW")
    ancho.set(qn("w:w"), "5000")
    ancho.set(qn("w:type"), "pct")
    for viejo in tblpr.findall(qn("w:tblW")):
        tblpr.remove(viejo)
    tblpr.append(ancho)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblpr.append(jc)
    ocupadas = set()
    combinar = []
    for r, fila in enumerate(encabezado):
        c = 0
        for celda in fila:
            texto, ncols, nfilas = (celda + (1, 1)[len(celda) - 1:]) if isinstance(celda, tuple) else (celda, 1, 1)
            while (r, c) in ocupadas:
                c += 1
            t.cell(r, c).text = texto
            for dr in range(nfilas):
                for dc in range(ncols):
                    ocupadas.add((r + dr, c + dc))
            if ncols > 1 or nfilas > 1:
                combinar.append((r, c, r + nfilas - 1, c + ncols - 1))
            c += ncols
    n_enc = len(encabezado)
    for r, fila in enumerate(filas):
        for c, valor in enumerate(fila):
            t.cell(n_enc + r, c).text = "" if valor is None else str(valor)
    for r in range(len(encabezado) + len(filas)):
        for c in range(ncol):
            celda = t.cell(r, c)
            relleno = (rellenos or {}).get((r - n_enc, c)) if r >= n_enc else None
            _formato_celda(celda, r < n_enc, relleno, c in izquierda)
    for r0, c0, r1, c1 in combinar:
        a = t.cell(r0, c0)
        texto = a.text
        a.merge(t.cell(r1, c1))
        set_cell_text(a, texto)
        for p in a.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9)
                run.bold = True
    if anchos:
        for fila in t.rows:
            for celda, w in zip(fila.cells, anchos):
                celda.width = w
    # Las filas de encabezado se repiten en cada pagina.
    for r in range(n_enc):
        trpr = t.rows[r]._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        trpr.append(rep)
    return t


# ---------------------------------------------------------------- contexto
class ContextoAire:
    def __init__(self, res: ResultadosAire, datos, hoy: Optional[dt.date] = None):
        self.res = res
        self.datos = datos
        self.proyecto = res.proyecto
        self.cliente = (res.proyecto.cliente or "").strip()
        self.estaciones = res.estaciones()
        self.nombres = {e.numero: (e.nombre or f"Estación {e.numero}").strip() for e in self.estaciones}
        self.contaminantes = res.contaminantes()
        fecha = None
        if getattr(datos, "fecha", ""):
            try:
                fecha = dt.date.fromisoformat(datos.fecha)
            except ValueError:
                fecha = None
        self.fecha = fecha or hoy or dt.date.today()
        self.periodos = {e.numero: self._periodo(e.numero) for e in self.estaciones}
        dias = [d for p in self.periodos.values() if p for d in p]
        self.inicio = min(dias) if dias else None
        self.fin = max(dias) if dias else None
        self.anio = self.fecha.year

    def _periodo(self, n):
        fechas = []
        for (c, k), s in self.res.manuales.items():
            if k == n:
                fechas += [m.inicio.date() for m in s.muestras if m.inicio] + [m.fecha for m in s.muestras if m.fecha]
        for (c, k), s in self.res.automaticos.items():
            if k == n and s.horas:
                fechas += [s.horas[0][0].date(), s.horas[-1][0].date()]
        for (c, k), s in self.res.cov.items():
            if k == n:
                fechas += [m.fecha for m in s.muestras if m.fecha]
        return (min(fechas), max(fechas)) if fechas else None

    @property
    def area(self) -> str:
        area = (getattr(self.datos, "area_estudio", "") or "").strip().rstrip(".")
        return area or f"el área del proyecto {self.proyecto.nombre_proyecto}".strip()

    def nombres_contaminantes(self, solo=None) -> str:
        return lista([INFO[c][0] for c in self.contaminantes if solo is None or c in solo])

    def las_estaciones(self) -> str:
        n = len(self.estaciones)
        return "la estación de monitoreo" if n == 1 else f"las {cantidad(n, 'estación', 'estaciones')} de monitoreo"


# ----------------------------------------------------------- partes fijas
def _portada_y_encabezado(doc, ctx: ContextoAire):
    d = ctx.datos
    titulo = [l.strip() for l in (d.titulo or "").splitlines() if l.strip()]
    if not titulo:
        titulo = [t for t in (ctx.proyecto.nombre_proyecto.strip(), ctx.cliente) if t]
    titulo = [t.upper() for t in titulo]
    if ctx.inicio:
        meses = [MESES[ctx.inicio.month - 1].upper()]
        if (ctx.fin.year, ctx.fin.month) != (ctx.inicio.year, ctx.inicio.month):
            meses.append(MESES[ctx.fin.month - 1].upper())
        periodo = f"{' – '.join(meses)} DE {ctx.fin.year}"
    else:
        periodo = f"{MESES[ctx.fecha.month - 1].upper()} DE {ctx.fecha.year}"
    for txbx in doc.element.body.iter(f"{_W}txbxContent"):
        ps = [Paragraph(p, None) for p in txbx.iter(f"{_W}p")]
        con_texto = [p for p in ps if p.text.strip()]
        texto = _sin_tildes(" ".join(p.text for p in con_texto)).strip()
        if not texto or "informe tecnico" in texto:
            continue
        if re.fullmatch(r"[a-z – -]+ de 20\d\d", texto) and any(m in texto for m in MESES):
            set_paragraph_text(con_texto[0], periodo)
        elif len(texto) > 20 and titulo:
            for p, valor in zip(con_texto, titulo + [""] * len(con_texto)):
                set_paragraph_text(p, valor)
            for p in con_texto[len(titulo):]:
                p._p.getparent().remove(p._p)

    fecha_txt = ctx.fecha.strftime("%d/%m/%Y")
    version = (d.version or "1.0").strip()
    acto = [l.strip() for l in (getattr(d, "acto_administrativo", "") or "").splitlines() if l.strip()]
    vistos = set()
    for seccion in doc.sections:
        for hdr in (seccion.header, seccion.first_page_header, seccion.even_page_header):
            if hdr.is_linked_to_previous or id(hdr._element) in vistos:
                continue
            vistos.add(id(hdr._element))
            for tabla in hdr.tables:
                for fila in tabla.rows:
                    for celda in fila.cells:
                        ps = celda.paragraphs
                        t0 = _sin_tildes(ps[0].text.strip()) if ps else ""
                        if t0.startswith("informe tecnico de calidad del aire") and len(ps) > 1:
                            set_paragraph_text(ps[1], " - ".join(titulo))
                        elif t0 == "expediente" and len(ps) > 1:
                            set_paragraph_text(ps[1], d.expediente.strip() or "No aplica")
                        elif t0 == "acto administrativo" and len(ps) > 1:
                            valores = acto or ["No aplica"]
                            for p, v in zip(ps[1:], valores + [None] * len(ps)):
                                if v is None:
                                    p._p.getparent().remove(p._p)
                                else:
                                    set_paragraph_text(p, v)
                        elif t0 == "elaborado" and len(ps) > 1:
                            set_paragraph_text(ps[1], fecha_txt)
                        elif t0.startswith("version:"):
                            set_paragraph_text(ps[0], f"Versión:  {version}")


def _parrafo(doc, *fragmentos, desde=None):
    """Primer parrafo del cuerpo que contiene todos los fragmentos (sin tildes)."""
    activo = desde is None
    for el in _cuerpo(doc):
        if not activo:
            activo = el is desde
            continue
        if el.tag == f"{_W}p" and all(f in _sin_tildes(_texto(el)) for f in fragmentos):
            return Paragraph(el, doc)
    return None


# Extension aproximada de Colombia (incluye San Andres y el mar territorial).
_LAT_COLOMBIA = (-5.0, 14.0)
_LON_COLOMBIA = (-83.0, -66.0)


def _coordenadas(e) -> dict:
    """Longitud/latitud (DMS) y Norte/Este (Origen Nacional) de una estacion.

    Las coordenadas planas digitadas se conservan tal cual (sin ida y vuelta por
    la proyeccion). Si el punto cae fuera de Colombia, `fuera` es True y no se
    calculan las otras: la conversion no es valida lejos del origen."""
    from .geo_colombia import a_geografica, es_origen_nacional, formatear_dms, geografica_a_origen_nacional

    salida = {"lon": e.longitud or "---", "lat": e.latitud or "---", "norte": "---", "este": "---", "dd": None,
              "fuera": False}
    x, y = parsear_coordenada(e.longitud), parsear_coordenada(e.latitud)
    if x is None or y is None:
        return salida
    plana = es_origen_nacional(x, y)
    if plana:
        salida.update(este=fmt(x, 3, True), norte=fmt(y, 3, True), lon="---", lat="---")
    lat, lon = a_geografica(x, y)
    if not (_LAT_COLOMBIA[0] <= lat <= _LAT_COLOMBIA[1] and _LON_COLOMBIA[0] <= lon <= _LON_COLOMBIA[1]):
        salida["fuera"] = True
        return salida
    if not plana:
        x, y = geografica_a_origen_nacional(lat, lon)
        salida.update(este=fmt(x, 3, True), norte=fmt(y, 3, True))
    salida.update(lon=formatear_dms(lon, True), lat=formatear_dms(lat, False), dd=(lat, lon))
    return salida


def _avisar_coordenadas(ctx, faltantes):
    fuera = [ctx.nombres[e.numero] for e in ctx.estaciones if _coordenadas(e)["fuera"]]
    if fuera:
        faltantes.append("Coordenadas fuera de Colombia (revise Este/Norte en Origen Nacional, p. ej. Este "
                         "4.000.000–5.700.000 y Norte 1.000.000–3.100.000): " + ", ".join(fuera))


def parsear_coordenada(texto) -> Optional[float]:
    """Grados decimales, metros (Origen Nacional) o DMS como 73°33'40.40"O."""
    if texto is None:
        return None
    t = str(texto).strip()
    if not t:
        return None
    m = re.match(r"^\s*(-?\d+(?:[.,]\d+)?)\s*°\s*(\d+(?:[.,]\d+)?)?\s*['’′]?\s*(\d+(?:[.,]\d+)?)?\s*(?:\"|”|″|'')?\s*([NSEOW])?\s*$",
                 t, flags=re.IGNORECASE)
    if m:
        g, mi, s, h = m.groups()
        v = abs(float(g.replace(",", "."))) + float((mi or "0").replace(",", ".")) / 60 \
            + float((s or "0").replace(",", ".")) / 3600
        negativo = g.startswith("-") or (h or "").upper() in ("S", "O", "W")
        return -v if negativo else v
    try:
        return float(t.replace(".", "").replace(",", ".")) if t.count(".") > 1 else float(t.replace(",", "."))
    except ValueError:
        return None


def _introduccion(doc, ctx: ContextoAire, faltantes):
    p = _parrafo(doc, "contrato con la firma")
    if p is not None:
        set_paragraph_text(p, (
            f"{ctx.cliente or 'El cliente'} contrató con la firma {EMPRESA} el monitoreo de calidad del aire para la "
            f"determinación de {ctx.nombres_contaminantes()}, en {ctx.las_estaciones()} ubicadas en {ctx.area}, "
            "con el fin de evaluar la calidad del aire en el área de influencia del proyecto."))
    # Tabla de estaciones.
    _, t = encontrar_tabla(doc, ["estacion de monitoreo", "id anla"], filas_encabezado=2)
    if t is not None:
        filas = []
        for e in ctx.estaciones:
            c = _coordenadas(e)
            filas.append([ctx.nombres[e.numero], e.codigo or "---", e.codigo_anla or "---", c["lon"], c["lat"],
                          c["norte"], c["este"]])
        _reemplazar_filas(t, 2, filas)
    else:
        faltantes.append("Tabla de estaciones de monitoreo (Introducción)")
    p = _parrafo(doc, "planes de monitoreo de")
    if p is not None:
        codigo = (ctx.proyecto.codigo or "").strip()
        texto = re.sub(r"con c[oó]digo.*$", f"con código {codigo}." if codigo else "correspondiente.", p.text)
        set_paragraph_text(p, texto)
    # Tabla de parametros medidos y nota del laboratorio de SO2.
    _, t = encontrar_tabla(doc, ["laboratorio encargado", "parametros"], filas_encabezado=1)
    if t is not None:
        _filtrar_filas(t, 1, ctx.contaminantes, columna=1)
    if SO2 not in ctx.contaminantes:
        nota = _parrafo(doc, "nota:", "dioxido de azufre", "equisam")
        if nota is not None:
            _quitar(nota._p)
    p = _parrafo(doc, "se llevo a cabo entre")
    if p is not None:
        set_paragraph_text(p, _texto_periodos(ctx))
    p = _parrafo(doc, "los resultados emitidos para")
    if p is not None:
        partes = []
        m24 = [c for c in (PM10, PM25, SO2) if c in ctx.contaminantes]
        if m24:
            partes.append(f"Los resultados emitidos para {ctx.nombres_contaminantes(m24)} fueron comparados con los "
                          "límites de tiempos de exposición de 24 horas")
        gases = [c for c in (NO2, CO, O3) if c in ctx.contaminantes]
        if gases:
            partes.append(f"los resultados para {ctx.nombres_contaminantes(gases)} fueron comparados con los límites "
                          "establecidos para tiempos de exposición de 1 hora y 8 horas, según corresponda, estos últimos "
                          "a partir de la media móvil, conforme a la Resolución 2254 de 2017")
        if partes:
            texto = ", mientras que, ".join(partes) + "."
            set_paragraph_text(p, texto[0].upper() + texto[1:])
        else:
            _quitar(p._p)


def _texto_periodos(ctx: ContextoAire) -> str:
    grupos: dict = {}
    for e in ctx.estaciones:
        if ctx.periodos.get(e.numero):
            grupos.setdefault(ctx.periodos[e.numero], []).append(ctx.nombres[e.numero])
    if not grupos:
        return "El monitoreo se llevó a cabo durante los días indicados en las tablas de resultados."
    if len(grupos) == 1:
        (d0, d1), _ = next(iter(grupos.items()))
        return f"El monitoreo se llevó a cabo {_rango_fechas(d0, d1)} en {ctx.las_estaciones()}."
    partes = [f"en {lista(n)} {_rango_fechas(d0, d1)}" for (d0, d1), n in grupos.items()]
    return "El monitoreo se llevó a cabo " + "; ".join(partes) + "."


def _reemplazar_filas(t: Table, encabezado: int, filas: list):
    """Deja tantas filas de datos como `filas` (copiando el formato de la primera)."""
    trs = list(t._tbl.iterchildren(f"{_W}tr"))
    datos = trs[encabezado:]
    if not datos:
        return
    modelo = datos[0]
    for tr in datos:
        t._tbl.remove(tr)
    for valores in filas:
        tr = copy.deepcopy(modelo)
        t._tbl.append(tr)
    for i, valores in enumerate(filas):
        celdas = t.rows[encabezado + i].cells
        for c, v in zip(celdas, valores):
            set_cell_text(c, v)


def _coincide(texto: str, contaminantes) -> bool:
    t = " " + _sin_tildes(texto) + " "
    return any(any(p in t for p in INFO[c][2]) for c in contaminantes)


def _filtrar_filas(t: Table, encabezado: int, contaminantes, columna: int = 0):
    """Quita las filas de contaminantes que no se midieron."""
    for fila in list(t.rows)[encabezado:]:
        celdas = fila.cells
        texto = celdas[min(columna, len(celdas) - 1)].text
        if not _coincide(texto, contaminantes):
            fila._tr.getparent().remove(fila._tr)


def _textos_fijos(doc, ctx: ContextoAire):
    """Parrafos fijos de la plantilla que nombran el proyecto original."""
    p = _parrafo(doc, "se describen los procedimientos relacionados")
    if p is not None:
        set_paragraph_text(p, (
            f"A continuación, se describen los procedimientos relacionados con los análisis de "
            f"{ctx.nombres_contaminantes()}, llevados a cabo en el monitoreo ejecutado en {ctx.area}. Más adelante se "
            "relacionan los métodos de monitoreo y análisis llevados a cabo en la ejecución del presente monitoreo."))
    p = _parrafo(doc, "fueron ubicadas de acuerdo con los criterios tecnicos")
    if p is not None:
        set_paragraph_text(p, (
            "Las estaciones de monitoreo a lo largo del área de influencia del estudio fueron ubicadas de acuerdo con "
            "los criterios técnicos del Manual de Diseño de Sistemas de Vigilancia de la Calidad del Aire y en función "
            f"del objetivo y condiciones de seguimiento a la calidad del aire en {ctx.area}."))


def _objetivos(doc, ctx: ContextoAire):
    p = _parrafo(doc, "evaluar la calidad de aire")
    if p is not None:
        set_paragraph_text(p, (
            f"Evaluar la calidad del aire por medio de la medición de {ctx.nombres_contaminantes()} en "
            f"{ctx.las_estaciones()} ubicadas en {ctx.area}."))
    normados = [c for c in ctx.contaminantes if c != COV]
    p = _parrafo(doc, "cuantificar las concentraciones")
    if p is not None:
        siglas = lista([INFO[c][1] for c in ctx.contaminantes])
        set_paragraph_text(p, f"Cuantificar las concentraciones de los parámetros de calidad del aire: {siglas} en el "
                              "área de influencia del proyecto.")
    p = _parrafo(doc, "comparar las concentraciones")
    if p is not None:
        if normados:
            set_paragraph_text(p, f"Comparar las concentraciones obtenidas de {lista([INFO[c][1] for c in normados])} "
                                  "con los límites máximos establecidos en la Resolución 2254 de 2017.")
        else:
            _quitar(p._p)


# Anchos (twips) de la tabla de descripcion de estaciones del informe ejemplo.
_ANCHOS_ESTACION = (1357, 1125, 1230, 2558, 2558)
_FOTO_ANCHO_CM, _FOTO_ALTO_CM = 5.4, 4.0


def _tabla_estaciones(doc, ctx: ContextoAire) -> Table:
    """Un bloque de 5 filas por estacion, como el informe ejemplo: encabezado
    (2 filas), datos, y foto (3 columnas) junto a la descripcion del punto."""
    from docx.shared import Twips

    t = construir_tabla(doc, [], [[""] * 5 for _ in range(5 * max(len(ctx.estaciones), 1))])
    combinar = []
    for k, e in enumerate(ctx.estaciones):
        r = 5 * k
        c = _coordenadas(e)
        d0, d1 = ctx.periodos.get(e.numero) or (None, None)
        textos = {
            (r, 0): "Ubicación", (r, 1): "Fecha de inicio del monitoreo", (r, 2): "Fecha de finalización del monitoreo",
            (r, 3): "Coordenadas planas Origen Nacional Único", (r + 1, 3): "Norte (Y)", (r + 1, 4): "Este (X)",
            (r + 2, 0): ctx.nombres[e.numero], (r + 2, 1): d0.isoformat() if d0 else "---",
            (r + 2, 2): d1.isoformat() if d1 else "---", (r + 2, 3): c["norte"], (r + 2, 4): c["este"],
            (r + 3, 3): "Descripción del punto de monitoreo", (r + 4, 3): e.descripcion or "---",
        }
        for (fila, col), texto in textos.items():
            t.cell(fila, col).text = texto
        encabezados = {(r, i) for i in range(5)} | {(r + 1, i) for i in range(5)} | {(r + 3, 3), (r + 3, 4)}
        for fila in range(r, r + 5):
            for col in range(5):
                celda = t.cell(fila, col)
                celda._tc.get_or_add_tcPr().clear()
                _formato_celda(celda, (fila, col) in encabezados, None, fila == r + 4)
                celda.width = Twips(_ANCHOS_ESTACION[col])
        combinar += [((r, 0), (r + 1, 0)), ((r, 1), (r + 1, 1)), ((r, 2), (r + 1, 2)), ((r, 3), (r, 4)),
                     ((r + 3, 3), (r + 3, 4)), ((r + 4, 3), (r + 4, 4)), ((r + 3, 0), (r + 4, 2))]
    for (r0, c0), (r1, c1) in combinar:
        a = t.cell(r0, c0)
        texto = a.text
        negrita = any(run.bold for p in a.paragraphs for run in p.runs)
        a.merge(t.cell(r1, c1))
        set_cell_text(a, texto)
        for p in a.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9)
                run.bold = negrita or None
    for k, e in enumerate(ctx.estaciones):
        if e.foto_ruta and os.path.exists(e.foto_ruta):
            celda = t.cell(5 * k + 3, 0)
            ancho = Cm(_FOTO_ANCHO_CM)
            try:
                from PIL import Image

                with Image.open(e.foto_ruta) as im:
                    w, h = im.size
                if h and w and _FOTO_ANCHO_CM * h / w > _FOTO_ALTO_CM:
                    ancho = Cm(_FOTO_ALTO_CM * w / h)
                celda.paragraphs[0].add_run().add_picture(e.foto_ruta, width=ancho)
            except Exception:  # noqa: BLE001 - una foto ilegible no detiene el informe
                pass
    # Ninguna fila es encabezado repetido: cada bloque trae el suyo.
    for fila in t.rows:
        for rep in fila._tr.findall(".//" + qn("w:tblHeader")):
            rep.getparent().remove(rep)
        # No partir un bloque entre paginas.
        trpr = fila._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
    return t


def _svca(doc, ctx: ContextoAire, carpeta: str, faltantes):
    """Descripcion de las estaciones (tabla) y figura de localizacion."""
    _avisar_coordenadas(ctx, faltantes)
    p = _parrafo(doc, "presenta las coordenadas")
    leyenda = next((el for el in _cuerpo(doc) if _es_leyenda(el)
                    and "descripcion de las estaciones" in _sin_tildes(_texto(el))), None)
    n_tabla = None
    if leyenda is not None:
        n_tabla = int(re.search(r"\d+", _texto(leyenda)).group())
        # Tablas de la plantilla (con fotos del proyecto original): se reemplazan.
        siguiente = leyenda.getnext()
        while siguiente is not None and _es_accesorio(siguiente):
            actual, siguiente = siguiente, siguiente.getnext()
            _quitar(actual)
        cur = Cursor(doc, leyenda)
        cur._poner(_tabla_estaciones(doc, ctx)._tbl)
        cur.fuente(anio=ctx.anio)
        sin_foto = [ctx.nombres[e.numero] for e in ctx.estaciones if not e.foto_ruta]
        if sin_foto:
            faltantes.append("Registro fotográfico: suba la foto de la estación en " + lista(sin_foto))
        if not any(e.descripcion for e in ctx.estaciones):
            faltantes.append("Descripción de las estaciones: agregue la descripción en cada estación")
    if p is not None and n_tabla:
        set_paragraph_text(p, f"La Tabla {n_tabla} presenta la descripción de las estaciones de monitoreo, las fechas "
                              "de monitoreo y sus coordenadas planas Origen Nacional.")
    # Figura de localizacion.
    leyenda_fig = next((el for el in _cuerpo(doc) if _es_leyenda(el)
                        and "localizacion de las estaciones" in _sin_tildes(_texto(el))), None)
    if leyenda_fig is None:
        return
    ruta = os.path.join(carpeta, "localizacion_aire.png")
    try:
        from .isophones import generar_mapa_localizacion

        puntos = []
        for e in ctx.estaciones:
            c = _coordenadas(e)
            if c["dd"]:
                puntos.append(SimpleNamespace(no_punto=e.numero, nombre=ctx.nombres[e.numero],
                                              longitud=str(c["dd"][1]), latitud=str(c["dd"][0])))
        if not puntos:
            raise ValueError("sin coordenadas")
        proyecto = SimpleNamespace(nombre_proyecto=ctx.proyecto.nombre_proyecto, cliente=ctx.cliente, puntos=puntos)
        generar_mapa_localizacion(proyecto, ruta, con_basemap=os.environ.get("ENGINE_BASEMAP", "1") != "0",
                                  titulo="LOCALIZACIÓN DE ESTACIONES - CALIDAD DEL AIRE")
    except Exception:  # noqa: BLE001
        ruta = None
    imagen = leyenda_fig.getnext()
    encontrados = blips(imagen) if imagen is not None else []
    if ruta and encontrados:
        reemplazar_imagen(doc, encontrados[0], ruta)
    else:
        faltantes.append("Figura de localización de las estaciones: registre coordenadas en las estaciones")
        if encontrados:
            for b in encontrados:
                dibujo = b
                while dibujo is not None and dibujo.tag != f"{_W}drawing":
                    dibujo = dibujo.getparent()
                if dibujo is not None:
                    dibujo.getparent().remove(dibujo)


def _equipos_y_validez(doc, ctx: ContextoAire, faltantes):
    _, t = encontrar_tabla(doc, ["parametro", "metodos de referencia"], filas_encabezado=1)
    if t is not None:
        _filtrar_filas(t, 1, ctx.contaminantes)
    _, t = encontrar_tabla(doc, ["documento", "descripcion", "codigo"], filas_encabezado=1)
    if t is not None:
        _filtrar_filas(t, 1, ctx.contaminantes, columna=1)
    # Instrumentos por estacion: la tabla de la plantilla es del proyecto original;
    # se deja la lista de instrumentos con columnas vacias para completar.
    leyenda = next((el for el in _cuerpo(doc) if _es_leyenda(el)
                    and "instrumentos y equipos utilizados" in _sin_tildes(_texto(el))), None)
    if leyenda is not None:
        vieja = leyenda.getnext()
        while vieja is not None and vieja.tag != f"{_W}tbl":
            vieja = vieja.getnext()
        if vieja is not None:
            instrumentos = [r.cells[0].text.strip() for r in Table(vieja, doc).rows[1:]
                            if _coincide(r.cells[0].text, ctx.contaminantes) or "meteorolog" in _sin_tildes(r.cells[0].text)]
            _quitar(vieja)
            Cursor(doc, leyenda).tabla([["Instrumento", *[ctx.nombres[e.numero] for e in ctx.estaciones]]],
                                       [[i] + [""] * len(ctx.estaciones) for i in instrumentos], izquierda=(0,))
            faltantes.append("Tabla de instrumentos y equipos por estación: completar los códigos de los equipos")
    p = _parrafo(doc, "la tabla 12 y la tabla 13")
    if p is not None:
        n = re.search(r"Tabla (\d+)", p.text)
        set_paragraph_text(p, f"La Tabla {n.group(1) if n else ''} presenta un resumen general de los equipos "
                              "utilizados en cada una de las estaciones de monitoreo.")
    # Porcentaje de datos validos.
    leyenda = next((el for el in _cuerpo(doc) if _es_leyenda(el)
                    and "porcentaje de datos validos" in _sin_tildes(_texto(el))), None)
    if leyenda is not None:
        vieja = leyenda.getnext()
        while vieja is not None and vieja.tag != f"{_W}tbl":
            vieja = vieja.getnext()
        if vieja is not None:
            _quitar(vieja)
        cols = [c for c in ctx.contaminantes]
        filas = [[ctx.nombres[e.numero], *[_pct_validos(ctx, c, e.numero) for c in cols]] for e in ctx.estaciones]
        Cursor(doc, leyenda).tabla([[("Estación de monitoreo", 1, 2), ("Porcentaje de datos válidos (%)", len(cols), 1)],
                                    [*[INFO[c][1] for c in cols]]], filas, izquierda=(0,))
    p = _parrafo(doc, "porcentajes de datos validos para cada estacion")
    if p is not None and leyenda is not None:
        n = re.search(r"\d+", _texto(leyenda)).group()
        set_paragraph_text(p, re.sub(r"en la Tabla \d+\.?\s*y la Tabla \d+", f"en la Tabla {n}", p.text))


def _pct_validos(ctx: ContextoAire, c: str, n: int) -> str:
    if c in (PM10, PM25, SO2):
        s = ctx.res.manuales.get((c, n))
        return f"{fmt(s.pct_validas, 0)}%" if s and s.pct_validas is not None else "---"
    if c == COV:
        series = [s for (k, e), s in ctx.res.cov.items() if e == n]
        return "100%" if series else "---"
    s = ctx.res.automaticos.get((c, n))
    return f"{fmt(s.pct_validos, 0)}%" if s and s.pct_validos is not None else "---"


# ------------------------------------------------------------ meteorologia
def _meteorologia(doc, ctx: ContextoAire, registros, carpeta: str, faltantes):
    titulo = next((el for el in _cuerpo(doc) if _es_titulo(el, 1, "meteorolog")), None)
    if titulo is None:
        faltantes.append("Capítulo de meteorología")
        return
    cur = Cursor(doc, titulo)
    if not registros or ctx.inicio is None:
        cur.parrafo("No se suministraron datos meteorológicos de la estación del área de influencia del proyecto "
                    "para el periodo de monitoreo.")
        if not registros:
            faltantes.append("Datos meteorológicos: suba el archivo de la estación para completar el capítulo")
        return
    sel = [r for r in registros if ctx.inicio <= r.fecha.date() <= ctx.fin]
    if not sel:
        cur.parrafo("El archivo de datos meteorológicos no contiene registros del periodo de monitoreo.")
        faltantes.append("Datos meteorológicos: el archivo no tiene datos de los días de monitoreo")
        return
    por_dia: dict = {}
    for r in sel:
        por_dia.setdefault(r.fecha.date(), []).append(r)

    def media(xs):
        xs = [x for x in xs if x is not None]
        return sum(xs) / len(xs) if xs else None

    dias = sorted(por_dia)
    temp = [media(r.temperatura for r in por_dia[d]) for d in dias]
    hum = [media(r.humedad for r in por_dia[d]) for d in dias]
    pres = [media(r.presion / 0.750062 if r.presion is not None else None for r in por_dia[d]) for d in dias]
    lluvia = [max([r.lluvia for r in por_dia[d] if r.lluvia is not None], default=None) for d in dias]
    viento = [media(r.viento for r in por_dia[d]) for d in dias]

    n_tabla = cur.numero("Tabla")
    cur.parrafo("La información meteorológica fue descargada a partir de la estación ubicada en el área de influencia "
                "del proyecto, la cual reporta datos continuos de temperatura (°C), precipitación (mm), humedad "
                "relativa (%), presión atmosférica (mbar), velocidad y dirección del viento durante los días de "
                "monitoreo. Estos datos son procesados con el fin de obtener un dato diario promedio de las variables.")
    cur.parrafo(f"En la Tabla {n_tabla} se consignan los datos promedio diarios (temperatura, humedad relativa, "
                "presión atmosférica y velocidad del viento) y la máxima de precipitación para los días de monitoreo.")
    cur.leyenda("Tabla", "Datos meteorológicos en el área de influencia del proyecto")
    filas = [[d.isoformat(), fmt(t), fmt(h), fmt(p), fmt(ll), fmt(v)]
             for d, t, h, p, ll, v in zip(dias, temp, hum, pres, lluvia, viento)]
    for etiqueta, f in (("Promedio", media), ("Máximo", lambda xs: max([x for x in xs if x is not None], default=None)),
                        ("Mínimo", lambda xs: min([x for x in xs if x is not None], default=None))):
        filas.append([etiqueta, *[fmt(f(serie)) for serie in (temp, hum, pres, lluvia, viento)]])
    cur.tabla([["Fecha", "Temperatura (°C)", "Humedad Relativa (%)", "Presión Atmosférica (mbar)",
                "Precipitación Máx. (mm)", "Velocidad del Viento (m/s)"]], filas)
    cur.fuente(anio=ctx.anio)

    variables = [
        ("Temperatura", temp, "°C", "la temperatura promedio", "Temperatura ambiente promedio durante los días de monitoreo",
         "Temperatura (°C)", "temperatura"),
        ("Humedad relativa", hum, "%", "la humedad relativa promedio",
         "Humedad relativa promedio durante los días de monitoreo", "Humedad relativa (%)", "humedad"),
        ("Presión atmosférica", pres, "mbar", "la presión atmosférica promedio",
         "Presión atmosférica promedio durante los días de monitoreo", "Presión (mbar)", "presion"),
    ]
    for titulo_v, serie, unidad, sujeto, leyenda, eje, clave in variables:
        validos = [(d, v) for d, v in zip(dias, serie) if v is not None]
        if not validos:
            continue
        prom = media(serie)
        d_max, v_max = max(validos, key=lambda x: x[1])
        d_min, v_min = min(validos, key=lambda x: x[1])
        cur.subtitulo(titulo_v)
        n = cur.numero("Gráfica")
        cur.parrafo(f"Durante el periodo comprendido {_rango_fechas(dias[0], dias[-1])}, {sujeto} fue de "
                    f"{fmt(prom)} {unidad}, con un valor máximo de {fmt(v_max)} {unidad}, registrado el "
                    f"{fecha_larga(d_max)}, y un valor mínimo de {fmt(v_min)} {unidad}, registrado el "
                    f"{fecha_larga(d_min)} (ver Gráfica {n}).")
        cur.leyenda("Gráfica", leyenda)
        cur.imagen(gr.meteo_diaria(dias, serie, eje, gr.nombre_archivo(carpeta, "meteo", clave)))
        cur.fuente(anio=ctx.anio)
    validos = [(d, v) for d, v in zip(dias, lluvia) if v is not None]
    if validos:
        cur.subtitulo("Precipitación")
        n = cur.numero("Gráfica")
        d_max, v_max = max(validos, key=lambda x: x[1])
        secos = sum(1 for _, v in validos if v == 0)
        cur.parrafo(f"La precipitación máxima diaria presentó valores entre {fmt(min(v for _, v in validos))} mm y "
                    f"{fmt(v_max)} mm durante el periodo evaluado; el valor más alto se registró el "
                    f"{fecha_larga(d_max)}. Se presentaron {secos} día(s) sin precipitación (ver Gráfica {n}).")
        cur.leyenda("Gráfica", "Precipitación máxima diaria durante los días de monitoreo")
        cur.imagen(gr.meteo_diaria(dias, lluvia, "Precipitación (mm)", gr.nombre_archivo(carpeta, "meteo", "lluvia"),
                                   color="#2E75B6", barras=True))
        cur.fuente(anio=ctx.anio)
    _viento(cur, ctx, sel, dias, viento, carpeta)


def _viento(cur: Cursor, ctx: ContextoAire, sel, dias, viento, carpeta):
    from .meteorologia import analizar, generar_graficas_meteo

    a = analizar(sel, set(dias), [])
    if a is None or not any(v is not None for v in viento):
        return
    try:
        sub = os.path.join(carpeta, "viento")
        os.makedirs(sub, exist_ok=True)
        graficas = generar_graficas_meteo(a, sub)
    except Exception:  # noqa: BLE001
        graficas = {}
    totales = {d: sum(v) for d, v in a.frecuencias.items()}
    predominante = max(totales, key=totales.get) if totales else None
    prom = sum(v for v in viento if v is not None) / len([v for v in viento if v is not None])
    cur.subtitulo("Velocidad y dirección del viento")
    n = cur.numero("Gráfica")
    texto = (f"Durante el periodo de monitoreo la velocidad promedio del viento fue de {fmt(prom)} m/s"
             + (f", con una procedencia principalmente desde el {predominante} ({fmt(totales[predominante], 1)} % de "
                "los registros)" if predominante and totales[predominante] > 0 else "")
             + f"; las calmas (velocidad inferior a 0,5 m/s) representaron el {fmt(a.calmas_pct, 1)} % de los "
             f"registros (ver Gráfica {n}).")
    cur.parrafo(texto)
    cur.leyenda("Gráfica", "Rosa de vientos durante los días de monitoreo")
    cur.imagen(graficas.get("viento_combinado"))
    cur.fuente(anio=ctx.anio)


# --------------------------------------------------------------- resultados
def _estaciones_con(ctx, series: dict):
    return [e for e in ctx.estaciones if e.numero in series]


def _texto_estacion_manual(nombre, sigla, est, bajo_lc, primero: bool) -> str:
    if bajo_lc:
        return (f"{'En' if not primero else 'Para'} {nombre}, los resultados diarios fueron inferiores a valores "
                f"comprendidos entre {fmt(est.minimo)} µg/m³ y {fmt(est.maximo)} µg/m³; el menor valor correspondió al "
                f"{fecha_larga(est.fecha_min)} y el mayor al {fecha_larga(est.fecha_max)}. El promedio aritmético "
                f"fue inferior a {fmt(est.promedio)} µg/m³.")
    return (f"{nombre} registró niveles de inmisión diarios de {sigla} que oscilaron entre {fmt(est.minimo)} µg/m³, "
            f"correspondiente al {fecha_larga(est.fecha_min)}, y {fmt(est.maximo)} µg/m³, reportado el "
            f"{fecha_larga(est.fecha_max)}. Para esta estación se determinó un promedio aritmético de "
            f"{fmt(est.promedio)} µg/m³ durante el periodo de monitoreo.")


def _comparacion(valores: dict, que: str, bajo_lc=False) -> Optional[str]:
    """`valores`: {estacion: promedio}."""
    if len(valores) < 2:
        return None
    orden = sorted(valores.items(), key=lambda x: -x[1])
    pref = "inferior a " if bajo_lc else ""
    resto = lista([f"{n} ({pref}{fmt(v, miles=True)} µg/m³)" for n, v in orden[1:]])
    return (f"Al comparar los valores promedio de {que} obtenidos en las estaciones de monitoreo, se observa que "
            f"{orden[0][0]} registró el mayor promedio ({pref}{fmt(orden[0][1], miles=True)} µg/m³), seguida de "
            f"{resto}.")


def _normativo(series: dict, limite: float, que: str, exposicion: str, nombres: dict) -> str:
    """`series`: {estacion: [(fecha, valor)]}."""
    excedencias = {n: [f for f, v in datos if v is not None and v > limite] for n, datos in series.items()}
    total = sum(len(x) for x in excedencias.values())
    if total == 0:
        return (f"Desde el punto de vista normativo, las concentraciones de {que} registradas en "
                f"{'la estación' if len(series) == 1 else 'todas las estaciones'} se encuentran por debajo del nivel "
                f"máximo permisible de {fmt(limite, 0, True)} µg/m³ para un tiempo de exposición de {exposicion}, "
                "establecido en la Resolución 2254 de 2017.")
    partes = [f"{nombres[n]} ({lista([fecha_larga(f) for f in fs])})" for n, fs in excedencias.items() if fs]
    return (f"Desde el punto de vista normativo, se presentaron {total} excedencia(s) del nivel máximo permisible de "
            f"{fmt(limite, 0, True)} µg/m³ para un tiempo de exposición de {exposicion} (Resolución 2254 de 2017): "
            + "; ".join(partes) + ".")


def _manual(cur: Cursor, ctx: ContextoAire, c: str, carpeta: str):
    series = {n: s for (k, n), s in ctx.res.manuales.items() if k == c}
    estaciones = _estaciones_con(ctx, series)
    sigla = INFO[c][1]
    limite = LIMITES[c]["24h"]
    cur.titulo(f"Parámetro: {TITULOS[c]}", 3)
    n_tabla = cur.numero("Tabla")
    fechas = sorted({f for n in series for f, _ in series[n].datos})
    filas = []
    for f in fechas:
        fila = [f.isoformat()]
        for e in estaciones:
            m = next((m for m in series[e.numero].muestras if m.fecha == f and m.concentracion is not None), None)
            fila.append(conc(m.concentracion, m.bajo_lc) if m else "---")
        filas.append(fila)
    estadisticas = {e.numero: series[e.numero].estadistica() for e in estaciones}
    for etiqueta, attr in (("Promedio aritmético", "promedio"), ("Máximo diario", "maximo"), ("Mínimo diario", "minimo")):
        filas.append([etiqueta, *[conc(getattr(estadisticas[e.numero], attr), series[e.numero].bajo_lc)
                                  if estadisticas[e.numero] else "---" for e in estaciones]])
    rellenos = {}
    for i, f in enumerate(fechas):
        for j, e in enumerate(estaciones, start=1):
            m = next((m for m in series[e.numero].muestras if m.fecha == f), None)
            if m and m.concentracion is not None and m.concentracion > limite:
                rellenos[(i, j)] = "F8CBAD"
    cur.leyenda("Tabla", f"Concentraciones diarias de {TITULOS[c]}")
    cur.tabla([[("Fecha final de monitoreo", 1, 2), (f"Nivel de inmisión (µg/m³) – {TITULOS[c]}", len(estaciones), 1)],
               [*[ctx.nombres[e.numero] for e in estaciones]]], filas, rellenos)
    todos_lc = all(series[e.numero].bajo_lc for e in estaciones)
    if any(series[e.numero].bajo_lc for e in estaciones):
        cur.nota("Nota: los resultados con el símbolo «<» corresponden a masas inferiores al límite de cuantificación "
                 "del método de análisis establecido por el laboratorio.")
    cur.fuente(anio=ctx.anio)
    n_graf = None
    if not todos_lc:
        n_graf = cur.numero("Gráfica")
        cur.leyenda("Gráfica", f"Concentraciones diarias de {TITULOS[c]} versus el límite máximo permisible")
        cur.imagen(gr.barras_diarias({ctx.nombres[e.numero]: series[e.numero].datos for e in estaciones}, limite,
                                     f"{sigla} (µg/m³)", gr.nombre_archivo(carpeta, "conc", c)))
        cur.fuente(anio=ctx.anio)
    ref = f"la Tabla {n_tabla}" + (f" y la Gráfica {n_graf}" if n_graf else "")
    for k, e in enumerate(estaciones):
        est = estadisticas[e.numero]
        if est is None:
            continue
        texto = _texto_estacion_manual(ctx.nombres[e.numero], sigla, est, series[e.numero].bajo_lc, k == 0)
        if k == 0:
            texto = f"De acuerdo con {ref}, " + (texto[0].lower() + texto[1:] if texto.startswith("Para ") else texto)
        cur.parrafo(texto)
    comp = _comparacion({ctx.nombres[e.numero]: estadisticas[e.numero].promedio for e in estaciones
                         if estadisticas[e.numero]}, sigla, todos_lc)
    if comp:
        cur.parrafo(comp)
    if todos_lc:
        cur.parrafo("Debido a que los resultados se encuentran precedidos por el símbolo «<», corresponden a "
                    "concentraciones inferiores al límite de cuantificación y no a concentraciones determinadas de "
                    "manera exacta.")
    cur.parrafo(_normativo({e.numero: series[e.numero].datos for e in estaciones}, limite, sigla, "24 horas",
                           ctx.nombres))


def _automatico(cur: Cursor, ctx: ContextoAire, c: str, carpeta: str):
    series = {n: s for (k, n), s in ctx.res.automaticos.items() if k == c}
    estaciones = _estaciones_con(ctx, series)
    sigla = INFO[c][1]
    exposiciones = [(clave, attr, texto) for clave, attr, texto in
                    (("1h", "max_horario", "Máximo horario"), ("8h", "max_8h", "Máximo octohorario"))
                    if clave in LIMITES[c]]
    cur.titulo(f"Parámetro: {TITULOS[c]}", 3)
    n_tabla = cur.numero("Tabla")
    fechas = sorted({d.fecha for s in series.values() for d in s.dias})
    por_dia = {e.numero: {d.fecha: d for d in series[e.numero].dias} for e in estaciones}
    filas, rellenos = [], {}
    for i, f in enumerate(fechas):
        fila = [f.isoformat()]
        j = 1
        for clave, attr, _ in exposiciones:
            for e in estaciones:
                d = por_dia[e.numero].get(f)
                v = getattr(d, attr) if d else None
                fila.append(fmt(v))
                if v is not None and v > LIMITES[c][clave]:
                    rellenos[(i, j)] = "F8CBAD"
                j += 1
        filas.append(fila)
    for etiqueta, fn in (("Promedio aritmético", lambda xs: sum(xs) / len(xs)), ("Máxima", max), ("Mínima", min)):
        fila = [etiqueta]
        for clave, attr, _ in exposiciones:
            for e in estaciones:
                xs = [getattr(d, attr) for d in series[e.numero].dias if getattr(d, attr) is not None]
                fila.append(fmt(fn(xs)) if xs else "---")
        filas.append(fila)
    cur.leyenda("Tabla", f"Resultados de las concentraciones de {TITULOS[c]}")
    cur.tabla([[("Fecha de monitoreo", 1, 2),
                *[(f"Nivel de inmisión (µg/m³) – {texto} – {sigla}", len(estaciones), 1) for _, _, texto in exposiciones]],
               [*[ctx.nombres[e.numero] for _ in exposiciones for e in estaciones]]], filas, rellenos)
    cur.fuente(anio=ctx.anio)

    # Todas las graficas son de barras, en el orden del informe ejemplo: por cada tiempo
    # de exposicion, la serie de cada estacion y luego los maximos diarios vs. el limite.
    graficas = []
    for clave, attr, texto in exposiciones:
        nombre_serie = "horarias" if clave == "1h" else "octohorarias"
        for e in estaciones:
            s = series[e.numero]
            pares = s.horas if clave == "1h" else s.medias_8h
            graficas.append((f"Concentraciones {nombre_serie} de {TITULOS[c]} – {ctx.nombres[e.numero]}",
                             gr.barras_horarias(pares, f"{sigla} (µg/m³)",
                                                gr.nombre_archivo(carpeta, nombre_serie, c, e.numero),
                                                "Concentración horaria" if clave == "1h" else "Media móvil de 8 horas")))
        graficas.append((f"Concentraciones {'máximas horarias' if clave == '1h' else 'máximas octohorarias'} de "
                         f"{TITULOS[c]} versus el límite máximo permisible",
                         gr.barras_diarias({ctx.nombres[e.numero]: [(d.fecha, getattr(d, attr)) for d in series[e.numero].dias]
                                            for e in estaciones}, LIMITES[c][clave], f"{sigla} (µg/m³)",
                                           gr.nombre_archivo(carpeta, "max", c, clave))))
    n_graf = cur.numero("Gráfica")
    for clave, attr, texto in exposiciones:
        nombre_exp = "concentraciones máximas horarias" if clave == "1h" else "medias móviles máximas de ocho horas"
        valores = {}
        for k, e in enumerate(estaciones):
            xs = [(d.fecha, getattr(d, attr)) for d in series[e.numero].dias if getattr(d, attr) is not None]
            if not xs:
                continue
            prom = sum(v for _, v in xs) / len(xs)
            f_max, v_max = max(xs, key=lambda x: x[1])
            f_min, v_min = min(xs, key=lambda x: x[1])
            valores[ctx.nombres[e.numero]] = prom
            inicio = f"De acuerdo con la Tabla {n_tabla}, en cuanto a las {nombre_exp} de {sigla}, " if k == 0 else ""
            frase = (f"{ctx.nombres[e.numero]} presentó valores comprendidos entre {fmt(v_min, miles=True)} µg/m³ "
                     f"({fecha_larga(f_min)}) y {fmt(v_max, miles=True)} µg/m³ ({fecha_larga(f_max)}), con un promedio "
                     f"aritmético de {fmt(prom, miles=True)} µg/m³.")
            cur.parrafo(inicio + frase)
        comp = _comparacion(valores, f"las {nombre_exp}")
        if comp:
            cur.parrafo(comp.replace(" µg/m³", " µg/m³"))
        cur.parrafo(_normativo({e.numero: [(d.fecha, getattr(d, attr)) for d in series[e.numero].dias]
                                for e in estaciones}, LIMITES[c][clave], sigla,
                               "una hora" if clave == "1h" else "ocho horas", ctx.nombres))
    ultima = n_graf + len(graficas) - 1
    cur.parrafo(f"Desde la Gráfica {n_graf} hasta la Gráfica {ultima} se presenta el comportamiento de las "
                f"concentraciones registradas en cada estación y su comparación con los límites máximos permisibles."
                if ultima > n_graf else f"En la Gráfica {n_graf} se presenta el comportamiento de las concentraciones.")
    for titulo, ruta in graficas:
        cur.leyenda("Gráfica", titulo)
        cur.imagen(ruta)
        cur.nota("Nota: gráfica en escala logarítmica, base 10.")
        cur.fuente(anio=ctx.anio)


def _cov(cur: Cursor, ctx: ContextoAire):
    cur.titulo(f"Parámetro: {TITULOS[COV]}", 3)
    primera = cur.numero("Tabla")
    tablas = 0
    for e in ctx.estaciones:
        series = {k: s for (k, n), s in ctx.res.cov.items() if n == e.numero}
        if not series:
            continue
        compuestos = [k for k in COMPUESTOS_COV if k in series]
        fechas = sorted({m.fecha for s in series.values() for m in s.muestras if m.fecha})
        filas = []
        for f in fechas:
            fila = [f.isoformat()]
            for k in compuestos:
                m = next((m for m in series[k].muestras if m.fecha == f), None)
                fila.append(conc(m.concentracion, m.bajo_lc, 3) if m else "---")
            filas.append(fila)
        if tablas == 0:
            cur.parrafo(f"En este numeral se presentan los resultados de los Compuestos Orgánicos Volátiles expresados "
                        f"como {lista(compuestos).lower()} a condiciones de referencia. Los resultados se presentan a "
                        "nivel informativo, dado que la Resolución 2254 de 2017 no establece un límite máximo permisible "
                        "para estos compuestos.")
        cur.leyenda("Tabla", f"Resultados COV – {ctx.nombres[e.numero]}")
        cur.tabla([[("Fecha de monitoreo", 1, 2), ("Niveles de inmisión de Compuestos Orgánicos Volátiles (COV)",
                                                    len(compuestos), 1)],
                   [*[f"{k} (µg/m³)" for k in compuestos]]], filas)
        if any(s.bajo_lc for s in series.values()):
            cur.nota("Nota: los resultados con el símbolo «<» fueron inferiores a los límites de cuantificación del "
                     "método de análisis establecidos por el laboratorio.")
        cur.fuente(anio=ctx.anio)
        tablas += 1
    if tablas:
        todos = all(s.bajo_lc for s in ctx.res.cov.values())
        rango = f"la Tabla {primera}" if tablas == 1 else f"la Tabla {primera} hasta la Tabla {primera + tablas - 1}"
        cur.parrafo(f"De acuerdo con {rango}, " + (
            "todos los resultados de COV fueron inferiores a los límites de cuantificación del método de análisis."
            if todos else "se observan las concentraciones de COV reportadas en cada estación de monitoreo."))


def _estadistica(cur: Cursor, ctx: ContextoAire):
    filas_var = [("Número de datos", "n", 0), ("Promedio aritmético", "promedio", 2),
                 ("Desviación estándar", "desviacion", 2), ("Coeficiente de variación", "cv", 2),
                 ("Mediana (m)", "mediana", 2), ("Q1", "q1", 2), ("Q3", "q3", 2), ("IQR=Q3-Q1", "iqr", 2),
                 ("m-3*IQR", "lim_inf", 2), ("m+3*IQR", "lim_sup", 2), ("Máxima concentración reportada", "maximo", 2),
                 ("Fecha máxima concentración reportada", "fecha_max", None),
                 ("Mínima concentración reportada", "minimo", 2), ("Fecha mínima concentración reportada", "fecha_min", None),
                 ("Concentraciones atípicas", "atipicos", 0)]
    primera = cur.numero("Tabla")
    bloques = []
    for c in (PM10, PM25, SO2):
        cols = [(ctx.nombres[e.numero], s.estadistica(), s.bajo_lc)
                for e in ctx.estaciones if (s := ctx.res.manuales.get((c, e.numero)))]
        if cols:
            bloques.append((INFO[c][1], [("", cols)]))
    for c in (NO2, CO, O3):
        grupos = []
        if c in (CO, O3):
            grupos.append(("Máximo octohorario", [(ctx.nombres[e.numero], s.estadistica_8h(), False) for e in ctx.estaciones
                                                  if (s := ctx.res.automaticos.get((c, e.numero)))]))
        if c in (CO, NO2):
            grupos.append(("Horario", [(ctx.nombres[e.numero], s.estadistica_1h(), False) for e in ctx.estaciones
                                       if (s := ctx.res.automaticos.get((c, e.numero)))]))
        grupos = [g for g in grupos if g[1]]
        if grupos:
            bloques.append((INFO[c][1], grupos))
    for k in COMPUESTOS_COV:
        cols = [(ctx.nombres[e.numero], s.estadistica(), s.bajo_lc)
                for e in ctx.estaciones if (s := ctx.res.cov.get((k, e.numero)))]
        if cols:
            bloques.append((f"COV – {k}", [("", cols)]))
    if not bloques:
        return
    cur.parrafo("El análisis estadístico permite identificar los datos atípicos de un conjunto de concentraciones y "
                "la consistencia de los datos, mediante el cálculo de las variables estadísticas (número de datos, "
                "promedio aritmético, desviación estándar, mediana, percentil 75 % (Q3), percentil 25 % (Q1), distancia "
                "intercuartílica (IQR = Q3 − Q1) y los límites m ± 3·IQR). Se considera atípica una concentración por "
                "fuera de dichos límites.")
    cur.parrafo(f"Desde la Tabla {primera} hasta la Tabla {primera + len(bloques) - 1} se presentan los resultados del "
                "análisis estadístico para cada parámetro evaluado." if len(bloques) > 1 else
                f"En la Tabla {primera} se presentan los resultados del análisis estadístico.")
    for nombre, grupos in bloques:
        columnas = [col for _, cols in grupos for col in cols]
        enc = [[("Variable estadística", 1, 2 if len(grupos) > 1 else 1),
                *([(f"{g} (µg/m³)", len(cols), 1) for g, cols in grupos] if len(grupos) > 1
                  else [n for n, _, _ in columnas])]]
        if len(grupos) > 1:
            enc.append([*[n for n, _, _ in columnas]])
        filas = []
        for etiqueta, attr, dec in filas_var:
            fila = [etiqueta]
            for _, est, lc in columnas:
                v = getattr(est, attr) if est else None
                if v is None:
                    fila.append("---")
                elif dec is None:
                    fila.append(v.isoformat() if hasattr(v, "isoformat") else str(v))
                elif dec == 0:
                    fila.append(str(int(v)))
                else:
                    fila.append(conc(v, lc and attr not in ("cv",), 3 if nombre.startswith("COV") else dec))
            filas.append(fila)
        cur.leyenda("Tabla", f"Resultados del análisis estadístico para {nombre}")
        cur.tabla(enc, filas, izquierda=(0,))
        cur.fuente(anio=ctx.anio)


def _ica(cur: Cursor, ctx: ContextoAire, carpeta: str):
    colores = {nombre: color for _, _, nombre, color in CATEGORIAS_ICA}
    hex_color = {"Verde": "00E400", "Amarillo": "FFFF00", "Naranja": "FF7E00", "Rojo": "FF0000",
                 "Púrpura": "8F3F97", "Marrón": "7E0023"}
    primera = cur.numero("Tabla")
    no_calc = [x for x in ("SO2", "COV") if x in ctx.contaminantes]
    texto = ("A continuación se presentan los resultados del cálculo del ICA en cada una de las estaciones y para cada "
             "uno de los parámetros evaluados, según la metodología descrita en el presente informe.")
    if no_calc:
        texto += (" Para " + lista(no_calc) + " no se calcula el ICA: los puntos de corte de la Resolución 2254 de "
                  "2017 para el SO2 están definidos para 1 hora (las muestras son de 24 horas) y para los COV no se "
                  "establecen puntos de corte.")
    cur.parrafo(texto)
    tablas = 0
    conteos: dict = {}
    for e in ctx.estaciones:
        dias = ctx.res.dias_ica(e.numero)
        if not dias:
            continue
        contaminantes = list(dias)
        fechas = sorted({f for xs in dias.values() for f, *_ in xs})
        filas, rellenos = [], {}
        for i, f in enumerate(fechas):
            fila = [f.isoformat()]
            for j, c in enumerate(contaminantes, start=1):
                x = next((x for x in dias[c] if x[0] == f), None)
                if x and x[2] is not None:
                    fila.append(f"{fmt(x[2], 0)} – {x[3]}")
                    rellenos[(i, j)] = hex_color[colores[x[3]]]
                    conteos.setdefault(c, {}).setdefault(ctx.nombres[e.numero], {})
                    conteos[c][ctx.nombres[e.numero]][x[3]] = conteos[c][ctx.nombres[e.numero]].get(x[3], 0) + 1
                else:
                    fila.append("---")
            filas.append(fila)
        cur.leyenda("Tabla", f"Índices de calidad del aire – {ctx.nombres[e.numero]}")
        cur.tabla([[("Fecha", 1, 2), ("Índice de calidad del aire (ICA) – categoría", len(contaminantes), 1)],
                   [INFO[c][1] for c in contaminantes]], filas, rellenos)
        cur.fuente(anio=ctx.anio)
        tablas += 1
    for c, por_estacion in conteos.items():
        cur.leyenda("Gráfica", f"Índice de Calidad del Aire (ICA) – {INFO[c][1]}")
        cur.imagen(gr.ica_categorias(por_estacion, gr.nombre_archivo(carpeta, "ica", c)))
        cur.fuente(anio=ctx.anio)
    return conteos, primera


# ------------------------------------------------------------- conclusiones
def _conclusiones(doc, ctx: ContextoAire, conteos: dict):
    titulo = next((el for el in _cuerpo(doc) if _es_titulo(el, 1, "conclusiones")), None)
    if titulo is None:
        return
    intro = titulo.getnext()
    while intro is not None and not _texto(intro).strip():
        intro = intro.getnext()
    modelo = next((el for el in _cuerpo(doc) if el.getparent() is not None and el.tag == f"{_W}p"
                   and el.find(f"{_W}pPr/{_W}numPr") is not None
                   and _es_despues(doc, el, titulo) and not _es_titulo(el, 1)), None)
    if intro is not None and modelo is not intro:
        set_paragraph_text(Paragraph(intro, doc), (
            f"De acuerdo con las mediciones ejecutadas durante el periodo de monitoreo para determinar la calidad del "
            f"aire en {ctx.las_estaciones()} ubicadas en {ctx.area}; {EMPRESA}, concluye lo siguiente:"))
    vinetas = []
    if ctx.inicio:
        vinetas.append(f"Se desarrolló el monitoreo de {ctx.nombres_contaminantes()} {_rango_fechas(ctx.inicio, ctx.fin)}.")
    for c in (PM10, PM25, SO2):
        series = {n: s for (k, n), s in ctx.res.manuales.items() if k == c}
        if not series:
            continue
        lim = LIMITES[c]["24h"]
        exc = sum(1 for s in series.values() for _, v in s.datos if v > lim)
        lc = " y fueron inferiores al límite de cuantificación del método de análisis" if all(
            s.bajo_lc for s in series.values()) else ""
        if exc:
            vinetas.append(f"Los resultados de {INFO[c][0]} presentaron {exc} excedencia(s) del límite máximo "
                           f"({lim} µg/m³) estipulado en la Resolución 2254 de 2017 para 24 horas.")
        else:
            vinetas.append(f"Los resultados de {INFO[c][0]} reportaron concentraciones por debajo del límite máximo "
                           f"({lim} µg/m³) estipulado en la Resolución 2254 de 2017 para 24 horas en "
                           f"{_en_estaciones(len(series))}{lc}.")
    for c in (NO2, CO, O3):
        series = {n: s for (k, n), s in ctx.res.automaticos.items() if k == c}
        if not series:
            continue
        partes, exc = [], 0
        for clave, attr in (("1h", "max_horario"), ("8h", "max_8h")):
            if clave in LIMITES[c]:
                lim = LIMITES[c][clave]
                partes.append(f"{fmt(lim, 0, True)} µg/m³ para {'1 hora' if clave == '1h' else '8 horas'}")
                exc += sum(1 for s in series.values() for d in s.dias if getattr(d, attr) is not None
                           and getattr(d, attr) > lim)
        if exc:
            vinetas.append(f"Para {INFO[c][0]} se presentaron {exc} excedencia(s) de los límites máximos permisibles "
                           f"({lista(partes)}) de la Resolución 2254 de 2017.")
        else:
            vinetas.append(f"Las concentraciones de {INFO[c][0]} se encontraron por debajo de los límites máximos "
                           f"permisibles ({lista(partes)}) establecidos en la Resolución 2254 de 2017 en "
                           f"{_en_estaciones(len(series))}.")
    if ctx.res.cov:
        todos = all(s.bajo_lc for s in ctx.res.cov.values())
        vinetas.append("Los resultados de Compuestos Orgánicos Volátiles (COV) "
                       + ("presentaron valores inferiores a los límites de cuantificación del método de análisis. "
                          if todos else "se reportan a nivel informativo. ")
                       + "La Resolución 2254 de 2017 no establece límites máximos permisibles para estos parámetros, "
                         "por lo cual no se emite juicio normativo.")
    for c, por_estacion in conteos.items():
        total: dict = {}
        for cats in por_estacion.values():
            for cat, n in cats.items():
                total[cat] = total.get(cat, 0) + n
        if not total:
            continue
        orden = sorted(total.items(), key=lambda x: -x[1])
        vinetas.append(f"El Índice de Calidad del Aire (ICA) para {INFO[c][1]} se ubicó principalmente en la categoría "
                       f"«{orden[0][0]}» ({orden[0][1]} de {sum(total.values())} registros diarios en el conjunto de "
                       "estaciones)" + (", seguida de " + lista([f"«{k}» ({v})" for k, v in orden[1:]]) if len(orden) > 1
                                       else "") + ".")
    if modelo is None:
        cur = Cursor(doc, intro if intro is not None else titulo)
        for v in vinetas:
            cur.parrafo(v, "List Paragraph")
        return
    anterior = modelo
    for k, v in enumerate(vinetas):
        el = modelo if k == 0 else copy.deepcopy(modelo)
        if k:
            anterior.addnext(el)
            anterior = el
        set_paragraph_text(Paragraph(el, doc), v)


def _renumerar(doc):
    """Numera las leyendas en orden (el numero visible del campo SEQ) y corrige
    las referencias 'Tabla N' / 'Figura N' / 'Gráfica N' del texto fijo de la
    plantilla, que cambian al quitar o agregar tablas."""
    mapa: dict = {}
    cuenta: dict = {}
    for el in doc.element.body.iterchildren():
        if not _es_leyenda(el):
            continue
        m = re.match(r"\s*(Tabla|Figura|Gráfica|Ecuación)\s+(\d+)", _texto(el))
        if not m:
            continue
        tipo, viejo = m.group(1), m.group(2)
        cuenta[tipo] = cuenta.get(tipo, 0) + 1
        nuevo = str(cuenta[tipo])
        mapa.setdefault(tipo, {}).setdefault(viejo, nuevo)
        # Resultado visible del campo: el primer texto numerico tras 'separate'.
        separado = False
        for nodo in el.iter():
            if nodo.tag == f"{_W}fldChar" and nodo.get(qn("w:fldCharType")) == "separate":
                separado = True
            elif separado and nodo.tag == f"{_W}t" and (nodo.text or "").strip().isdigit():
                nodo.text = nodo.text.replace(nodo.text.strip(), nuevo)
                break
    patron = re.compile(r"\b(Tabla|Figura|Gráfica)(\s+No\.)?\s+(\d+)")
    for el in doc.element.body.iterchildren():
        if _es_titulo(el, 1, "meteorolog"):
            break  # el texto generado ya trae los numeros correctos
        if el.tag != f"{_W}p" or _es_leyenda(el):
            continue
        ts = list(el.iter(f"{_W}t"))
        completo = "".join(t.text or "" for t in ts)
        cambios = []
        for m in patron.finditer(completo):
            nuevo = mapa.get(m.group(1), {}).get(m.group(3))
            if nuevo and nuevo != m.group(3):
                cambios.append((m.start(3), m.end(3), nuevo))
        for ini, fin, nuevo in reversed(cambios):
            pos = 0
            primero = True
            for t in ts:
                texto = t.text or ""
                a0, b0 = pos, pos + len(texto)
                pos = b0
                a, b = max(ini, a0), min(fin, b0)
                if a < b:
                    t.text = texto[:a - a0] + (nuevo if primero else "") + texto[b - a0:]
                    primero = False


def _es_despues(doc, el, ref) -> bool:
    for x in doc.element.body.iterchildren():
        if x is ref:
            return True
        if x is el:
            return False
    return False


def _en_estaciones(n: int) -> str:
    return "la estación de monitoreo" if n == 1 else f"las {cantidad(n, 'estación', 'estaciones')} de monitoreo"


# ------------------------------------------------------------------ principal
def generar_informe_aire(res: ResultadosAire, ruta_plantilla: str, ruta_salida: str, carpeta: str, datos,
                         registros_meteo=None, hoy: Optional[dt.date] = None) -> tuple[str, list]:
    """Genera el informe Word. Devuelve (ruta, partes a revisar)."""
    doc = docx.Document(ruta_plantilla)
    limpiar_plantilla(doc)
    ctx = ContextoAire(res, datos, hoy)
    faltantes: list = []
    os.makedirs(carpeta, exist_ok=True)

    _portada_y_encabezado(doc, ctx)
    _cuadro_control(doc, SimpleNamespace(datos=datos, fecha=ctx.fecha))
    _introduccion(doc, ctx, faltantes)
    _objetivos(doc, ctx)
    _textos_fijos(doc, ctx)
    if ctx.cliente:
        _cliente(doc, SimpleNamespace(cliente=ctx.cliente, fecha=ctx.fecha, datos=datos), faltantes)
    else:
        faltantes.append("Información del cliente: el proyecto no tiene cliente")
    _svca(doc, ctx, carpeta, faltantes)
    _equipos_y_validez(doc, ctx, faltantes)
    _meteorologia(doc, ctx, registros_meteo, carpeta, faltantes)

    resultados = next((el for el in _cuerpo(doc) if _es_titulo(el, 1, "resultados")), None)
    niveles = next((el for el in _cuerpo(doc) if _es_titulo(el, 2, "niveles de inmision")), None)
    if resultados is not None and niveles is not None:
        cur = Cursor(doc, resultados)
        normados = [c for c in ctx.contaminantes if c != COV]
        cur.parrafo(f"En el siguiente numeral se presentan los resultados reportados para {ctx.nombres_contaminantes()}"
                    + (f", su respectiva comparación normativa con la Resolución 2254 de 2017" if normados else "")
                    + ", el análisis estadístico y el cálculo del Índice de Calidad del Aire (ICA).")
        cur.parrafo("En el Anexo 4 se relacionan las respectivas memorias de cálculo correspondientes a cada parámetro "
                    "evaluado.")
        cur = Cursor(doc, niveles)
        for c in ctx.contaminantes:
            if c in (PM10, PM25, SO2):
                _manual(cur, ctx, c, carpeta)
            elif c in (NO2, CO, O3):
                _automatico(cur, ctx, c, carpeta)
            elif c == COV:
                _cov(cur, ctx)
    else:
        faltantes.append("Capítulo de resultados de la plantilla")
    estad = next((el for el in _cuerpo(doc) if _es_titulo(el, 2, "analisis estadistico")), None)
    if estad is not None:
        _estadistica(Cursor(doc, estad), ctx)
    conteos = {}
    ica_t = next((el for el in _cuerpo(doc) if _es_titulo(el, 2, "indice de calidad")), None)
    if ica_t is not None:
        conteos, _ = _ica(Cursor(doc, ica_t), ctx, carpeta)
    _conclusiones(doc, ctx, conteos)
    _renumerar(doc)
    _anio_fuentes(doc, SimpleNamespace(fecha=ctx.fecha))
    actualizar_indices_al_abrir(doc)
    quitar_relaciones_huerfanas(doc)
    if not (getattr(datos, "area_estudio", "") or "").strip():
        faltantes.append("Datos del informe: no se indicó el área de estudio; se usó el nombre del proyecto")
    doc.save(ruta_salida)
    return ruta_salida, faltantes
