"""Exportacion de los resultados de calidad del aire a Excel."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .aire import (
    AUTOMATICOS, CO, COMPUESTOS_COV, LIMITES, MANUALES, NO2, NOMBRES, O3, ResultadosAire,
)

_TITULO = Font(bold=True, size=12)
_NEGRITA = Font(bold=True)
_ENCABEZADO = PatternFill("solid", fgColor="D9E1F2")
_FUERA = PatternFill("solid", fgColor="F8CBAD")
_BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FMT = "#,##0.00"
_FMT_LC = '"<"#,##0.00'
_FMT_COV = '"<"#,##0.000'

_FILAS_EST = (
    ("Número de datos", "n", "0"), ("Promedio aritmético", "promedio", None), ("Desviación estándar", "desviacion", None),
    ("Coeficiente de variación", "cv", None), ("Mediana (m)", "mediana", None), ("Q1", "q1", None), ("Q3", "q3", None),
    ("IQR=Q3-Q1", "iqr", None), ("m-3*IQR", "lim_inf", None), ("m+3*IQR", "lim_sup", None),
    ("Máxima concentración reportada", "maximo", None), ("Fecha máxima", "fecha_max", "yyyy-mm-dd"),
    ("Mínima concentración reportada", "minimo", None), ("Fecha mínima", "fecha_min", "yyyy-mm-dd"),
    ("Concentraciones atípicas", "atipicos", "0"),
)


def _encabezado(ws, fila, valores, col=1):
    for i, v in enumerate(valores):
        c = ws.cell(fila, col + i, v)
        c.font, c.fill, c.border, c.alignment = _NEGRITA, _ENCABEZADO, _BORDE, _CENTRO


def _celda(ws, fila, col, valor, fmt=None, relleno=None):
    c = ws.cell(fila, col, valor)
    c.border = _BORDE
    c.alignment = Alignment(horizontal="center")
    if fmt and isinstance(valor, (int, float)):
        c.number_format = fmt
    elif fmt and fmt.startswith("yyyy"):
        c.number_format = fmt
    if relleno:
        c.fill = relleno
    return c


def _anchos(ws, primera=16, resto=22):
    ws.column_dimensions["A"].width = primera
    for i in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = resto


def _tabla_diaria(ws, fila, titulo, estaciones, series_por_estacion, valor, limite, fmt_por_estacion):
    """Tabla fecha x estacion (como las del informe) y su resumen."""
    ws.cell(fila, 1, titulo).font = _NEGRITA
    fila += 1
    _encabezado(ws, fila, ["Fecha final de monitoreo", *[e.nombre for e in estaciones], "Límite (µg/m³)"])
    fechas = sorted({f for s in series_por_estacion.values() for f, _ in valor(s)})
    for f in fechas:
        fila += 1
        _celda(ws, fila, 1, f, "yyyy-mm-dd")
        for j, e in enumerate(estaciones, start=2):
            s = series_por_estacion.get(e.numero)
            v = dict(valor(s)).get(f) if s else None
            relleno = _FUERA if v is not None and limite is not None and v > limite else None
            _celda(ws, fila, j, v if v is not None else "---", fmt_por_estacion(s), relleno)
        _celda(ws, fila, len(estaciones) + 2, limite)
    for etiqueta, funcion in (("Promedio aritmético", lambda xs: sum(xs) / len(xs)), ("Máximo", max), ("Mínimo", min)):
        fila += 1
        _celda(ws, fila, 1, etiqueta).font = _NEGRITA
        for j, e in enumerate(estaciones, start=2):
            s = series_por_estacion.get(e.numero)
            xs = [v for _, v in valor(s) if v is not None] if s else []
            _celda(ws, fila, j, funcion(xs) if xs else "---", fmt_por_estacion(s))
    return fila + 2


def _tabla_estadistica(ws, fila, titulo, columnas):
    """`columnas`: [(encabezado, Estadistica|None, formato)]."""
    ws.cell(fila, 1, titulo).font = _NEGRITA
    fila += 1
    _encabezado(ws, fila, ["Variable estadística", *[c[0] for c in columnas]])
    for etiqueta, attr, fmt_fijo in _FILAS_EST:
        fila += 1
        _celda(ws, fila, 1, etiqueta).alignment = Alignment(horizontal="left")
        for j, (_, est, fmt) in enumerate(columnas, start=2):
            v = getattr(est, attr) if est else None
            _celda(ws, fila, j, v if v is not None else "---", fmt_fijo or fmt)
    return fila + 2


def exportar_resultados_aire(res: ResultadosAire, ruta: str):
    wb = Workbook()
    estaciones = res.estaciones()

    # ------------------------------------------------------------- resumen
    ws = wb.active
    ws.title = "Resumen"
    ws.cell(1, 1, f"Calidad del aire – {res.proyecto.nombre_proyecto or res.proyecto.codigo}").font = _TITULO
    _encabezado(ws, 3, ["Contaminante", "Estación", "Tiempo de exposición", "Datos", "Promedio (µg/m³)",
                        "Máximo (µg/m³)", "Límite (µg/m³)", "Excedencias", "Muestras válidas (%)"])
    fila = 3
    for c in MANUALES:
        for e in estaciones:
            s = res.manuales.get((c, e.numero))
            est = s.estadistica() if s else None
            if not est:
                continue
            limite = LIMITES[c]["24h"]
            fila += 1
            fmt = _FMT_LC if s.bajo_lc else _FMT
            for j, v in enumerate([c, e.nombre, "24 horas", est.n, est.promedio, est.maximo, limite,
                                   sum(1 for _, x in s.datos if x > limite), s.pct_validas], start=1):
                _celda(ws, fila, j, v, fmt if j in (5, 6) else ("0.0" if j == 9 else None))
    for c in AUTOMATICOS:
        for e in estaciones:
            s = res.automaticos.get((c, e.numero))
            if not s:
                continue
            for exp, attr in (("1 hora", "max_horario"), ("8 horas", "max_8h")):
                limite = LIMITES[c].get("1h" if exp == "1 hora" else "8h")
                if limite is None:
                    continue
                xs = [getattr(d, attr) for d in s.dias if getattr(d, attr) is not None]
                if not xs:
                    continue
                fila += 1
                for j, v in enumerate([c, e.nombre, exp + " (máximo diario)", len(xs), sum(xs) / len(xs), max(xs),
                                       limite, sum(1 for x in xs if x > limite), None], start=1):
                    _celda(ws, fila, j, v, _FMT if j in (5, 6) else None)
    _anchos(ws, 16, 20)
    ws.column_dimensions["B"].width = 42

    # --------------------------------------------------- manuales (24 h)
    for c in MANUALES:
        series = {n: s for (k, n), s in res.manuales.items() if k == c}
        if not series:
            continue
        ws = wb.create_sheet(c)
        ws.cell(1, 1, NOMBRES[c]).font = _TITULO
        fila = _tabla_diaria(ws, 3, "Nivel de inmisión (µg/m³) – concentraciones diarias", estaciones, series,
                             lambda s: s.datos, LIMITES[c]["24h"], lambda s: _FMT_LC if s and s.bajo_lc else _FMT)
        fila = _tabla_estadistica(ws, fila, "Análisis estadístico", [
            (e.nombre, series[e.numero].estadistica(), _FMT_LC if series[e.numero].bajo_lc else _FMT)
            for e in estaciones if e.numero in series])
        ws.cell(fila, 1, "Memoria de cálculo").font = _NEGRITA
        fila += 1
        _encabezado(ws, fila, ["Estación", "Inicio", "Fin", "Filtro / muestra", "Tiempo (min)", "Temp. (°C)",
                               "Presión (mmHg)", "Caudal", "Masa (µg)", "Volumen actual (m³)",
                               "Conc. actual (µg/m³)", "Conc. referencia (µg/m³)", "Válida"])
        for e in estaciones:
            s = series.get(e.numero)
            for m in (s.muestras if s else []):
                fila += 1
                valores = [e.nombre, m.inicio, m.fin, m.codigo, m.minutos, m.temperatura, m.presion, m.caudal,
                           m.masa, m.volumen, m.concentracion_actual, m.concentracion, "SI" if m.valida else "NO"]
                for j, v in enumerate(valores, start=1):
                    fmt = "yyyy-mm-dd hh:mm" if j in (2, 3) else ("0.0000" if j == 8 else _FMT)
                    _celda(ws, fila, j, v, fmt, None if m.valida else _FUERA)
        _anchos(ws)

    # ----------------------------------------------- automaticos (horarios)
    for c in AUTOMATICOS:
        series = {n: s for (k, n), s in res.automaticos.items() if k == c}
        if not series:
            continue
        ws = wb.create_sheet(c)
        ws.cell(1, 1, NOMBRES[c]).font = _TITULO
        fila = 3
        if c in (CO, NO2):
            fila = _tabla_diaria(ws, fila, "Nivel de inmisión (µg/m³) – máximo horario", estaciones, series,
                                 lambda s: [(d.fecha, d.max_horario) for d in s.dias], LIMITES[c]["1h"],
                                 lambda s: _FMT)
        if c in (CO, O3):
            fila = _tabla_diaria(ws, fila, "Nivel de inmisión (µg/m³) – máxima media móvil de 8 horas", estaciones,
                                 series, lambda s: [(d.fecha, d.max_8h) for d in s.dias if d.max_8h is not None],
                                 LIMITES[c]["8h"], lambda s: _FMT)
        columnas = []
        for e in estaciones:
            s = series.get(e.numero)
            if not s:
                continue
            if c in (CO, O3):
                columnas.append((f"{e.nombre} (8 h)", s.estadistica_8h(), _FMT))
            if c in (CO, NO2):
                columnas.append((f"{e.nombre} (1 h)", s.estadistica_1h(), _FMT))
        fila = _tabla_estadistica(ws, fila, "Análisis estadístico (todos los datos horarios / medias móviles)", columnas)
        _anchos(ws)

    # ------------------------------------------------------------------ COV
    if res.cov:
        ws = wb.create_sheet("COV")
        ws.cell(1, 1, NOMBRES["COV"]).font = _TITULO
        fila = 3
        for e in estaciones:
            series = {k: s for (k, n), s in res.cov.items() if n == e.numero}
            if not series:
                continue
            compuestos = [k for k in COMPUESTOS_COV if k in series]
            ws.cell(fila, 1, e.nombre).font = _NEGRITA
            fila += 1
            _encabezado(ws, fila, ["Fecha", *[f"{k} (µg/m³)" for k in compuestos]])
            for f in sorted({m.fecha for s in series.values() for m in s.muestras}):
                fila += 1
                _celda(ws, fila, 1, f, "yyyy-mm-dd")
                for j, k in enumerate(compuestos, start=2):
                    m = next((m for m in series[k].muestras if m.fecha == f), None)
                    _celda(ws, fila, j, m.concentracion if m else "---", _FMT_COV if m and m.bajo_lc else "#,##0.000")
            fila += 2
        _anchos(ws)

    # ------------------------------------------------------------------ ICA
    ws = wb.create_sheet("ICA")
    ws.cell(1, 1, "Índice de calidad del aire (ICA)").font = _TITULO
    fila = 3
    for e in estaciones:
        dias = res.dias_ica(e.numero)
        if not dias:
            continue
        ws.cell(fila, 1, e.nombre).font = _NEGRITA
        fila += 1
        contaminantes = list(dias)
        _encabezado(ws, fila, ["Fecha", *[f"{c} (µg/m³)" for c in contaminantes], *[f"ICA {c}" for c in contaminantes],
                               *[f"Categoría {c}" for c in contaminantes]])
        fechas = sorted({f for xs in dias.values() for f, *_ in xs})
        for f in fechas:
            fila += 1
            _celda(ws, fila, 1, f, "yyyy-mm-dd")
            for j, c in enumerate(contaminantes):
                fila_c = next((x for x in dias[c] if x[0] == f), None)
                _celda(ws, fila, 2 + j, fila_c[1] if fila_c else "---", _FMT)
                _celda(ws, fila, 2 + len(contaminantes) + j, fila_c[2] if fila_c and fila_c[2] is not None else "---", "0")
                _celda(ws, fila, 2 + 2 * len(contaminantes) + j, fila_c[3] if fila_c and fila_c[3] else "---")
        fila += 2
    _anchos(ws, 14, 16)

    if res.advertencias:
        ws = wb.create_sheet("Advertencias")
        for i, a in enumerate(res.advertencias, start=1):
            ws.cell(i, 1, a)
        ws.column_dimensions["A"].width = 140
    wb.save(ruta)
