"""Calidad del aire: lectura de las plantillas FP y calculos (Res. 2254 de 2017)."""
import os
import sys
import tempfile
import unittest
from datetime import date

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core.aire import (  # noqa: E402
    CO, COV, NO2, O3, PM10, PM25, SO2, ErrorPlantillaAire, ProyectoAire, estadistica, ica, leer_pm10_hivol,
    leer_reporte_analizador, percentil_exc, procesar_aire, redondear,
)
from core.meteorologia import _norm  # noqa: E402
from tests import plantillas_aire as fp  # noqa: E402

# Datos horarios reales de CO (ppm) de Finca Los Camachos, 08 y 09/02/2026.
CO_HORAS = [1.62, 1.26, 1.95, 1.39, 1.3, 1.56, 2.67, 2.5, 2.69, 2.47, 1.02, 2.4, 1.44, 2.16, 1.25, 1.16, 2.19,
            2.15, 2.64, 2.82, 2.36, 1.9, 1.04, 1.84, 1.1, 0.9, 1.65, 1.5, 1.87, 1.35, 2.26, 1.68, 2.62, 1.72, 0.99,
            1.84, 1.0, 1.44, 1.95, 1.0, 0.83, 2.74, 1.41, 2.06, 1.49, 1.74, 2.57, 0.18]
NO2_HORAS = [0.8, 1.2, 1.49, 0.9] * 6


class TestCalculos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        rutas = {k: os.path.join(cls.tmp.name, f"{k}.xlsx") for k in ("pm10", "pm25", "so2", "cov", "auto")}
        fp.fp031_pm10(rutas["pm10"])
        fp.fp032_pm25(rutas["pm25"])
        fp.fp033_so2(rutas["so2"])
        fp.fp035_cov(rutas["cov"])
        fp.fp021_automaticos(rutas["auto"], CO_HORAS, NO2_HORAS)
        cls.rutas = rutas
        cls.res = procesar_aire(ProyectoAire(plantillas={
            PM10: rutas["pm10"], PM25: rutas["pm25"], SO2: rutas["so2"], COV: rutas["cov"], "AUTOMATICOS": rutas["auto"],
        }))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_pm10_hivol_como_el_informe(self):
        s = self.res.manuales[(PM10, 1)]
        self.assertEqual([m.fecha for m in s.muestras], [date(2026, 2, 8), date(2026, 2, 9)])
        self.assertEqual([round(m.concentracion, 2) for m in s.muestras], [48.92, 51.39])
        self.assertAlmostEqual(s.muestras[0].volumen, 977.99, places=2)
        self.assertTrue(all(m.valida for m in s.muestras))

    def test_pm25_con_temperatura_ambiente(self):
        m = self.res.manuales[(PM25, 1)].muestras[0]
        # La FP-032 V03 usaba el caudal en lugar de la temperatura (16,26 µg/m3).
        self.assertAlmostEqual(m.concentracion, 16.93, places=2)
        self.assertAlmostEqual(m.minutos, 1430)

    def test_so2_y_cov_bajo_el_limite_de_cuantificacion(self):
        so2 = self.res.manuales[(SO2, 1)]
        self.assertAlmostEqual(so2.muestras[0].concentracion, 9.36, places=2)
        self.assertTrue(so2.bajo_lc)
        self.assertAlmostEqual(self.res.cov[("Tolueno", 1)].muestras[0].concentracion, 7.679, places=3)
        self.assertAlmostEqual(self.res.cov[("m,p-Xileno", 1)].muestras[0].concentracion, 15.358, places=3)
        self.assertTrue(self.res.cov[("Tolueno", 1)].bajo_lc)

    def test_gases_maximo_horario_y_media_movil(self):
        dias = self.res.automaticos[(CO, 1)].dias
        self.assertEqual(round(dias[0].max_horario, 2), 3214.80)  # 2,82 ppm x 1.140
        self.assertEqual(round(dias[0].max_8h, 2), 2472.38)
        # La media movil del 09/02 toma horas del 08/02 (como la FP-021).
        self.assertEqual(redondear(dias[1].max_8h), 2258.63)
        no2 = self.res.automaticos[(NO2, 1)]
        self.assertEqual(round(no2.dias[0].max_horario, 2), 2.80)
        self.assertEqual(no2.nombre_estacion, "")
        self.assertEqual(self.res.automaticos[(CO, 1)].nombre_estacion, "Finca Los Camachos -Tempestuosa")

    def test_redondeo_como_excel(self):
        self.assertEqual(redondear(2258.6249999999995), 2258.63)
        self.assertEqual(redondear(0.125), 0.13)

    def test_ica_y_estaciones(self):
        self.assertEqual([round(v, 2) for v in ica(PM10, 48.92)[:1]], [44.48])
        self.assertEqual(ica(PM25, 16.93)[1], "Aceptable")
        self.assertIsNone(ica("O3", 400))
        dias = self.res.dias_ica(1)
        self.assertEqual(dias[CO][0][3], "Buena")
        self.assertEqual([e.nombre for e in self.res.estaciones()], ["Finca Los Camachos -Tempestuosa E1_Lis_Tmpst"])
        self.assertEqual(self.res.contaminantes(), [PM10, PM25, SO2, NO2, CO, COV])

    def test_plantilla_equivocada(self):
        with self.assertRaises(ErrorPlantillaAire):
            leer_pm10_hivol(self.rutas["pm25"])
        res = procesar_aire(ProyectoAire(plantillas={PM10: self.rutas["so2"]}))
        self.assertTrue(any("FP-031" in a for a in res.advertencias))


class TestInformeWord(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import docx

        from core.aire import EstacionAire
        from core.informe_aire import generar_informe_aire
        from core.models import DatosInforme

        cls.tmp = tempfile.TemporaryDirectory()
        r = {k: os.path.join(cls.tmp.name, f"{k}.xlsx") for k in ("pm10", "so2", "cov", "auto")}
        fp.fp031_pm10(r["pm10"])
        fp.fp033_so2(r["so2"])
        fp.fp035_cov(r["cov"])
        fp.fp021_automaticos(r["auto"], CO_HORAS, NO2_HORAS)
        from PIL import Image

        foto = os.path.join(cls.tmp.name, "foto.jpg")
        Image.new("RGB", (400, 300), "#4a7").save(foto)
        res = procesar_aire(ProyectoAire(
            nombre_proyecto="Planta Norte", codigo="EC-900-26", cliente="Industrias XYZ S.A.S.",
            estaciones=[EstacionAire(1, "Barrio La Esperanza", "E1_Norte", "", "-74.0721", "4.7110", "Zona residencial",
                                     foto_ruta=foto)],
            plantillas={PM10: r["pm10"], SO2: r["so2"], COV: r["cov"], "AUTOMATICOS": r["auto"]}))
        datos = DatosInforme(area_estudio="el área de influencia de la Planta Norte", fecha="2026-03-15",
                             expediente="ANLA: LAM 0001", acto_administrativo="Resolución 10 del 1/01/2025")
        os.environ["ENGINE_BASEMAP"] = "0"
        salida = os.path.join(cls.tmp.name, "informe.docx")
        _, cls.faltantes = generar_informe_aire(res, os.path.join(RAIZ, "templates", "informe_aire_template.docx"),
                                                salida, os.path.join(cls.tmp.name, "g"), datos)
        cls.doc = docx.Document(salida)
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        partes = ["".join(t.text or "" for t in p.iter(f"{w}t")) for p in cls.doc.element.body.iter(f"{w}p")]
        for s_ in cls.doc.sections:
            partes += ["".join(t.text or "" for t in p.iter(f"{w}t")) for p in s_.header._element.iter(f"{w}p")]
        cls.texto = "\n".join(partes)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sin_datos_del_proyecto_de_la_plantilla(self):
        for viejo in ("Ecopetrol", "ECOPETROL", "Lisama", "Llanito", "Camachos", "Rubí", "LAM 2249", "1653"):
            self.assertNotIn(viejo, self.texto, viejo)

    def test_resultados_y_textos(self):
        self.assertIn("Barrio La Esperanza registró niveles de inmisión diarios de PM10 que oscilaron entre 48,92",
                      self.texto)
        self.assertIn("Resolución 10 del 1/01/2025", self.texto)
        self.assertIn("Parámetro: Monóxido de Carbono (CO)", self.texto)
        self.assertIn("<9,36", self.texto)
        self.assertNotIn("Parámetro: Partículas Menores a 2,5", self.texto)  # sin plantilla de PM2.5
        leyendas = [p.text for p in self.doc.paragraphs if p.style.name == "Caption" and p.text.startswith("Tabla")]
        self.assertEqual([int(t.split()[1].rstrip(".")) for t in leyendas], list(range(1, len(leyendas) + 1)))

    def test_tabla_de_estaciones_como_el_informe_ejemplo(self):
        t = next(t for t in self.doc.tables if t.cell(0, 0).text == "Ubicación")
        self.assertEqual(len(t.rows), 5)
        self.assertEqual(t.cell(2, 0).text, "Barrio La Esperanza")
        self.assertEqual(t.cell(3, 3).text, "Descripción del punto de monitoreo")
        self.assertEqual(t.cell(4, 3).text, "Zona residencial")
        self.assertEqual(len(t.cell(3, 0)._tc.xpath(".//pic:pic")), 1)  # foto de la estacion
        self.assertIn("Resolución 1004 del 24 de agosto de 2026", self.texto)
        self.assertNotIn("1096 del 11 de octubre", self.texto)

    def test_meteorologia_faltante_se_informa(self):
        self.assertTrue(any("meteorológicos" in f for f in self.faltantes))


class TestEstadistica(unittest.TestCase):
    def test_cuartiles_como_excel(self):
        datos = [48.92, 51.39, 47.22, 47.80, 49.39, 45.31, 41.85, 28.44, 26.97, 27.45, 29.79, 23.83, 21.02, 21.40,
                 30.46, 23.53, 24.89, 26.53]
        e = estadistica([(date(2026, 2, i + 1), v) for i, v in enumerate(datos)])
        self.assertEqual((round(e.promedio, 2), round(e.desviacion, 2), round(e.mediana, 2)), (34.23, 11.24, 29.12))
        # Datos ya redondeados a 2 decimales: el informe (con todos los decimales) da 24,63 y 47,37.
        self.assertEqual((round(e.q1, 2), round(e.q3, 2)), (24.62, 47.36))
        self.assertEqual(e.atipicos, 0)
        self.assertEqual(percentil_exc([1, 2, 3, 4], 0.25), 1.25)


class TestCoordenadas(unittest.TestCase):
    def test_planas_se_conservan_y_fuera_de_colombia_se_avisa(self):
        from types import SimpleNamespace

        from core.informe_aire import _coordenadas
        c = _coordenadas(SimpleNamespace(longitud="4874000", latitud="2214000"))
        self.assertEqual((c["este"], c["norte"], c["fuera"]), ("4.874.000,000", "2.214.000,000", False))
        self.assertTrue(c["lat"].endswith("N") and c["lon"].endswith("O"))
        c = _coordenadas(SimpleNamespace(longitud="7727100", latitud="13198600"))
        self.assertEqual((c["este"], c["norte"], c["lon"], c["fuera"], c["dd"]),
                         ("7.727.100,000", "13.198.600,000", "---", True, None))
        c = _coordenadas(SimpleNamespace(longitud="-74.0721", latitud="4.7110"))
        self.assertFalse(c["fuera"])
        self.assertNotEqual(c["este"], "---")


class TestMeteorologia(unittest.TestCase):
    def test_encabezados_con_tildes_danadas(self):
        self.assertEqual(_norm("PresiÃ³n absoluta(hpa)"), "presion absoluta(hpa)")
        self.assertEqual(_norm("DirecciÃ³n del viento(Â°)"), "direccion del viento()")


if __name__ == "__main__":
    unittest.main()


class TestLeerReporteAnalizador(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta_ok = os.path.join(cls.tmp.name, "auto_ok.xlsx")
        cls.ruta_vacia = os.path.join(cls.tmp.name, "auto_vacia.xlsx")
        cls.ruta_incompleta = os.path.join(cls.tmp.name, "auto_incompleta.xlsx")

        import openpyxl
        from datetime import datetime
        # Happy path file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Time", "CO (ppm)"])
        ws.append([datetime(2026, 2, 8, 0, 0), 1.5])
        ws.append([datetime(2026, 2, 8, 1, 0), 2.0])
        ws.append(["invalid date", 3.0])
        ws.append([datetime(2026, 2, 8, 3, 0), "invalid value"])
        wb.save(cls.ruta_ok)

        # Empty file
        wb = openpyxl.Workbook()
        wb.save(cls.ruta_vacia)

        # Incomplete columns
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Fecha", "NO2"])
        wb.save(cls.ruta_incompleta)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_leer_reporte_valido(self):
        from datetime import datetime
        serie = leer_reporte_analizador(self.ruta_ok, CO, 1, "Estacion 1")
        self.assertEqual(serie.nombre_estacion, "Estacion 1")
        self.assertEqual(serie.contaminante, CO)
        self.assertEqual(serie.estacion, 1)
        # 1.5 ppm CO * 1140 (factor) = 1710
        # 2.0 ppm CO * 1140 (factor) = 2280
        self.assertEqual(serie.horas, [
            (datetime(2026, 2, 8, 0, 0), 1710.0),
            (datetime(2026, 2, 8, 1, 0), 2280.0)
        ])

    def test_reporte_sin_columnas(self):
        with self.assertRaisesRegex(ErrorPlantillaAire, "El reporte del analizador no tiene las columnas 'Time' y 'CO'"):
            leer_reporte_analizador(self.ruta_vacia, CO, 1, "")

        with self.assertRaisesRegex(ErrorPlantillaAire, "El reporte del analizador no tiene las columnas 'Time' y 'O3'"):
            leer_reporte_analizador(self.ruta_incompleta, O3, 1, "")
