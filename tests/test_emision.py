"""Emision de ruido: calculo (plantilla FP-007) e informe Word (ER-753-25)."""
import datetime as dt
import os
import sys
import tempfile
import unittest

import docx
from docx.text.paragraph import Paragraph

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core.emision import (  # noqa: E402
    EMISION,
    RESIDUAL,
    ItemBarrido,
    cumple_emision,
    estandares_emision,
    nivel_emision,
    procesar_emision,
)
from core.excel_export import exportar_resultados_emision  # noqa: E402
from core.informe_emision import generar_informe_emision  # noqa: E402
from core.models import ArchivoMemoria, DatosInforme, Proyecto, Punto  # noqa: E402
from core.norms import TABLA_RES_0627  # noqa: E402
from engine.tests.memorias import crear_memoria  # noqa: E402

PLANTILLA = os.path.join(RAIZ, "templates", "informe_emision_template.docx")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
INDUSTRIAL = next(s.etiqueta for s in TABLA_RES_0627 if "industriales" in s.etiqueta)
RESIDENCIAL = next(s.etiqueta for s in TABLA_RES_0627 if "residenciales" in s.etiqueta)


class TestCalculo(unittest.TestCase):
    def test_formula_de_emision_como_fp007(self):
        # Valores del informe ER-753-25 (P1 diurno y nocturno).
        self.assertAlmostEqual(nivel_emision(62.8, 45.9), 62.7, places=1)
        self.assertAlmostEqual(nivel_emision(58.5, 54.5), 56.3, places=1)
        # Diferencia <= 3 dB: la emision es del orden del residual (se reporta el residual).
        self.assertEqual(nivel_emision(60.0, 58.4), 58.4)
        # Residual mayor que la medicion: se reporta la medicion.
        self.assertEqual(nivel_emision(55.0, 57.0), 55.0)

    def test_estandares_del_articulo_9(self):
        self.assertEqual(estandares_emision(INDUSTRIAL), (75, 75))
        self.assertEqual(estandares_emision(RESIDENCIAL), (65, 55))
        # Empate cumple (estandar >= emision), a diferencia de ambiental.
        self.assertEqual(cumple_emision(75.0, 75), "Si")
        self.assertEqual(cumple_emision(75.1, 75), "No")


def _proyecto(carpeta):
    p = Proyecto(tipo="emision", nombre_proyecto="Planta Sur", codigo_informe="ER-900-26", cliente="Industrias XYZ S.A.S.",
                 informe=DatosInforme(area_estudio="el área de influencia de la Planta Sur", municipio="Yumbo",
                                      departamento="Valle del Cauca", fecha="2026-01-20"))
    coords = [(4903197.168, 2341245.159), (4903045.22, 2340841.757)]
    for i, (x, y) in enumerate(coords):
        pt = Punto(i + 1, ["Portería", "Parqueadero"][i], sector=INDUSTRIAL, longitud=str(x), latitud=str(y),
                   incertidumbre=0.0043, fuentes="Tránsito de camiones y motos.")
        for esquema, nivel, hora in (("DH", 70 + 8 * i, 10), ("NDH", 60 + i, 22)):
            ruta = os.path.join(carpeta, f"m{i}{esquema}.xlsx")
            crear_memoria(ruta, nivel, nivel + 1, inicio=dt.datetime(2025, 12, 16 + i, hora))
            pt.archivos[esquema] = {EMISION: ArchivoMemoria(EMISION, ruta)}
        pt.archivos["DNH"] = {EMISION: ArchivoMemoria(EMISION, os.path.join(carpeta, "dnh.xlsx"))} if i else {}
        crear_memoria(os.path.join(carpeta, "dnh.xlsx"), 66, 67, inicio=dt.datetime(2025, 12, 20, 11))
        p.puntos.append(pt)
    # Residual medido en el punto 2 de dia: LAeq 76 -> 78 - 76 = 2 dB: del orden del residual.
    crear_memoria(os.path.join(carpeta, "res.xlsx"), 76, 77, inicio=dt.datetime(2025, 12, 17, 12))
    p.puntos[1].archivos["DH"][RESIDUAL] = ArchivoMemoria(RESIDUAL, os.path.join(carpeta, "res.xlsx"))
    for k, leq in enumerate((62, 75, 58)):
        ruta = os.path.join(carpeta, f"b{k}.xlsx")
        crear_memoria(ruta, leq, leq + 1, inicio=dt.datetime(2025, 11, 15, 9, 10 * k))
        p.barrido.append(ItemBarrido(f"Barrido {k + 1}", "Encendido", ruta, seleccionado=k == 1))
    return p


class TestInformeEmision(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.res = procesar_emision(_proyecto(cls.tmp.name))
        cls.salida = os.path.join(cls.tmp.name, "informe.docx")
        _, cls.faltantes = generar_informe_emision(cls.res, PLANTILLA, cls.salida)
        cls.doc = docx.Document(cls.salida)
        cls.texto = "\n".join(Paragraph(el, cls.doc).text for el in cls.doc.element.body.iter(f"{W}p"))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_residual_medido_y_l90(self):
        r = self.res.por_punto[2]["DH"]
        self.assertFalse(r.residual_es_l90)
        self.assertAlmostEqual(r.residual, 76.0)
        self.assertTrue(r.del_orden_del_residual)
        self.assertEqual(r.emision, 76.0)
        r1 = self.res.por_punto[1]["DH"]
        self.assertTrue(r1.residual_es_l90)
        self.assertAlmostEqual(r1.residual, r1.medicion.l90_corregido)

    def test_barrido_ordenado(self):
        self.assertEqual([b.nombre for b in self.res.barrido], ["Barrido 2", "Barrido 1", "Barrido 3"])

    def test_no_queda_el_proyecto_de_la_plantilla(self):
        fuera_meteo = self.texto.split("METEOROLOGÍA")[0] + self.texto.split("RESULTADOS del barrido")[-1]
        for viejo in ("Refinería", "GRB", "Ecopetrol", "nueve (09)", "Portería Planta GLP", "ER-753-25"):
            self.assertNotIn(viejo, fuera_meteo, viejo)

    def test_textos_y_tablas(self):
        self.assertIn("se establecieron dos (2) puntos de monitoreo denominados: Portería y Parqueadero", self.texto)
        self.assertIn("fueron ejecutadas los días 16, 17 y 20 de diciembre de 2025", self.texto)
        self.assertIn("con un límite máximo de emisión de ruido de 75 dB(A)", self.texto)
        self.assertIn("En Parqueadero la diferencia aritmética entre el LRAeq,1h y el LRAeq,1h residual es igual o "
                      "inferior a 3 dB(A)", self.texto)
        self.assertIn("Parqueadero (periodo diurno, día hábil)", self.texto)
        # La jornada diurna de dia no habil duplica sus tablas con la leyenda correspondiente.
        leyendas = [p for p in self.texto.split("\n") if p.startswith("Tabla") and "– Día no hábil" in p]
        self.assertGreaterEqual(len(leyendas), 4, leyendas)
        relacion = next(t for t in self.doc.tables if "ID Punto de monitoreo" in t.rows[0].cells[-1].text or "ID Punto de monitoreo" in " ".join(c.text for c in t.rows[0].cells[:2]))
        self.assertEqual([c.text for c in relacion.rows[2].cells[:2]], ["Portería", "P1"])
        calculo = next(t for t in self.doc.tables if "Medidor 1" in " ".join(c.text for c in t.rows[0].cells[:2]))
        self.assertEqual(calculo.rows[2].cells[0].text, "P1")
        self.assertIn('w:fill="C2D69B"', calculo.rows[3]._tr.xml)  # P2: diferencia <= 3 dB

    def test_portada_y_excel(self):
        cuadros = [" ".join(Paragraph(p, None).text for p in tx.iter(f"{W}p"))
                   for tx in self.doc.element.body.iter(f"{W}txbxContent")]
        self.assertIn("DICIEMBRE 2025", cuadros)
        ruta = os.path.join(self.tmp.name, "r.xlsx")
        exportar_resultados_emision(self.res, ruta)
        import openpyxl

        self.assertEqual(openpyxl.load_workbook(ruta).sheetnames, ["DH", "DNH", "NDH", "Comparacion Norma", "Barrido"])


if __name__ == "__main__":
    unittest.main()
