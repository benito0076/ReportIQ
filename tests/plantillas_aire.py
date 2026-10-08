"""Plantillas FP minimas para las pruebas de calidad del aire.

Reproducen solo las celdas que lee core.aire, con datos de campo reales del
monitoreo EC-042-26 (estacion 1, primeros dias), de modo que los resultados
esperados son los del informe EC-041/042-26."""
from datetime import datetime, time, timedelta

import openpyxl

INICIO = datetime(2026, 2, 7)


def _libro(hojas):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb, [wb.create_sheet(h) for h in hojas]


def fp031_pm10(ruta, estaciones=1):
    """Hi-Vol: calibracion del 07/02 y dos muestras (08/02 = 48,92; 09/02 = 51,39 µg/m3)."""
    wb, hojas = _libro([f"CA-{n}" for n in range(1, estaciones + 1)] + ["RESUMEN"])
    for ws in hojas[:-1]:
        ws["I10"] = "Finca Los Camachos -Tempestuosa E1_Lis_Tmpst"
        ws["E17"], ws["J18"] = 747.8, 23.9
        ws["K24"], ws["K25"] = 0.99975, -0.0054
        for r, (dh, lectura) in enumerate([(2.1, 2.3), (2.5, 2.8), (3, 3.2), (3.3, 3.5), (3.7, 4)], start=36):
            ws.cell(r, 5, dh)
            ws.cell(r, 8, lectura)
        ws["B72"], ws["C72"] = "ID", "Fecha de inicio del monitoreo"
        filas = [
            (time(16, 24), time(15, 35), 17953.56, 17976.73, 3.4, 3.3, 28.5, 30.5, 750.3, 752.5, 71336, 2.8291, 2.8757),
            (time(15, 55), time(15, 20), 17976.73, 18000.13, 3.4, 3.2, 30.5, 30.8, 747.4, 749.6, 71337, 2.8377, 2.8867),
        ]
        for i, f in enumerate(filas):
            r = 74 + i
            ws.cell(r, 2, i + 1)
            ws.cell(r, 3, INICIO + timedelta(days=i))
            ws.cell(r, 4, INICIO + timedelta(days=i + 1))
            for col, v in zip((5, 6, 7, 8), f[:4]):
                ws.cell(r, col, v)
            for col, v in zip((16, 17, 20, 21, 23, 24, 26, 27, 28), f[4:]):
                ws.cell(r, col, v)
    wb.save(ruta)


def fp032_pm25(ruta):
    """Bajo volumen: 08/02 = 394 µg en 23,82 m3 (16,93 µg/m3 con la temperatura ambiente)."""
    wb, (ws,) = _libro(["CA-1"])
    ws["G41"] = "Finca Los Camachos -Tempestuosa E1_Lis_Tmpst"
    ws["B45"], ws["C45"] = "ID", "Inicial"
    ws["B47"], ws["C47"], ws["D47"] = 1, INICIO, time(16, 15)
    ws["E47"], ws["F47"] = INICIO + timedelta(days=1), time(16, 5)
    ws["I47"], ws["J47"], ws["L47"], ws["M47"] = 28.3, 28.9, 750.3, 752.5
    ws["O47"], ws["P47"], ws["V47"], ws["W47"], ws["X47"] = 16.67, 16.64, 71142, 0.13583, 0.136224
    wb.save(ruta)


def fp033_so2(ruta):
    """SO2: masa en el limite de cuantificacion (2,5 µg) -> <9,36 µg/m3 el 08/02."""
    wb, (ws,) = _libro(["CA-1"])
    ws["B10"] = "Finca Los Camachos_E1_Lis_Tmpst"
    ws["B33"], ws["C33"] = "ID", "Fecha de inicio del monitoreo"
    valores = {2: 1, 3: INICIO, 4: INICIO + timedelta(days=1), 5: time(16, 40), 6: time(15, 45), 7: 7188.81,
               8: 7211.89, 10: 200, 11: 195, 14: 28.6, 15: 28.9, 17: 750.3, 18: 752.5, 20: 2086, 21: 2.5}
    for col, v in valores.items():
        ws.cell(35, col, v)
    wb.save(ruta)


def fp035_cov(ruta):
    """COV: tubo de alto flujo en el limite de cuantificacion -> <7,679 µg/m3 de tolueno el 08/02."""
    wb, (ws,) = _libro(["CA1"])
    ws["H27"] = "Finca Los Camachos -Tempestuosa"
    for r, titulo in ((29, "VOC´S  - TOLUENO"), (40, "VOC´S - m,p - XILENO")):
        ws.cell(r, 2, titulo)
        ws.cell(r + 1, 2, "ID")
        fila = r + 2
        valores = {2: 1, 3: datetime(2026, 2, 8), 4: datetime(2026, 2, 8), 5: time(16), 6: time(17), 12: 60,
                   13: 750.3, 14: 752.5, 16: 28.2, 17: 28.5, 31: 66.7, 32: 66.5}
        valores[9] = 0.06 if "XILENO" in titulo else 0.03
        for col, v in valores.items():
            ws.cell(fila, col, v)
    wb.save(ruta)


def fp021_automaticos(ruta, horas_co, horas_no2):
    """FP-021 con una estacion de CO (ppm) y NO2 (ppb): fecha solo en la primera hora de cada dia."""
    wb, (co, no2) = _libro(["ESTACION 1_CO", "ESTACIÓN 1_NO2"])
    co["B5"], co["G5"] = "Nombre de la estación de monitoreo", "Finca Los Camachos -Tempestuosa"
    co["C8"], co["D8"], co["E8"] = "Fecha ", "Hora ", "Concentraciones de CO (ppm)"
    for i, v in enumerate(horas_co):
        t = datetime(2026, 2, 8) + timedelta(hours=i)
        if t.hour == 0:
            co.cell(9 + i, 3, t)
        co.cell(9 + i, 4, datetime(1900, 1, 1, t.hour))
        co.cell(9 + i, 5, str(v))  # el analizador exporta algunos valores como texto
    no2["B9"], no2["C9"], no2["D9"] = "Fecha ", "Hora ", "Concentración de NO₂ (PPB)"
    for i, v in enumerate(horas_no2):
        t = datetime(2026, 2, 8) + timedelta(hours=i)
        if t.hour == 0:
            no2.cell(10 + i, 2, t)
        no2.cell(10 + i, 3, time(t.hour))
        no2.cell(10 + i, 4, v)
    wb.save(ruta)


def reporte_analizador(ruta, datos, time_col="Time", gas_col="CO ppm"):
    """Exportacion de un analizador automatico."""
    wb, (ws,) = _libro(["Hoja1"])
    cols = []
    if time_col:
        cols.append(time_col)
    if gas_col:
        cols.append(gas_col)
    for c, val in enumerate(cols, start=1):
        ws.cell(1, c, val)

    for r, f in enumerate(datos, start=2):
        for c, val in enumerate(f, start=1):
            ws.cell(r, c, val)

    wb.save(ruta)
