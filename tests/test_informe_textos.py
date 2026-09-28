"""Redaccion automatica del informe Word: portada, encabezado, resumen,
objetivos, cliente, analisis de resultados, fuentes de ruido y conclusiones."""
import datetime as dt
import os
import re
import sys
import tempfile
import unittest

import docx
from docx.text.paragraph import Paragraph

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core import informe_textos as it  # noqa: E402
from core.models import DIRECCIONES, ArchivoMemoria, DatosInforme, Proyecto, Punto  # noqa: E402
from core.pipeline import procesar_proyecto  # noqa: E402
from core.report_generator import generar_informe  # noqa: E402
from engine.tests.memorias import crear_memoria  # noqa: E402

PLANTILLA = os.path.join(RAIZ, "templates", "informe_template.docx")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SECTOR_D = ("D - Zona suburbana o rural de tranquilidad y ruido moderado - Rural Habitada destinada a "
            "explotacion agropecuaria.")
SECTOR_C = "C - Ruido Intermedio Restringido - Zonas con usos permitidos de oficinas."
INICIO = {"DH": dt.datetime(2025, 7, 30, 10), "DNH": dt.datetime(2025, 8, 2, 10),
          "NDH": dt.datetime(2025, 7, 31, 1), "NDNH": dt.datetime(2025, 8, 2, 22)}


def _resultados(carpeta, niveles, sectores, fuentes=None):
    proyecto = Proyecto(
        nombre_proyecto="Planta Norte", codigo_informe="ER-900-25", cliente="Aguas del Valle S.A.S.",
        informe=DatosInforme(area_estudio="el área de la Planta Norte", municipio="Tuluá",
                             departamento="Valle del Cauca", titulo="Planta Norte\nAguas del Valle S.A.S.",
                             expediente="EXP-12", version="1.1", fecha="2025-08-20", cliente_nit="900.1",
                             elaboro_nombre="Ana Pérez", elaboro_cargo="Coordinadora"))
    coords = [(4919855.0, 2309786.0), (4921029.0, 2305271.0), (4919530.0, 2307990.0)]
    for i, sector in enumerate(sectores):
        punto = Punto(no_punto=i + 1, nombre=f"P{i + 1}", sector=sector, longitud=str(coords[i][0]),
                      latitud=str(coords[i][1]), incertidumbre=0.004,
                      fuentes=(fuentes or [""] * len(sectores))[i])
        for esquema, valores in niveles.items():
            for k, direccion in enumerate(DIRECCIONES):
                ruta = os.path.join(carpeta, f"m{i}_{esquema}_{direccion}.xlsx")
                crear_memoria(ruta, valores[i], valores[i] + 1,
                              inicio=INICIO[esquema] + dt.timedelta(hours=2 * i, minutes=13 * k))
                punto.archivos[esquema][direccion] = ArchivoMemoria(direccion, ruta)
        proyecto.puntos.append(punto)
    return procesar_proyecto(proyecto)


def _textos(doc):
    return [Paragraph(el, doc).text for el in doc.element.body.iter(f"{W}p")]


class TestRedaccion(unittest.TestCase):
    def test_numeros_y_fechas(self):
        self.assertEqual(it.cantidad(5, "punto", "puntos"), "cinco (5) puntos")
        self.assertEqual(it.cantidad(1, "punto", "puntos"), "un (1) punto")
        self.assertEqual(it.numero_en_letras(34), "treinta y cuatro")
        self.assertEqual(it.texto_dias([dt.date(2026, 6, 28), dt.date(2026, 6, 29), dt.date(2026, 6, 30)]),
                         "los días 28, 29 y 30 de junio de 2026")
        self.assertEqual(it.texto_dias([dt.date(2025, 7, 31), dt.date(2025, 8, 1)]),
                         "los días 31 de julio y 1 de agosto de 2025")
        self.assertEqual(it.texto_dias([dt.date(2025, 12, 31), dt.date(2026, 1, 1)]),
                         "los días 31 de diciembre de 2025 y 1 de enero de 2026")

    def test_fuentes_frecuentes_ignora_negaciones(self):
        puntos = [Punto(1, "A", fuentes="Aves y tránsito de motos. No se registraron actividades de construcción."),
                  Punto(2, "B", fuentes="Canto de aves y ganado.")]
        frecuentes = it.fuentes_frecuentes(puntos)
        self.assertEqual(frecuentes[0], "la fauna local (aves, insectos y otros animales)")
        self.assertNotIn("actividades de construcción", frecuentes)


class TestInformeCompleto(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        niveles = {"DH": [50.0, 70.0, 60.0], "DNH": [52.0, 66.0, 61.0],
                   "NDH": [40.0, 50.0, 44.0], "NDNH": [41.0, 49.0, 46.0]}
        res = _resultados(cls.tmp.name, niveles, [SECTOR_D, SECTOR_C, SECTOR_D],
                          ["Se percibieron aves y el paso de vehículos.", "Maquinaria de la planta.", ""])
        cls.salida = os.path.join(cls.tmp.name, "informe.docx")
        _, cls.faltantes = generar_informe(res, PLANTILLA, cls.salida)
        cls.doc = docx.Document(cls.salida)
        cls.texto = "\n".join(_textos(cls.doc))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_no_queda_texto_del_proyecto_de_la_plantilla(self):
        # Sin archivo meteorologico el capitulo de meteorologia queda como en la plantilla (se avisa).
        meteo = {Paragraph(el, self.doc).text for el in it.seccion(self.doc, "meteorolog", 1) if el.tag == f"{W}p"}
        texto = "\n".join(t for t in _textos(self.doc) if t not in meteo)
        for viejo in ("Colorado", "Ecopetrol", "ECOPETROL", "cinco (5)", "San Vicente", "ER 731-26", "RA34"):
            self.assertNotIn(viejo, texto, viejo)
        self.assertTrue(any("meteorolog" in f.lower() for f in self.faltantes), self.faltantes)

    def test_resumen_y_objetivos(self):
        self.assertIn("se monitorearon tres (3) puntos, los cuales se encuentran ubicados en el municipio de "
                      "Tuluá en el departamento de Valle del Cauca", self.texto)
        self.assertIn("se realizaron los días 30 y 31 de julio y 2 y 3 de agosto de 2025", self.texto)
        self.assertIn("con código ER-900-25", self.texto)
        self.assertIn("Determinar los niveles de presión sonora de ruido ambiental en tres (3) puntos de "
                      "monitoreo, ubicados en el área de la Planta Norte.", self.texto)
        # Dos sectores distintos: se indica el de cada punto.
        self.assertIn("(puntos P1 y P3)", self.texto)
        self.assertIn("(punto P2)", self.texto)

    def test_analisis_con_cumplimiento_parcial(self):
        self.assertIn("se encontró que 4 de los 6 resultados superaron dicho estándar", self.texto)
        self.assertIn("P2 en día hábil (70,0 dB(A))", self.texto)
        self.assertIn("(Ver Gráfica 5).", self.texto)
        self.assertIn("En P2 se registraron niveles más altos durante la jornada hábil, mientras que en P1 y P3 los "
                      "valores más altos correspondieron a la jornada no hábil.", self.texto)

    def test_fuentes_y_conclusiones(self):
        self.assertIn("En el punto P1 se percibieron aves y el paso de vehículos.", self.texto)
        self.assertIn("En el punto P2 se identificaron las siguientes fuentes de ruido: maquinaria de la planta.",
                      self.texto)
        self.assertTrue(any("P3" in f for f in self.faltantes), self.faltantes)
        self.assertIn("para Aguas del Valle S.A.S. en el área de la Planta Norte, se concluye que:", self.texto)
        self.assertIn("resultados superaron los estándares establecidos", self.texto)

    def test_portada_encabezado_y_control(self):
        cuadros = [" ".join(Paragraph(p, None).text for p in tx.iter(f"{W}p"))
                   for tx in self.doc.element.body.iter(f"{W}txbxContent")]
        self.assertIn("PLANTA NORTE AGUAS DEL VALLE S.A.S.", cuadros)
        self.assertIn("SEGUNDO SEMESTRE 2025", cuadros)
        encabezado = " ".join(c.text for t in self.doc.sections[0].header.tables for f in t.rows for c in f.cells)
        self.assertIn("PLANTA NORTE - AGUAS DEL VALLE S.A.S.", encabezado)
        self.assertIn("EXP-12", encabezado)
        self.assertIn("20/08/2025", encabezado)
        self.assertIn("Versión:  1.1", encabezado)
        control = [[c.text for c in f.cells] for f in self.doc.tables[0].rows]
        self.assertEqual(control[0][1], "1.1")
        self.assertEqual(control[2][:3], ["Ana Pérez", "Coordinadora", "2025-08-20"])
        self.assertIn("updateFields", self.doc.settings.element.xml)

    def test_cliente_y_anio_de_fuentes(self):
        self.assertIn("La Tabla 1 relaciona la información general de Aguas del Valle S.A.S.", self.texto)
        self.assertIn("Fuente: Aguas del Valle S.A.S., 2025.", self.texto)
        self.assertIn("Fuente: Ambienciq Ingenieros S.A.S., 2025.", self.texto)
        self.assertIn("Fuente: Resolución 627 de 2006, MAVDT.", self.texto)
        self.assertNotRegex(self.texto, r"Fuente: Ambienciq Ingenieros S\.A\.S\., 2026")

    def test_tabla_de_estandares_resalta_los_sectores_usados(self):
        _, tabla = it.encontrar_tabla(self.doc, ["sector", "subsector", "dia", "noche"], filas_encabezado=2)
        resaltadas = []
        for fila in list(tabla.rows)[2:]:
            shd = fila._tr.tc_lst[1].find(f".//{W}shd")
            if shd is not None and shd.get(f"{W}fill") == it.COLOR_RESALTADO:
                resaltadas.append(fila.cells[1].text)
        self.assertEqual(len(resaltadas), 2, resaltadas)
        self.assertTrue(any("oficinas" in r for r in resaltadas))
        self.assertTrue(any("agropecuaria" in r for r in resaltadas))
        # Valores de la Res. 0627 (sector B residencial: 65/50).
        residencial = next(f for f in tabla.rows if "residenciales" in f.cells[1].text)
        self.assertEqual([residencial.cells[2].text, residencial.cells[3].text], ["65", "50"])


if __name__ == "__main__":
    unittest.main()
