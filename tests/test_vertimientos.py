"""Vertimientos: reporte de laboratorio, FP-004 y comparacion con la Resolucion 0631 de 2015."""
import os
import re
import sys
import tempfile
import unittest
from datetime import time

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

import openpyxl  # noqa: E402

from core import res0631  # noqa: E402
from core.vertimientos import (  # noqa: E402
    ProyectoVertimiento, PuntoVertimiento, interpretar_reporte, leer_fp004, numero, parametro_de,
    procesar_vertimiento,
)

# Texto con el espaciado del reporte FT-024 (datos ficticios).
REPORTE = """
                                         REPORTE DE RESULTADOS No. 9001-26
                                               INFORMACIÓN DEL CLIENTE
   RAZÓN SOCIAL:                  EMPRESA DE PRUEBA SAS              PROYECTO:            COT-001-26
   NIT:                           900000000-1                        SOLICITADO POR:      Persona de Contacto
                                               INFORMACIÓN DE LA MUESTRA
   NÚMERO DE MUESTRA:             9001-26                            TIPO DE MUESTREO:     Compuesta
   PUNTO DE MUESTREO:             Salida PTAR                        PLAN Y MÉTODO DE MUESTREO:   EA-900-26
   CIUDAD/MUNICIPIO               BOGOTA                             FECHA Y HORA DE MUESTREO:    2026-03-12 19:00
   DEPARTAMENTO                   Cundinamarca                       FECHA Y HORA DE RECEPCIÓN:   2026-03-13 08:00
   TIPO DE MUESTRA:               Agua Residual no Doméstica         FECHA DE EMISIÓN DEL INFORME: 2026-04-17
                                                     RESULTADOS
     ENSAYO REALIZADO                     MÉTODO               LCM       FECHA DE          REPORTE     INCERTIDUMBRE     UNIDAD
                                                                         ANÁLISIS                         (μc),%
   ACEITES Y GRASAS                      SM 5520 C             4,00      2026-03-31          <4,00          2,6          mg/L
   DEMANDA BIOQUÍMICA DE OXÍGENO      SM 5210 B, SM 4500-O G   3,00      2026-03-18          180            4,2        mg O2/L
   (DBO5)
   HAPS - BENZO(A)PIRENO                SM 6440 B, C          0,00400    2026-03-24        <0,00400         3,3          mg/L
FT-024 Versión 13                Cra 28 N° 75-37 PBX: 2310500 Bogotá D.C., Colombia                  Página 1 de 2
                                                     RESULTADOS
   NITRÓGENO TOTAL                 Análisis de Aguas. J. Rodier   0,602      2026-03-31          105             8          mg N/L
                                             Numeral 9.6
   pH MÁXIMO                            SM 4500-H+ B            2,00      2026-03-12          9,40           2,2      Unidades de pH
   pH MÍNIMO                            SM 4500-H+ B            2,00      2026-03-12          6,10           2,2      Unidades de pH
   TEMPERATURA  MÁXIMO                    SM 2550 B             N.A       2026-03-12          18,8           1,8           °C
   LCM: Límite de cuantificación del Método / N.E: No especifica
                                               ENSAYOS SUBCONTRATADOS
     ENSAYO REALIZADO                     MÉTODO               LCM       FECHA DE          REPORTE     INCERTIDUMBRE     UNIDAD
   CIANURO TOTAL*                     SM 4500 CN - B, C, F      0,1      2026-03-17          0,30          0,092       mg CN-/L
   OBSERVACIONES:
""".splitlines()


def crear_fp004(ruta: str):
    wb = openpyxl.Workbook()
    wb.active.title = "FP-004 V02"
    for nombre, titulo in (("Entrada", "ENTRADA DEL SISTEMA DE TRATAMIENTO"), ("Salida", "SALIDA DEL SISTEMA")):
        ws = wb.create_sheet(nombre)
        ws["B3"] = titulo
        for c, t in zip("BCDE", ("Hora", "Volumen (mL)", "Tiempo (s)", "Caudal (mL/s)")):
            ws[f"{c}4"] = t
        ws["P24"], ws["Q24"] = "Tamaño de muestra (mL)", 8000
        for i, (h, vol, t, ph, temp, ss) in enumerate([(time(10, 0), 3000, 10, 7.1, 18.2, 0.1),
                                                       (time(10, 30), 2000, 10, 7.4, 18.6, None),
                                                       (time(11, 0), 3000, 20, 7.0, 18.4, 1.3)]):
            r = 6 + i
            ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"], ws[f"H{r}"] = h, vol, t, ph, temp
            if ss is not None:
                ws[f"N{r}"] = ss
                ws[f"N{r}"].number_format = "\\<0.0" if ss == 0.1 else "0.0"
        ws["B23"] = "Suma"
    wb.save(ruta)


class TestReporte(unittest.TestCase):
    def test_reporte_de_resultados(self):
        inf = interpretar_reporte(REPORTE)
        self.assertEqual((inf.numero, inf.muestra, inf.punto, inf.tipo_muestreo), ("9001-26", "9001-26",
                                                                                 "Salida PTAR", "Compuesta"))
        self.assertEqual(inf.fecha_muestreo.isoformat(), "2026-03-12T19:00:00")
        self.assertEqual(inf.cliente["nit"], "900000000-1")
        claves = [r.parametro.clave for r in inf.resultados]
        self.assertEqual(claves, ["aceites_grasas", "dbo5", "hap_benzo_a_pireno", "nitrogeno_total", "ph_max",
                                  "ph_min", "temperatura_max", "cianuro"])
        dbo = inf.resultado("dbo5")
        self.assertEqual((dbo.ensayo, dbo.metodo, dbo.reporte, dbo.unidad),
                         ("DEMANDA BIOQUÍMICA DE OXÍGENO (DBO5)", "SM 5210 B, SM 4500-O G", "180", "mg O2/L"))
        self.assertEqual(inf.resultado("nitrogeno_total").metodo, "Análisis de Aguas. J. Rodier Numeral 9.6")
        self.assertEqual(inf.resultado("hap_benzo_a_pireno").parametro.nombre, "Benzo(a)pireno")
        self.assertTrue(inf.resultado("aceites_grasas").menor_que_lcm)
        self.assertTrue(inf.resultado("cianuro").subcontratado)

    def test_numeros_y_parametros(self):
        self.assertEqual(numero("<0,00100"), 0.001)
        self.assertEqual(numero("1.200,50"), 1200.5)
        self.assertEqual(parametro_de("CONDUCTIVIDAD ELÉCTRICA MÁXIMO").clave, "conductividad_max")
        self.assertEqual(parametro_de("SURFACTANTES ANIÓNICOS COMO SAAM").clave, "saam")


class TestRes0631(unittest.TestCase):
    def test_limites_y_articulo_16(self):
        lim = res0631.limite("art12_alimentos_animales", "dbo5")
        self.assertEqual((lim.texto, lim.maximo), ("100,00", 100.0))
        aj = res0631.limite("art12_alimentos_animales", "dbo5", alcantarillado=True)
        self.assertEqual((aj.texto, aj.marca), ("150,00", "*"))
        self.assertEqual(res0631.limite("art12_alimentos_animales", "cianuro", True).texto, "0,20")
        self.assertEqual(res0631.limite("art12_alimentos_animales", "ph_min", True).minimo, 5.0)
        self.assertEqual(res0631.limite("art13_sabores_fragancias", "formaldehido").tipo, "ayr")
        self.assertEqual(res0631.limite("art13_sabores_fragancias", "formaldehido", True).tipo, "ne")
        self.assertEqual(res0631.limite("art13_sabores_fragancias", "temperatura_max").maximo, 40.0)
        self.assertEqual(res0631.limite("art10_carbon", "hap_naftaleno", consumo_humano=True).maximo, 0.01)

    def test_evaluacion(self):
        lim = res0631.limite("art12_alimentos_animales", "ph_max")
        self.assertEqual(res0631.evaluar("ph_max", 9.4, False, lim), "No cumple")
        self.assertEqual(res0631.evaluar("ph_min", 6.1, False, lim), "Cumple")
        self.assertEqual(res0631.evaluar("dbo5", 4.0, True, res0631.limite("art12_alimentos_animales", "dbo5")),
                         "Cumple")
        self.assertIsNone(res0631.evaluar("acidez", 10, False, res0631.AYR))

    def test_catalogo_completo(self):
        self.assertEqual(len(res0631.ACTIVIDADES), 76)
        self.assertEqual({g["articulo"] for g in res0631.actividades_por_articulo()}, set(range(8, 16)))


class TestProcesamiento(unittest.TestCase):
    def test_fp004_y_comparacion(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, "fp004.xlsx")
            crear_fp004(ruta)
            campos = leer_fp004(ruta)
            self.assertEqual([c.hoja for c in campos], ["Entrada", "Salida"])
            sal = campos[1]
            self.assertEqual([round(m.caudal_mls, 1) for m in sal.mediciones], [300.0, 200.0, 150.0])
            self.assertEqual([m.sedimentables for m in sal.mediciones], ["<0,1", None, "1,3"])
            self.assertAlmostEqual(sal.alicuota_ml(sal.mediciones[0]), 8000 * 300 / 650)

            import core.vertimientos as v
            original = v.leer_informe_laboratorio
            v.leer_informe_laboratorio = lambda _ruta: interpretar_reporte(REPORTE)
            try:
                res = procesar_vertimiento(ProyectoVertimiento(
                    puntos=[PuntoVertimiento("Salida sistema de tratamiento", informe_pdf="x.pdf")], fp004=ruta,
                    actividades=["art12_alimentos_animales"], alcantarillado=True))
            finally:
                v.leer_informe_laboratorio = original
        self.assertEqual(res.puntos[0].campo.hoja, "Salida")
        self.assertEqual([c.titulo for c in res.columnas],
                         ["Art. 12 Elaboración de alimentos preparados para animales", "Art. 12 Ajustado al Art. 16"])
        incumple = {f.parametro.clave for _, f, _ in res.incumplimientos()}
        self.assertEqual(incumple, {"dbo5", "ph_max", "cianuro"})  # DBO 180 > 150; pH 9,4 > 9; CN 0,30 > 0,20
        self.assertTrue(any("Salida PTAR" in a for a in res.advertencias))


if __name__ == "__main__":
    unittest.main()


class TestInformeWord(unittest.TestCase):
    def test_informe_word(self):
        import docx
        from types import SimpleNamespace

        import core.vertimientos as v
        from core.informe_vertimientos import generar_informe_vertimiento, minus

        os.environ["ENGINE_BASEMAP"] = "0"
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, "fp004.xlsx")
            crear_fp004(ruta)
            original = v.leer_informe_laboratorio
            v.leer_informe_laboratorio = lambda _ruta: interpretar_reporte(REPORTE)
            try:
                res = procesar_vertimiento(ProyectoVertimiento(
                    codigo="EA-900-26", cliente="EMPRESA DE PRUEBA SAS",
                    puntos=[PuntoVertimiento("Salida sistema de tratamiento", informe_pdf="x.pdf", latitud="4.96498",
                                             longitud="-73.95357", descripcion="Salida de la PTAR.")],
                    fp004=ruta, actividades=["art12_alimentos_animales"], alcantarillado=True))
            finally:
                v.leer_informe_laboratorio = original
            datos = SimpleNamespace(municipio="Bogotá", departamento="Cundinamarca", version="1.0",
                                    fecha="2026-04-28", **{k: "" for k in (
                                        "area_estudio", "titulo", "expediente", "acto_administrativo", "cliente_nit",
                                        "cliente_direccion", "cliente_contacto", "cliente_ciudad",
                                        "cliente_departamento", "cliente_actividad", "elaboro_nombre", "elaboro_cargo",
                                        "autorizo_nombre", "autorizo_cargo")})
            salida = os.path.join(d, "informe.docx")
            plantilla = os.path.join(RAIZ, "templates", "informe_vertimientos_template.docx")
            generar_informe_vertimiento(res, plantilla, salida, os.path.join(d, "graficas"), datos)
            doc = docx.Document(salida)
        texto = "\n".join(p.text for p in doc.paragraphs)
        titulos = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        for t in ("RESULTADOS DE PARÁMETROS IN SITU", "RESULTADOS MEDIDOS EN EL LABORATORIO", "Caudal",
                  "Cianuro total", "DBO5 y DQO", "Hidrocarburos aromáticos policíclicos"):
            self.assertIn(t, titulos)
        self.assertIn("Resolución 1004 del 24 de agosto de 2026", texto)
        self.assertNotIn("1096", texto)
        self.assertIn("cianuro total en el punto «Salida sistema de tratamiento»", texto)
        self.assertIn("Tabla 9. Resultados del análisis de laboratorio vs Resolución 0631 de 2015", texto)
        self.assertIn("La Tabla 9 presenta los límites de referencia", texto)
        self.assertIsNone(re.search(r"\b[Dd]e el\b", texto))
        self.assertEqual(minus("Demanda Bioquímica de Oxígeno (DBO5)"), "demanda bioquímica de oxígeno (DBO5)")


def test_catalogo_web_sincronizado():
    """web/src/lib/res0631-actividades.json (formulario de la web) debe coincidir con core.
    Regenerarlo: python -c "import json; from core import res0631; json.dump(res0631.actividades_por_articulo(),
    open('web/src/lib/res0631-actividades.json', 'w'), ensure_ascii=False, indent=1)"."""
    import json
    import os

    from core import res0631

    ruta = os.path.join(os.path.dirname(__file__), "..", "web", "src", "lib", "res0631-actividades.json")
    with open(ruta, encoding="utf-8") as f:
        assert json.load(f) == res0631.actividades_por_articulo()
