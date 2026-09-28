"""El texto escrito en la plantilla Word conserva su formato: las
descripciones de varias lineas usan la misma fuente y tamano, y las filas
de dias de la tabla meteorologica no heredan el fondo de las de resumen."""
import os
import sys
import unittest
from datetime import date

import docx

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.docx_utils import encontrar_tabla, encontrar_tablas, reconstruir_filas, set_cell_lines  # noqa: E402

PLANTILLA = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates",
                         "informe_template.docx")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _fondo(celda):
    shd = celda._tc.find(f".//{W}shd")
    fill = shd.get(f"{W}fill") if shd is not None else None
    return None if fill in (None, "auto", "FFFFFF") else fill


class TestFormatoWord(unittest.TestCase):
    def test_descripcion_varias_lineas_mismo_formato(self):
        doc = docx.Document(PLANTILLA)
        tabla = encontrar_tablas(doc, ["nombre del punto", "longitud/este", "latitud/norte"], filas_encabezado=1)[0][1]
        celda = tabla.cell(6, 1)
        set_cell_lines(celda, ["Primera linea.", "", "Segunda linea mas larga.", "Altitud: 152 m.s.n.m."])
        parrafos = celda.paragraphs
        self.assertEqual([p.text for p in parrafos],
                         ["Primera linea.", "", "Segunda linea mas larga.", "Altitud: 152 m.s.n.m."])
        formatos = {(p.style.name, p.runs[0].font.size, p.runs[0].font.name) for p in parrafos if p.runs}
        self.assertEqual(len(formatos), 1, formatos)

    def test_tabla_meteorologica_dias_sin_fondo(self):
        doc = docx.Document(PLANTILLA)
        _, tabla = encontrar_tabla(doc, ["fecha de monitoreo", "temperatura", "humedad"], filas_encabezado=1)
        dias = [[date(2025, 5, d).isoformat(), "1", "2", "3", "4", "5"] for d in range(1, 8)]
        resumen = [["Promedio", "1", "2", "3", "4", "5"], ["Máximo", "1", "2", "3", "4", "5"],
                   ["Mínimo", "1", "2", "3", "4", "5"]]
        reconstruir_filas(tabla, 1, dias + resumen, es_resumen=lambda v: v[0] in ("Promedio", "Máximo", "Mínimo"))
        filas = list(tabla.rows)[1:]
        self.assertEqual([f.cells[0].text for f in filas], [d[0] for d in dias] + ["Promedio", "Máximo", "Mínimo"])
        for f in filas[:7]:
            self.assertIsNone(_fondo(f.cells[0]), f.cells[0].text)
        for f in filas[7:]:
            self.assertEqual(_fondo(f.cells[0]), "BAD405", f.cells[0].text)


if __name__ == "__main__":
    unittest.main()
