"""Pruebas de la API del motor: descarga de archivos por URL, procesamiento
y generacion de entregables (Excel, Word y anexos)."""
import functools
import io
import json
import os
import sys
import tempfile
import threading
import unittest
import zipfile
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import unquote

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, RAIZ)
os.environ["ENGINE_API_KEY"] = "clave-de-prueba"
os.environ["ENGINE_BASEMAP"] = "0"

from fastapi.testclient import TestClient  # noqa: E402

from core.corrections import promedio_energetico  # noqa: E402
from engine.app import app  # noqa: E402
from engine.tests.memorias import crear_memoria  # noqa: E402

DIRECCIONES = ["Vertical", "Norte", "Sur", "Este", "Oeste"]
AUTH = {"Authorization": "Bearer clave-de-prueba"}
SECTOR_B = ("B - Tranquilidad y ruido moderado - Zonas residenciales o exclusivamente "
            "destinadas para desarrollo habitacional, hoteleria y hospedajes.")
COORDENADAS = [(4919855.125, 2309786.167), (4921029.181, 2305271.928), (4919539.506, 2308027.992)]


class _Silencioso(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


class TestApiMotor(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        handler = functools.partial(_Silencioso, directory=cls.tmp.name)
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
        cls.base = f"http://127.0.0.1:{cls.server.server_address[1]}"
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.tmp.cleanup()

    def _proyecto(self, n_puntos=3, laeq=55.0):
        puntos = []
        for i in range(n_puntos):
            memorias = {"DH": {}}
            for j, d in enumerate(DIRECCIONES):
                nombre = f"P{i + 1}_{d}.xlsx"
                crear_memoria(os.path.join(self.tmp.name, nombre), laeq=laeq + i + j * 0.5, laieq=laeq + i + 1)
                memorias["DH"][d] = {"url": f"{self.base}/{nombre}", "nombre": nombre}
            este, norte = COORDENADAS[i]
            puntos.append({
                "no_punto": i + 1, "nombre": f"P{i + 1}", "sector": SECTOR_B,
                "este": str(este), "norte": str(norte), "incertidumbre": 0.0043,
                "memorias": memorias,
            })
        return {"nombre_proyecto": "Prueba", "codigo_informe": "ER-001-26", "cliente": "Cliente", "puntos": puntos}

    def test_requiere_clave(self):
        r = self.client.post("/v1/procesar", json={"proyecto": {"puntos": []}})
        self.assertEqual(r.status_code, 401)

    def test_procesar(self):
        r = self.client.post("/v1/procesar", json={"proyecto": self._proyecto()}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        data = r.json()
        self.assertEqual(len(data["puntos"]), 3)
        p1 = data["puntos"][0]
        self.assertEqual(p1["estandar_diurno"], 65)
        dh = p1["esquemas"]["DH"]
        self.assertEqual(len(dh["direcciones"]), 5)
        esperado = promedio_energetico([d["lraeq_corregido"] for d in dh["direcciones"]])
        self.assertAlmostEqual(dh["lraeq_resultante"], esperado, places=3)
        self.assertEqual(p1["cumple"]["DH"], "Si")
        self.assertEqual(data["equipos_detectados"], [{"serial": "0006599", "modelo": "SoundTrack LxT1"}])

    def test_direccion_faltante_genera_advertencia(self):
        proyecto = self._proyecto(n_puntos=1)
        del proyecto["puntos"][0]["memorias"]["DH"]["Oeste"]
        r = self.client.post("/v1/procesar", json={"proyecto": proyecto}, headers=AUTH)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(any("Oeste" in a["direccion"] for a in r.json()["advertencias"]))

    def test_memoria_invalida_muestra_nombre_original(self):
        proyecto = self._proyecto(n_puntos=1)
        with open(os.path.join(self.tmp.name, "roto.xlsx"), "wb") as f:
            f.write(b"no es un excel")
        proyecto["puntos"][0]["memorias"]["DH"]["Norte"] = {"url": f"{self.base}/roto.xlsx", "nombre": "roto.xlsx"}
        r = self.client.post("/v1/procesar", json={"proyecto": proyecto}, headers=AUTH)
        mensajes = [a["mensaje"] for a in r.json()["advertencias"]]
        self.assertTrue(any("roto.xlsx" in m and "ruido_" not in m for m in mensajes), mensajes)

    def test_descarga_fallida(self):
        proyecto = self._proyecto(n_puntos=1)
        proyecto["puntos"][0]["memorias"]["DH"]["Norte"]["url"] = f"{self.base}/no-existe.xlsx"
        r = self.client.post("/v1/procesar", json={"proyecto": proyecto}, headers=AUTH)
        self.assertEqual(r.status_code, 502)

    def test_generar_excel(self):
        r = self.client.post("/v1/generar", json={"proyecto": self._proyecto(), "tipo": "excel"}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(unquote(r.headers["x-nombre-archivo"]), "ER-001-26 - Resultados.xlsx")
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertIn("DH", wb.sheetnames)
        self.assertIn("Comparacion Norma", wb.sheetnames)

    def test_generar_word(self):
        proyecto = self._proyecto()
        proyecto["informe"] = {"area_estudio": "el área de la Planta Sur", "expediente": "EXP-7",
                               "elaboro_nombre": "Ana Pérez"}
        proyecto["puntos"][0]["fuentes"] = "Se percibió el tránsito de camiones."
        cuerpo = {
            "proyecto": proyecto, "tipo": "word",
            "equipos": [{"nombre": "Sonometro X", "codigo": "612-C", "serial": "6599"}],
        }
        r = self.client.post("/v1/generar", json=cuerpo, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        import docx
        doc = docx.Document(io.BytesIO(r.content))
        texto = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertIn("612-C", texto)
        self.assertIn("Ana Pérez", texto)
        parrafos = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("ubicados en el área de la Planta Sur", parrafos)
        self.assertIn("En el punto P1 se percibió el tránsito de camiones.", parrafos)
        self.assertIsInstance(json.loads(unquote(r.headers["x-advertencias"])), list)

    def _proyecto_emision(self):
        from datetime import datetime

        puntos = []
        for i in range(2):
            memorias = {}
            for esquema, hora in (("DH", 10), ("NDH", 22)):
                nombre = f"E{i + 1}_{esquema}.xlsx"
                crear_memoria(os.path.join(self.tmp.name, nombre), laeq=70.0 + i, laieq=71.0 + i,
                              inicio=datetime(2026, 3, 10, hora))
                memorias[esquema] = {"Emision": {"url": f"{self.base}/{nombre}", "nombre": nombre}}
            este, norte = COORDENADAS[i]
            puntos.append({"no_punto": i + 1, "nombre": f"Porteria {i + 1}", "sector": SECTOR_B,
                           "este": str(este), "norte": str(norte), "incertidumbre": 0.0043, "memorias": memorias})
        crear_memoria(os.path.join(self.tmp.name, "residual.xlsx"), laeq=69.0, laieq=70.0,
                      inicio=datetime(2026, 3, 10, 12))
        puntos[0]["memorias"]["DH"]["Residual"] = {"url": f"{self.base}/residual.xlsx", "nombre": "residual.xlsx"}
        crear_memoria(os.path.join(self.tmp.name, "barrido1.xlsx"), laeq=66.0, laieq=67.0)
        return {"tipo": "emision", "nombre_proyecto": "Emision", "codigo_informe": "ER-002-26", "cliente": "Cliente",
                "puntos": puntos,
                "barrido": [{"nombre": "Barrido 1", "seleccionado": True,
                             "archivo": {"url": f"{self.base}/barrido1.xlsx", "nombre": "barrido1.xlsx"}}]}

    def test_emision_procesar_y_generar(self):
        proyecto = self._proyecto_emision()
        r = self.client.post("/v1/procesar", json={"proyecto": proyecto}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        datos = r.json()
        self.assertEqual(datos["tipo"], "emision")
        dh = datos["puntos"][0]["esquemas"]["DH"]
        self.assertEqual(dh["residual_origen"], "medido")
        self.assertTrue(dh["del_orden_del_residual"])  # 70 - 69 = 1 dB
        self.assertEqual(datos["puntos"][1]["esquemas"]["DH"]["residual_origen"], "L90")
        self.assertEqual(datos["puntos"][0]["estandar_diurno"], 65)  # emision, sector B residencial
        self.assertEqual(datos["barrido"][0]["nombre"], "Barrido 1")

        for tipo, ext in (("word", "docx"), ("excel", "xlsx"), ("anexos", "zip")):
            r = self.client.post("/v1/generar", json={"proyecto": proyecto, "tipo": tipo}, headers=AUTH)
            self.assertEqual(r.status_code, 200, r.text)
            self.assertTrue(unquote(r.headers["x-nombre-archivo"]).endswith(f".{ext}"))
        nombres = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
        self.assertIn("graficas/emision_DH.png", nombres)
        self.assertFalse(any(n.startswith("isofonas/") for n in nombres))

    def test_generar_word_plantilla_invalida(self):
        with open(os.path.join(self.tmp.name, "plantilla_mala.docx"), "wb") as f:
            f.write(b"")
        cuerpo = {"proyecto": self._proyecto(n_puntos=1), "tipo": "word",
                  "plantilla": {"url": f"{self.base}/plantilla_mala.docx", "nombre": "p.docx"}}
        r = self.client.post("/v1/generar", json=cuerpo, headers=AUTH)
        self.assertEqual(r.status_code, 422)

    def test_generar_anexos(self):
        r = self.client.post("/v1/generar", json={"proyecto": self._proyecto(), "tipo": "anexos"}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        nombres = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
        self.assertIn("graficas/grafica_diurna.png", nombres)
        self.assertIn("isofonas/isofonas_DH.png", nombres)
        self.assertIn("isofonas/isofonas_DH.pdf", nombres)

    def test_fallo_del_mapa_satelital_se_informa(self):
        from unittest import mock

        import contextily

        from engine import service

        def falla(*_a, **_k):
            raise ConnectionError("sin acceso a server.arcgisonline.com")

        with mock.patch.object(service, "CON_MAPA_BASE", True), \
                mock.patch.object(contextily, "add_basemap", falla):
            for tipo in ("anexos", "word"):
                r = self.client.post("/v1/generar", json={"proyecto": self._proyecto(), "tipo": tipo}, headers=AUTH)
                self.assertEqual(r.status_code, 200, r.text)
                avisos = json.loads(unquote(r.headers["x-advertencias"]))
                satelital = [a for a in avisos if a.startswith("Mapa satelital no disponible")]
                self.assertEqual(len(satelital), 1, avisos)
                self.assertIn("sin acceso a server.arcgisonline.com", satelital[0])

    def test_word_y_anexos_con_datos_meteorologicos(self):
        from datetime import date

        from tests.test_meteorologia import crear_archivo

        # Las memorias de prueba son del 10/03/2026: solo ese dia debe usarse.
        crear_archivo(os.path.join(self.tmp.name, "meteo.xlsx"), dias=(date(2026, 3, 9), date(2026, 3, 10)))
        proyecto = self._proyecto()
        proyecto["meteorologia"] = {"url": f"{self.base}/meteo.xlsx", "nombre": "meteo.xlsx"}

        r = self.client.post("/v1/generar", json={"proyecto": proyecto, "tipo": "word"}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        import docx
        doc = docx.Document(io.BytesIO(r.content))
        tabla = next(t for t in doc.tables if "Fecha de Monitoreo" in t.rows[0].cells[0].text)
        self.assertEqual([row.cells[0].text for row in tabla.rows[1:]], ["2026-03-10", "Promedio", "Máximo", "Mínimo"])
        texto = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("10 de marzo de 2026, único día de monitoreo", texto)
        self.assertNotIn("WRPLOT", texto)
        # Grafica de viento: una sola imagen en linea (la rosa y las clases) y la tabla por direccion llena.
        viento = next(t for t in doc.tables if "Porcentaje (%)" in t._tbl.xml)
        xml = viento._tbl.xml
        self.assertNotIn("<wp:anchor", xml)
        self.assertIn("<wp:inline", xml)
        anidada = next(t for row in viento.rows for c in row.cells for t in c.tables)
        self.assertEqual(anidada.rows[-1].cells[0].text, "Calms")
        self.assertNotEqual(anidada.rows[1].cells[1].text, "2,67")  # ya no es el valor de la plantilla
        avisos = json.loads(unquote(r.headers["x-advertencias"]))
        self.assertTrue(any("se descartaron 5 fila(s)" in a for a in avisos), avisos)

        r = self.client.post("/v1/generar", json={"proyecto": proyecto, "tipo": "anexos"}, headers=AUTH)
        nombres = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
        self.assertIn("meteorologia/meteo_rosa_vientos.png", nombres)
        self.assertIn("meteorologia/meteo_temperatura.png", nombres)

    def test_meteorologia_sin_datos_de_los_dias_de_medicion(self):
        from datetime import date

        from tests.test_meteorologia import crear_archivo

        crear_archivo(os.path.join(self.tmp.name, "meteo_otro.xlsx"), dias=(date(2025, 1, 1),))
        proyecto = self._proyecto()
        proyecto["meteorologia"] = {"url": f"{self.base}/meteo_otro.xlsx", "nombre": "meteo_otro.xlsx"}
        r = self.client.post("/v1/generar", json={"proyecto": proyecto, "tipo": "word"}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        avisos = json.loads(unquote(r.headers["x-advertencias"]))
        self.assertTrue(any("no tiene registros validos para los dias de medicion" in a for a in avisos), avisos)


    def test_calidad_del_aire(self):
        from tests import plantillas_aire as fp

        fp.fp031_pm10(os.path.join(self.tmp.name, "FP031.xlsx"))
        fp.fp033_so2(os.path.join(self.tmp.name, "FP033.xlsx"))
        proyecto = {
            "codigo": "EC-042-26",
            "estaciones": [{"numero": 1, "nombre": "Finca Los Camachos", "codigo": "E1_Lis_Tmpst"}],
            "plantillas": {"PM10": {"url": f"{self.base}/FP031.xlsx", "nombre": "FP031.xlsx"},
                           "SO2": {"url": f"{self.base}/FP033.xlsx", "nombre": "FP033.xlsx"}},
        }
        r = self.client.post("/v1/aire/procesar", json={"proyecto": proyecto}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        datos = r.json()
        self.assertEqual(datos["tipo"], "aire")
        self.assertEqual(datos["contaminantes"], ["PM10", "SO2"])
        self.assertEqual(datos["estaciones"][0]["nombre"], "Finca Los Camachos")
        pm10 = datos["manuales"][0]
        self.assertEqual([m["fecha"] for m in pm10["muestras"]], ["2026-02-08", "2026-02-09"])
        self.assertAlmostEqual(pm10["muestras"][0]["concentracion"], 48.9217, places=3)
        self.assertTrue(datos["manuales"][1]["bajo_lc"])
        self.assertEqual(datos["ica"]["1"]["PM10"][0]["categoria"], "Buena")

        r = self.client.post("/v1/aire/generar", json={"proyecto": proyecto, "tipo": "excel"}, headers=AUTH)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("EC-042-26", unquote(r.headers["X-Nombre-Archivo"]))
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Resumen", "PM10", "SO2", "ICA"])


if __name__ == "__main__":
    unittest.main()
